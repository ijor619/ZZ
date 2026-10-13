"""
Zig_Zag lounge — бот заказов для кальянной (единый файл).

Запуск:      python bot.py
Зависимости: pip install -r requirements.txt
Меню:        menu.json (лежит рядом с bot.py)
Настройки:   переменные окружения BOT_TOKEN, CHANNEL_ID, ADMIN_IDS, CURRENCY
             (на Bothost — раздел «Переменные»; локально — файл .env)
Картинки:    promo.jpg («🔥 Выгода») и rules.jpg («📜 Правила») рядом с bot.py

ГДЕ ЧТО МЕНЯТЬ (ищите название раздела через Ctrl+F)
  Позиции и цены меню ............ файл menu.json (не код)
  Картинки «Выгода»/«Правила» .... файлы promo.jpg / rules.jpg (не код)
  Кешбэк, лимиты заказов ......... НАСТРОЙКИ
  Акции и счастливые часы,
  брони (места, часы, цены аренды),
  предзаказ, сегменты гостей ..... ПРАВИЛА ЗАВЕДЕНИЯ
  Тексты кнопок и сообщений ...... ТЕКСТЫ (приветствие, кнопки, статусы)
                                   + тексты бронирования/предзаказа — в своих разделах
  Кнопки и раскладка ............. КЛАВИАТУРЫ
  Меню команд Telegram, /help .... _register_commands, cmd_help

ОГЛАВЛЕНИЕ (разделы — строки вида «# ===== НАЗВАНИЕ =====»)
  НАСТРОЙКИ                — токен, канал, админы, кешбэк, лимиты
  ПРАВИЛА ЗАВЕДЕНИЯ        — акции, брони, предзаказ, склад (STOCK_RENAMES), графики
  БД                       — таблицы SQLite и все запросы (db_*)
  МЕНЮ                     — загрузка menu.json, склад из меню
  АКЦИИ («Выгода»)         — расчёт счастливых часов, состав акций
  ТЕКСТЫ                   — тексты гостю, статусы заказов, кнопки
  ЛОЯЛЬНОСТЬ               — баллы: начисление, списание, «Моя карта»
  АНАЛИТИКА ПО КАРТАМ ЛОЯЛЬНОСТИ — /loyalty: сегменты, деньги, визиты
  ХРАНИЛИЩЕ ДИАЛОГОВ (FSM) — корзина/стол переживают перезапуск
  КЛАВИАТУРЫ               — все inline/reply-клавиатуры
  ХЕНДЕЛЕРЫ                — дальше обработчики (ПОРЯДОК ВАЖЕН — не переставлять):
    ГРАФИКИ (PNG)          — картинки к /stats и /loyalty
    админ-команды          — /stats /orders /timing /backup
    /stock и /recipe       — склад, таблица, Excel, 📜 история, рецепты; /bookings
    кнопки статусов в канале — «Готов», оплата, отмена
    /start, ввод стола     — вход гостя
  СЦЕНАРИЙ: ВЫБОР → ЛОЯЛЬНОСТЬ — карта гостя, меню, корзина, отправка заказа
  БРОНИРОВАНИЕ             — 📅 бронь: тип → дата → время → гости
  КАРТА ЛОЯЛЬНОСТИ         — /card /guest /loyalty
  ЗАКАЗ ПО АКЦИИ («Выгода») — отдельная корзина акций
  ПРЕДЗАКАЗ                — 📝 заказ к приходу / к брони; /help
  ЗАПУСК                   — main(), проверка файлов при старте
"""

from __future__ import annotations

import asyncio
import calendar
import html
import json
import logging
import os
import re
import sqlite3
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Mapping

from aiogram import Bot, Dispatcher, F, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.exceptions import TelegramBadRequest
from aiogram.filters import Command, CommandObject, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.base import BaseStorage, StateType, StorageKey
from aiogram.types import (
    BotCommand,
    BotCommandScopeAllGroupChats,
    BotCommandScopeAllPrivateChats,
    BotCommandScopeChat,
    BotCommandScopeDefault,
    BufferedInputFile,
    CallbackQuery,
    FSInputFile,
    InputMediaPhoto,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    Message,
    ReplyKeyboardMarkup,
    ReplyKeyboardRemove,
)

try:
    from dotenv import load_dotenv

    load_dotenv()
except ImportError:
    pass

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
)
logger = logging.getLogger("zigzag")

BOT_VERSION = "2026-10-13-v12"

# Все временные метки бота — московские (UTC+3, переход на летнее время отменён)
_MSK = timezone(timedelta(hours=3), name="MSK")


def _now_dt() -> datetime:
    """Текущие дата и время по Москве (наивные — как и все метки в БД)."""
    return datetime.now(_MSK).replace(tzinfo=None)

# ============================== НАСТРОЙКИ ==================================

BOT_TOKEN = os.getenv("BOT_TOKEN", "").strip()
CURRENCY = (os.getenv("CURRENCY", "₽") or "₽").strip()


def _parse_channel_id(raw: str):
    """CHANNEL_ID может быть числом (-100...) или @username."""
    if raw.lstrip("-").isdigit():
        return int(raw)
    return raw


CHANNEL_ID = _parse_channel_id(os.getenv("CHANNEL_ID", "").strip())
ADMIN_IDS: set[int] = {
    int(x)
    for x in os.getenv("ADMIN_IDS", "").replace(";", ",").replace(" ", ",").split(",")
    if x.strip().lstrip("-").isdigit()
}


def _data_dir() -> str:
    """Каталог данных: DATA_DIR (Bothost даёт /app/data), иначе /app/data в Docker, иначе cwd."""
    d = os.getenv("DATA_DIR") or ("/app/data" if os.path.isdir("/app") else ".")
    try:
        os.makedirs(d, exist_ok=True)
    except OSError:
        d = "."  # диск только на чтение — пишем рядом с кодом
    return d


DB_PATH = os.path.join(_data_dir(), "orders.db")


def _env_int(name: str, default: int, lo: int, hi: int) -> int:
    """Целое из переменной окружения с проверкой диапазона (мусор → default)."""
    try:
        v = int(str(os.getenv(name, default)).strip())
    except (TypeError, ValueError):
        return default
    return v if lo <= v <= hi else default


# --- Лояльность (кешбэк баллами, 1 балл = 1 ₽) ---
LOYALTY_PERCENT = _env_int("LOYALTY_PERCENT", 5, 0, 50)        # % кешбэка от оплаты
BONUS_MAX_PAY_PCT = _env_int("BONUS_MAX_PAY_PCT", 50, 0, 100)  # макс. % чека баллами

# --- Защита от спама и злоупотреблений (админы из ADMIN_IDS не ограничены) ---
ORDER_MAX_ITEMS = 20         # максимум позиций (штук) в одном заказе
ORDER_COOLDOWN_SEC = 60      # не чаще одного заказа в минуту от гостя
RSV_MAX_ACTIVE_PER_USER = 2  # активных броней («Ожидает»/«Подтверждена») на гостя
FSM_TTL_HOURS = 12           # через сколько часов бездействия сбрасывать стол/корзину
RESERVE_HOURS = 24           # резерв склада держат только открытые заказы за 24 ч


# ======================= ПРАВИЛА ЗАВЕДЕНИЯ (меняйте здесь) ====================
# Всё, что владелец может захотеть поменять, — в одном месте. После правки —
# перезапуск бота. Тексты кнопок и сообщений — в разделе «ТЕКСТЫ».

# --- Аналитика гостей: пороги сегментов ---
SEG_SLEEP_DAYS = 30        # «😴 давно не был» — дней без визита
SEG_VIP_SPENT = 50000      # «💎 ценный гость» — потратил от, ₽
SEG_VIP_VISITS = 10        # …или визитов от
SEG_OFTEN_30 = 3           # «🔥 ходит часто» — визитов за последние 30 дней от

# --- Акции («🎁 Заказ по акции»): счастливые часы и цены ---
HH_WEEKDAYS = (0, 1, 2, 3, 4)   # пн–пт
HH_START_H, HH_END_H = 15, 18   # счастливые часы 15:00–18:00 (МСК)
HH_GRACE_MIN = 10               # оформить набранное можно до 18:10
# «Любой кальян по полной цене + чай в подарок» на афише входит в счастливые
# часы. True — разрешить подарочный чай в любое время.
PROMO_TEA_ALWAYS = False
PROMO_TEA_HOOKAH_CATS = ("hookah", "hookah_huka")         # какие кальяны «любые»
PROMO_TEA_CATS = ("tea_green", "tea_black", "tea_karkade")  # какой чай в подарок
PROMO_ITEMS: dict[str, dict] = {
    "pr_hh1": {"name": "☀️ Классический кальян (счастливые часы)",
               "price": 1300, "comps": [("h1", 1)], "hh": True},
    "pr_hh2": {"name": "☀️ Премиальный кальян (счастливые часы)",
               "price": 1600, "comps": [("h2", 1)], "hh": True},
    "pr_cb1": {"name": "🔥 Комбо: 2 классических кальяна на ХУКА Про",
               "price": 2500, "comps": [("hu1", 2)], "hh": False},
    "pr_cb2": {"name": "🔥 Комбо: классический + премиальный на ХУКА Про",
               "price": 2800, "comps": [("hu1", 1), ("hu2", 1)], "hh": False},
    "pr_cb3": {"name": "🔥 Комбо: 2 премиальных кальяна на ХУКА Про",
               "price": 3000, "comps": [("hu2", 2)], "hh": False},
}

# --- Бронирование: вместимость, часы работы, длительности, цены аренды ---
RSV_KINDS = {
    "table": "🪑 Стол",
    "bar": "🍸 Место за барной стойкой",
    "vip": "🔑 VIP-комната",
    "ps": "🎮 Аренда PlayStation",
}
# Сколько броней категории может идти ОДНОВРЕМЕННО («Ожидает» + «Подтверждена»).
# Занятость считается по интервалам времени: VIP 15:00–17:00 и 18:00–22:00 в один
# день — можно. После закрытия брони (оплата/отмена/не пришёл) место свободно.
RSV_CAPACITY = {"table": 8, "bar": 5, "vip": 1, "ps": 2}
RSV_CLOSE_H = 3                      # закрытие в 03:00 (работаем 15:00–03:00)
RSV_BUFFER_MIN = {"vip": 30}         # уборка VIP-комнаты между бронями, минут
# сколько стол считается занятым: «1–2 часа» → 2 ч, «2–3 часа» → 3 ч, «3+» → до закрытия
RSV_DUR_OCCUPY_MIN = {"1-2": 120, "2-3": 180, "3+": None}
# сколько минут должно оставаться до закрытия, чтобы вариант стола показывался
RSV_DUR_NEED_MIN = {"1-2": 60, "2-3": 120, "3+": 180}
RSV_BAR_OCCUPY_MIN = 180             # место за стойкой (без часов) — расчётно 3 ч
RSV_BAR_NEED_MIN = 60                # за стойку — не позже чем за час до закрытия
# VIP/PlayStation: пакет (2/4/6 ч) должен целиком закончиться до закрытия
RSV_DURATIONS = {
    # 🪑 Стол
    "1-2": "1–2 часа",
    "2-3": "2–3 часа",
    "3+": "3 и более",
    # 🔑 VIP-комната
    "vip-2": "2 часа - 1 500 ₽",
    "vip-4": "4 часа - 2 500 ₽",
    "vip-6": "6 часов - 3 000 ₽",
    # 🎮 PlayStation
    "ps-2": "2 часа - 400 ₽",
    "ps-4": "4 часа - 600 ₽",
    "ps-6": "6 часов - 700 ₽",
}
# какие длительности показывать каждому виду брони
RSV_DUR_BY_KIND = {
    "table": ["1-2", "2-3", "3+"],
    "bar": ["1-2", "2-3", "3+"],
    "vip": ["vip-2", "vip-4", "vip-6"],
    "ps": ["ps-2", "ps-4", "ps-6"],
}
# ⚠️ Цена аренды записана в ДВУХ местах: в подписи кнопки выше (RSV_DURATIONS)
# и в RSV_BASE_PRICE ниже (по ней считается счёт). Меняете цену — меняйте обе;
# при расхождении бот предупредит админов при запуске.
# подписи под кнопками длительности
RSV_DUR_CAPTIONS = {
    "vip": "При заказе от 6000 ₽ аренда - бессрочно и бесплатно!",
    "ps": "При заказе от 3500 ₽ аренда - бессрочно и бесплатно!",
}
# --- почасовая аренда VIP/PlayStation: правка ±15 мин администратором ---
RSV_BOOKED_MIN = {
    "vip-2": 120, "vip-4": 240, "vip-6": 360,
    "ps-2": 120, "ps-4": 240, "ps-6": 360,
}
RSV_BASE_PRICE = {
    "vip-2": 1500, "vip-4": 2500, "vip-6": 3000,
    "ps-2": 400, "ps-4": 600, "ps-6": 700,
}
RSV_STEP_MIN = 15
RSV_STEP_PRICE = {"vip": 188, "ps": 50}   # ₽ за 15 мин (от тарифа пакета 2 ч)
RSV_ADJ_KINDS = ("vip", "ps")
RSV_MIN_MIN = 15
RSV_DAYS = 30          # на сколько дней вперёд можно бронировать (1 месяц)
RSV_LEAD_MIN = 10      # минимальный запас до времени брони, минут
WDAYS = ["пн", "вт", "ср", "чт", "пт", "сб", "вс"]

# --- Предзаказ ---
PRE_LEAD_MIN = 30          # время прихода — не раньше чем через 30 мин
PRE_SUBMIT_MIN = 15        # при отправке до прихода должно оставаться ≥15 мин
PRE_LAST_BEFORE_CLOSE = 60 # последний слот — за час до закрытия (02:00)
PRE_MAX_ACTIVE = 2         # одновременно активных предзаказов у гостя

# --- Склад: массовые правки текстом / Excel ---
SX_MAX_OPS = 300                 # строк правок за один раз
SX_MAX_FILE = 2 * 1024 * 1024    # размер Excel-файла, байт
SX_MAX_QTY = 1_000_000
SX_MAX_SUM = 100_000_000
SX_MAX_COST = 10_000_000
SH_PAGE = 20                     # строк на странице «📜 История» склада
# Переименовали позицию в menu.json — склад сам её не узнает (он хранит своё
# название). Впишите «старое слово → новое»: при запуске бот заменит его в
# названиях позиций склада (целым словом; если такое название уже есть — пропустит
# и напишет админам). После замены старого слова не остаётся — повторно не сработает.
STOCK_RENAMES = {
    "Стелла Артура": "Стелла Артуа",
    "Буд": "Бад",
}

# --- Картинки в /stats и /loyalty (переменная CHARTS=0 — выключить) ---
CHARTS_ENABLED = os.getenv("CHARTS", "1").strip() not in ("0", "false", "no", "off")


def _is_admin(user_id: int | None) -> bool:
    return bool(user_id) and int(user_id) in ADMIN_IDS

# ================================== БД ======================================


def _connect() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def _now() -> str:
    return _now_dt().strftime("%Y-%m-%d %H:%M:%S")


def db_init() -> None:
    with _connect() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS orders (
                id        INTEGER PRIMARY KEY AUTOINCREMENT,
                table_no  TEXT    NOT NULL,
                total     INTEGER NOT NULL,
                status    TEXT    NOT NULL DEFAULT 'accepted',
                created_at TEXT   NOT NULL,
                updated_at TEXT   NOT NULL,
                message_id INTEGER,
                user_id   INTEGER,
                username  TEXT,
                phone     TEXT,
                issued_at TEXT,
                pay_method TEXT,
                stock_deducted INTEGER NOT NULL DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS order_items (
                id       INTEGER PRIMARY KEY AUTOINCREMENT,
                order_id INTEGER NOT NULL REFERENCES orders(id) ON DELETE CASCADE,
                item_id  TEXT,
                name     TEXT    NOT NULL,
                price    INTEGER NOT NULL,
                qty      INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS ingredients (
                id      INTEGER PRIMARY KEY AUTOINCREMENT,
                name    TEXT NOT NULL UNIQUE,
                unit    TEXT NOT NULL DEFAULT 'шт',
                qty     REAL NOT NULL DEFAULT 0,
                cost    REAL NOT NULL DEFAULT 0,
                min_qty REAL NOT NULL DEFAULT 0,
                created_at TEXT,
                updated_at TEXT
            );
            CREATE TABLE IF NOT EXISTS recipes (
                id        INTEGER PRIMARY KEY AUTOINCREMENT,
                item_id   TEXT NOT NULL,
                ingredient_id INTEGER NOT NULL REFERENCES ingredients(id) ON DELETE CASCADE,
                qty       REAL NOT NULL,
                UNIQUE (item_id, ingredient_id)
            );
            CREATE TABLE IF NOT EXISTS stock_moves (
                id            INTEGER PRIMARY KEY AUTOINCREMENT,
                ingredient_id INTEGER NOT NULL,
                delta         REAL NOT NULL,
                reason        TEXT NOT NULL,
                ref           TEXT,
                created_at    TEXT
            );
            CREATE TABLE IF NOT EXISTS reservations (
                id         INTEGER PRIMARY KEY AUTOINCREMENT,
                kind       TEXT    NOT NULL,
                date       TEXT    NOT NULL,
                time       TEXT    NOT NULL,
                guests     INTEGER NOT NULL,
                duration   TEXT,
                phone      TEXT,
                guest_name TEXT,
                user_id    INTEGER,
                username   TEXT,
                status     TEXT    NOT NULL DEFAULT 'new',
                created_at TEXT    NOT NULL,
                updated_at TEXT    NOT NULL,
                message_id INTEGER,
                actual_min INTEGER,
                price      INTEGER,
                pay_method TEXT
            );
            """
        )
        try:
            conn.execute("ALTER TABLE reservations ADD COLUMN duration TEXT")
        except sqlite3.OperationalError:
            pass  # колонка уже есть (обновление существующей БД)
        for _stmt in (
            "ALTER TABLE orders ADD COLUMN phone TEXT",
            "ALTER TABLE reservations ADD COLUMN phone TEXT",
            "ALTER TABLE orders ADD COLUMN issued_at TEXT",
            "ALTER TABLE orders ADD COLUMN pay_method TEXT",
            "ALTER TABLE order_items ADD COLUMN item_id TEXT",
            "ALTER TABLE orders ADD COLUMN stock_deducted INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE reservations ADD COLUMN actual_min INTEGER",
            "ALTER TABLE reservations ADD COLUMN price INTEGER",
            "ALTER TABLE reservations ADD COLUMN pay_method TEXT",
            # заказ по акции / оплата баллами / кешбэк
            "ALTER TABLE orders ADD COLUMN kind TEXT NOT NULL DEFAULT 'regular'",
            "ALTER TABLE orders ADD COLUMN bonus_pending INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE orders ADD COLUMN bonus_used INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE orders ADD COLUMN bonus_accrued INTEGER NOT NULL DEFAULT 0",
            # предзаказ: время прихода и (необязательно) бронь, к которой он
            "ALTER TABLE orders ADD COLUMN arrive_at TEXT",
            "ALTER TABLE orders ADD COLUMN rsv_id INTEGER",
            # когда закрыт чек аренды (визит и выручка — по этому дню)
            "ALTER TABLE reservations ADD COLUMN paid_at TEXT",
        ):
            try:
                conn.execute(_stmt)
            except sqlite3.OperationalError:
                pass  # колонка уже есть
        # брони, оплаченные до появления paid_at: время оплаты = последнее изменение
        conn.execute(
            "UPDATE reservations SET paid_at = COALESCE(updated_at, created_at)"
            " WHERE status = 'paid' AND paid_at IS NULL"
        )
        conn.executescript(
            """
            -- карта лояльности гостя (ключ — Telegram id; телефон = номер карты)
            CREATE TABLE IF NOT EXISTS guests (
                user_id    INTEGER PRIMARY KEY,
                phone      TEXT,
                phone_verified INTEGER NOT NULL DEFAULT 0,
                username   TEXT,
                first_name TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE UNIQUE INDEX IF NOT EXISTS ux_guests_phone
                ON guests(phone) WHERE phone IS NOT NULL;
            CREATE INDEX IF NOT EXISTS ix_guests_username ON guests(username);
            -- журнал баллов: баланс = SUM(delta)
            CREATE TABLE IF NOT EXISTS bonus_moves (
                id         INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id    INTEGER NOT NULL,
                delta      INTEGER NOT NULL,
                kind       TEXT NOT NULL,
                amount     INTEGER,
                order_id   INTEGER,
                admin_id   INTEGER,
                note       TEXT,
                created_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS ix_bonus_user ON bonus_moves(user_id);
            -- состояние диалогов (FSM) — переживает перезапуск бота
            CREATE TABLE IF NOT EXISTS fsm (
                k          TEXT PRIMARY KEY,
                state      TEXT,
                data       TEXT NOT NULL DEFAULT '{}',
                updated_at REAL NOT NULL
            );
            CREATE INDEX IF NOT EXISTS ix_orders_user ON orders(user_id, created_at);
            CREATE INDEX IF NOT EXISTS ix_rsv_user ON reservations(user_id, status);
            CREATE INDEX IF NOT EXISTS ix_stock_moves_ing ON stock_moves(ingredient_id, id);
            -- позиции меню, для которых склад уже создавался (заглушка — один раз)
            CREATE TABLE IF NOT EXISTS stock_seeded (
                item_id    TEXT PRIMARY KEY,
                created_at TEXT
            );
            """
        )


def db_add_order(
    table_no: str,
    total: int,
    user_id: int | None,
    username: str | None,
    phone: str | None = None,
    kind: str = "regular",
    arrive_at: str | None = None,
    rsv_id: int | None = None,
) -> int:
    now = _now()
    with _connect() as conn:
        cur = conn.execute(
            "INSERT INTO orders (table_no, total, status, created_at, updated_at,"
            " user_id, username, phone, kind, arrive_at, rsv_id)"
            " VALUES (?, ?, 'accepted', ?, ?, ?, ?, ?, ?, ?, ?)",
            (str(table_no), int(total), now, now, user_id, username, phone, kind,
             arrive_at, rsv_id),
        )
        return int(cur.lastrowid)


def db_add_items(order_id: int, items: list[dict]) -> None:
    with _connect() as conn:
        conn.executemany(
            "INSERT INTO order_items (order_id, item_id, name, price, qty)"
            " VALUES (?, ?, ?, ?, ?)",
            [
                (order_id, i.get("id"), i["name"], int(i["price"]), int(i["qty"]))
                for i in items
            ],
        )


def db_delete_order(order_id: int) -> None:
    with _connect() as conn:
        conn.execute("DELETE FROM order_items WHERE order_id = ?", (order_id,))
        conn.execute("DELETE FROM orders WHERE id = ?", (order_id,))


def db_set_message_id(order_id: int, message_id: int) -> None:
    with _connect() as conn:
        conn.execute("UPDATE orders SET message_id = ? WHERE id = ?", (message_id, order_id))


def db_set_status(order_id: int, status: str) -> None:
    now = _now()
    with _connect() as conn:
        if status == "issued":
            # запоминаем время выдачи — для статистики /timing
            conn.execute(
                "UPDATE orders SET status = ?, updated_at = ?, issued_at = ? WHERE id = ?",
                (status, now, now, order_id),
            )
        elif status in ("accepted", "ready", "cancelled"):
            # заказ вернули в работу/отменили — сбрасываем время выдачи
            conn.execute(
                "UPDATE orders SET status = ?, updated_at = ?, issued_at = NULL WHERE id = ?",
                (status, now, order_id),
            )
        else:
            conn.execute(
                "UPDATE orders SET status = ?, updated_at = ? WHERE id = ?",
                (status, now, order_id),
            )


def db_set_pay_method(order_id: int, method: str) -> None:
    with _connect() as conn:
        conn.execute("UPDATE orders SET pay_method = ? WHERE id = ?", (method, order_id))


def db_claim_stock_flag(order_id: int, new_val: int) -> bool:
    """Атомарно переключает флаг списания склада. True — переключили мы
    (значит, списывать/возвращать нужно именно сейчас); False — уже было.
    Защищает от двойного списания при одновременных нажатиях."""
    old_val = 0 if new_val else 1
    with _connect() as conn:
        cur = conn.execute(
            "UPDATE orders SET stock_deducted = ? WHERE id = ? AND stock_deducted = ?",
            (int(new_val), order_id, old_val),
        )
        return cur.rowcount == 1


def db_last_order_age_sec(user_id: int) -> float | None:
    """Сколько секунд назад гость оформил последний заказ (None — не было)."""
    with _connect() as conn:
        row = conn.execute(
            "SELECT created_at FROM orders WHERE user_id = ? ORDER BY id DESC LIMIT 1",
            (user_id,),
        ).fetchone()
    if not row:
        return None
    try:
        t = datetime.strptime(row["created_at"], "%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError):
        return None
    return (_now_dt() - t).total_seconds()


def db_timing_today() -> tuple[list[dict], list[dict]]:
    """Статистика для /timing: выданные сегодня и в работе прямо сейчас."""
    today = _now_dt().strftime("%Y-%m-%d")
    with _connect() as conn:
        issued = conn.execute(
            "SELECT id, table_no, created_at, issued_at FROM orders"
            " WHERE created_at LIKE ? AND issued_at IS NOT NULL"
            " AND kind != 'pre'"  # предзаказ готовят к приходу — скорость не про него
            " ORDER BY issued_at",
            (today + "%",),
        ).fetchall()
        active = conn.execute(
            "SELECT id, table_no, created_at, status FROM orders"
            " WHERE created_at LIKE ? AND status IN ('accepted', 'ready')"
            " AND kind != 'pre'"
            " ORDER BY created_at",
            (today + "%",),
        ).fetchall()
    return [dict(r) for r in issued], [dict(r) for r in active]


def db_get_order(order_id: int) -> dict | None:
    with _connect() as conn:
        row = conn.execute("SELECT * FROM orders WHERE id = ?", (order_id,)).fetchone()
        if row is None:
            return None
        items = conn.execute(
            "SELECT item_id, name, price, qty FROM order_items WHERE order_id = ?",
            (order_id,),
        ).fetchall()
    order = dict(row)
    order["items"] = [dict(i) for i in items]
    return order


def db_recent_orders(limit: int = 10) -> list[dict]:
    with _connect() as conn:
        rows = conn.execute(
            "SELECT id, table_no, total, status, created_at FROM orders"
            " ORDER BY id DESC LIMIT ?",
            (limit,),
        ).fetchall()
    return [dict(r) for r in rows]


# ------------------------------ брони (reservation db) ------------------------
def db_add_reservation(
    kind: str,
    date_s: str,
    time_s: str,
    guests: int,
    duration: str,
    guest_name: str,
    user_id: int | None,
    username: str | None,
    phone: str | None = None,
) -> int:
    now = _now()
    with _connect() as conn:
        cur = conn.execute(
            "INSERT INTO reservations (kind, date, time, guests, duration,"
            " guest_name, user_id, username, phone, status, created_at, updated_at)"
            " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'new', ?, ?)",
            (
                kind,
                date_s,
                time_s,
                int(guests),
                duration,
                guest_name,
                user_id,
                username,
                phone,
                now,
                now,
            ),
        )
        return int(cur.lastrowid)


def db_get_reservation(rid: int) -> dict | None:
    with _connect() as conn:
        row = conn.execute("SELECT * FROM reservations WHERE id = ?", (rid,)).fetchone()
    return dict(row) if row else None


def db_delete_reservation(rid: int) -> None:
    with _connect() as conn:
        conn.execute("DELETE FROM reservations WHERE id = ?", (rid,))


def db_set_rsv_status(rid: int, status: str) -> None:
    with _connect() as conn:
        conn.execute(
            "UPDATE reservations SET status = ?, updated_at = ? WHERE id = ?",
            (status, _now(), rid),
        )


def db_set_rsv_message_id(rid: int, message_id: int) -> None:
    with _connect() as conn:
        conn.execute(
            "UPDATE reservations SET message_id = ? WHERE id = ?",
            (message_id, rid),
        )


def db_set_rsv_actual(rid: int, actual_min: int, price: int) -> None:
    with _connect() as conn:
        conn.execute(
            "UPDATE reservations SET actual_min = ?, price = ? WHERE id = ?",
            (int(actual_min), int(price), rid),
        )


def db_set_rsv_paid(rid: int, method: str, actual_min: int, price: int) -> None:
    with _connect() as conn:
        conn.execute(
            "UPDATE reservations SET status = 'paid', pay_method = ?,"
            " actual_min = ?, price = ?, updated_at = ?, paid_at = ? WHERE id = ?",
            (method, int(actual_min), int(price), _now(), _now(), rid),
        )


def db_rsv_user_active(user_id: int) -> int:
    """Активные брони гостя («Ожидает»/«Подтверждена») на сегодня и позже
    (вчерашняя дата тоже — её ночные слоты ещё могут быть впереди)."""
    since = (_now_dt() - timedelta(days=1)).strftime("%Y-%m-%d")
    with _connect() as conn:
        row = conn.execute(
            "SELECT COUNT(*) AS n FROM reservations WHERE user_id = ?"
            " AND status IN ('new', 'confirmed') AND date >= ?",
            (user_id, since),
        ).fetchone()
    return int(row["n"])


def db_add_reservation_checked(user_limit: int | None, **kw) -> tuple[int | None, str]:
    """Проверка лимитов и вставка брони одной синхронной операцией (без await
    между проверкой и записью — гонки двух гостей за последнее место нет).
    Возвращает (id, "") или (None, "full" | "user_limit")."""
    kind, date_s, uid = kw["kind"], kw["date_s"], kw.get("user_id")
    time_s, code = kw["time_s"], kw.get("duration") or ""
    try:
        late = _arrival(date_s, time_s) <= _now_dt()
    except ValueError:
        return None, "full"
    if (
        late
        or not rsv_fits_closing(kind, date_s, time_s, code)
        or not _rsv_free_in(db_rsv_active_rows(kind, date_s), kind, date_s, time_s, code)
    ):
        return None, "full"
    if user_limit is not None and uid and db_rsv_user_active(uid) >= user_limit:
        return None, "user_limit"
    return db_add_reservation(**kw), ""


# ---------------------------------------------- брони по интервалам времени
def _rsv_close_dt(date_s: str) -> datetime:
    """Закрытие рабочего дня date_s: следующий календарный день, 03:00."""
    d = datetime.strptime(date_s, "%Y-%m-%d") + timedelta(days=1)
    return d.replace(hour=RSV_CLOSE_H, minute=0, second=0, microsecond=0)


def _rsv_need_min(kind: str, code: str) -> int:
    """Сколько минут должно оставаться до закрытия, чтобы вариант был доступен."""
    if kind in RSV_ADJ_KINDS:
        return RSV_BOOKED_MIN.get(code, 120)
    if kind == "bar":
        return RSV_BAR_NEED_MIN
    return RSV_DUR_NEED_MIN.get(code, 60)


def rsv_interval(kind: str, date_s: str, time_s: str, code: str,
                 actual_min: int | None = None) -> tuple[datetime, datetime]:
    """Сколько бронь занимает места: (начало, конец). VIP/PS — пакет или факт
    после правки ±15 мин; стол/бар — расчётно, но не дольше закрытия."""
    start = _arrival(date_s, time_s)
    if kind in RSV_ADJ_KINDS:
        occ = int(actual_min or 0) or RSV_BOOKED_MIN.get(code or "", 120)
        return start, start + timedelta(minutes=occ)
    if kind == "bar":
        occ = RSV_BAR_OCCUPY_MIN
    else:
        occ = RSV_DUR_OCCUPY_MIN.get(code or "", RSV_DUR_OCCUPY_MIN["2-3"])
    close = _rsv_close_dt(date_s)
    end = close if occ is None else min(close, start + timedelta(minutes=occ))
    return start, end


def rsv_fits_closing(kind: str, date_s: str, time_s: str, code: str) -> bool:
    """Помещается ли вариант до закрытия в 03:00."""
    try:
        start = _arrival(date_s, time_s)
    except ValueError:
        return False
    return start + timedelta(minutes=_rsv_need_min(kind, code)) <= _rsv_close_dt(date_s)


def db_rsv_active_rows(kind: str, date_s: str, exclude_id: int | None = None) -> list[dict]:
    """Активные брони категории на рабочий день («Ожидает»/«Подтверждена»)."""
    with _connect() as conn:
        rows = conn.execute(
            "SELECT id, kind, date, time, duration, actual_min FROM reservations"
            " WHERE kind = ? AND date = ? AND status IN ('new', 'confirmed') AND id != ?",
            (kind, date_s, int(exclude_id or 0)),
        ).fetchall()
    return [dict(r) for r in rows]


def _rsv_free_in(rows: list[dict], kind: str, date_s: str, time_s: str, code: str,
                 actual_min: int | None = None) -> bool:
    """True — в интервале новой брони одновременно занято меньше RSV_CAPACITY
    мест. Для VIP к концу каждой брони добавляется время на уборку."""
    cap = RSV_CAPACITY.get(kind)
    if cap is None:
        return True
    buf = timedelta(minutes=RSV_BUFFER_MIN.get(kind, 0))
    s0, e0 = rsv_interval(kind, date_s, time_s, code, actual_min)
    e0 += buf
    ivs: list[tuple[datetime, datetime]] = []
    for r in rows:
        try:
            rs, r_end = rsv_interval(r["kind"], r["date"], r["time"],
                                     r.get("duration") or "", r.get("actual_min"))
        except (ValueError, TypeError, KeyError):
            continue
        r_end += buf
        if rs < e0 and s0 < r_end:
            ivs.append((rs, r_end))
    if len(ivs) < cap:
        return True
    # максимум одновременных броней достигается в начале новой или чьей-то брони
    for p in [s0] + [rs for rs, _ in ivs if rs > s0]:
        if sum(1 for rs, r_end in ivs if rs <= p < r_end) >= cap:
            return False
    return True


def rsv_options(kind: str, date_s: str, time_s: str, rows: list[dict] | None = None) -> list[str]:
    """Доступные варианты длительности на это время (у бара — [""])."""
    if rows is None:
        rows = db_rsv_active_rows(kind, date_s)
    codes = [""] if kind == "bar" else RSV_DUR_BY_KIND.get(kind, [])
    return [
        c for c in codes
        if rsv_fits_closing(kind, date_s, time_s, c)
        and _rsv_free_in(rows, kind, date_s, time_s, c)
    ]


def rsv_slots_available(kind: str, date_s: str) -> list[str]:
    """Время, на которое ещё можно забронировать: не прошло, помещается до
    закрытия и есть свободное место хотя бы для одного варианта длительности."""
    try:
        datetime.strptime(date_s, "%Y-%m-%d")
    except ValueError:
        return []
    now = _now_dt() + timedelta(minutes=RSV_LEAD_MIN)
    rows = db_rsv_active_rows(kind, date_s)
    return [
        sl for sl in _slots()
        if _arrival(date_s, sl) > now and rsv_options(kind, date_s, sl, rows)
    ]


def rsv_capacity_full(kind: str, date_s: str) -> bool:
    """True — на эту дату для категории не осталось ни одного доступного времени."""
    return not rsv_slots_available(kind, date_s)


def db_upcoming_reservations(limit: int = 10) -> list[dict]:
    today = _biz_today().strftime("%Y-%m-%d")  # до 03:00 идёт вчерашняя смена
    with _connect() as conn:
        rows = conn.execute(
            "SELECT * FROM reservations WHERE date >= ? ORDER BY date, time, id LIMIT ?",
            (today, limit),
        ).fetchall()
    return [dict(r) for r in rows]


# ================================== МЕНЮ ====================================

MENU_PATH = Path(__file__).with_name("menu.json")

MENU_ERROR: str | None = None   # меню не загрузилось — причина (админам при старте)
MENU_WARNINGS: list[str] = []    # битые позиции, которые пропустили
STOCK_WARNINGS: list[str] = []   # переименования склада, которые не удались


def _clean_menu(m: dict) -> dict:
    """Проверка menu.json после ручных правок: позиция без id/названия/цены или с
    повтором id не роняет бота — пропускается, а админ видит, что исправить."""
    seen: set[str] = set()
    cats = []
    for ci, c in enumerate(m.get("categories") or [], 1):
        if not isinstance(c, dict) or not c.get("id") or not c.get("name"):
            MENU_WARNINGS.append(f"категория №{ci}: нет id или name — пропущена")
            continue
        items = []
        for it in c.get("items") or []:
            label = f"«{c['name']}» → {it.get('name') or it.get('id') or '?'}" \
                if isinstance(it, dict) else f"«{c['name']}» → ?"
            if not isinstance(it, dict) or not it.get("id") or not it.get("name"):
                MENU_WARNINGS.append(f"{label}: нет id или name — пропущена")
                continue
            try:
                price = int(it.get("price"))
                if price < 0:
                    raise ValueError
            except (TypeError, ValueError):
                MENU_WARNINGS.append(f"{label}: цена не число — пропущена")
                continue
            iid = str(it["id"])
            if iid in seen:
                MENU_WARNINGS.append(f"{label}: id «{iid}» уже есть в меню — пропущена")
                continue
            seen.add(iid)
            items.append({**it, "id": iid, "name": str(it["name"]), "price": price})
        cats.append({**c, "id": str(c["id"]), "name": str(c["name"]),
                     "emoji": c.get("emoji") or "•", "items": items})
    return {**m, "categories": cats}


def _load_menu() -> dict:
    """Меню из menu.json рядом с bot.py. Если файла нет или в нём ошибка — меню
    пустое (брони и карта работают), причина — в логах и админам при старте."""
    global MENU_ERROR
    MENU_ERROR = None
    MENU_WARNINGS.clear()
    try:
        with MENU_PATH.open(encoding="utf-8") as f:
            m = json.load(f)
        if isinstance(m, dict) and m.get("categories"):
            m = _clean_menu(m)
            for w in MENU_WARNINGS:
                logger.warning("menu.json: %s", w)
            return m
        MENU_ERROR = "в menu.json нет категорий"
    except FileNotFoundError:
        MENU_ERROR = "файл menu.json не найден рядом с bot.py"
    except json.JSONDecodeError as e:
        MENU_ERROR = f"ошибка в menu.json: строка {e.lineno}, символ {e.colno} ({e.msg})"
    except OSError as e:
        MENU_ERROR = f"menu.json не читается: {e}"
    logger.error("❌ Меню не загружено — %s. Каталог пуст: исправьте файл и перезапустите бота",
                 MENU_ERROR)
    return {"categories": []}


MENU = _load_menu()


def seed_stock_from_menu() -> None:
    """Строки склада для НОВЫХ позиций меню: позиция = ингредиент (единица «шт»),
    рецепт 1 порция = 1 шт. qty=999 — заглушка, замените реальным остатком.
    Каждая позиция меню обрабатывается один раз (таблица stock_seeded): если
    заглушку или рецепт удалили вручную, после перезапуска они НЕ вернутся."""
    STOCK_WARNINGS[:] = stock_apply_renames()   # сначала — новые названия из STOCK_RENAMES
    created_ings = created_recipes = 0
    now = _now()
    with _connect() as conn:
        seeded = {r[0] for r in conn.execute("SELECT item_id FROM stock_seeded")}
        for cat in MENU.get("categories", []):
            for it in cat.get("items", []):
                if it["id"] in seeded:
                    continue  # уже обрабатывали — ручные правки склада не трогаем
                conn.execute(
                    "INSERT OR IGNORE INTO stock_seeded (item_id, created_at) VALUES (?, ?)",
                    (it["id"], now),
                )
                has = conn.execute(
                    "SELECT 1 FROM recipes WHERE item_id = ? LIMIT 1", (it["id"],)
                ).fetchone()
                if has:
                    continue  # рецепт уже настроен — не вмешиваемся
                row = conn.execute(
                    "SELECT id FROM ingredients WHERE name = ?", (it["name"],)
                ).fetchone()
                if row:
                    ing_id = int(row["id"])
                else:
                    cur = conn.execute(
                        "INSERT INTO ingredients"
                        " (name, unit, qty, cost, min_qty, created_at, updated_at)"
                        " VALUES (?, 'шт', 999, 0, 0, ?, ?)",
                        (it["name"], now, now),
                    )
                    ing_id = int(cur.lastrowid)
                    created_ings += 1
                conn.execute(
                    "INSERT OR IGNORE INTO recipes (item_id, ingredient_id, qty)"
                    " VALUES (?, ?, 1)",
                    (it["id"], ing_id),
                )
                created_recipes += 1
    if created_ings or created_recipes:
        logger.info(
            "Склад пополнен из меню: +%d позиций, +%d рецептов",
            created_ings, created_recipes,
        )


def esc(text) -> str:
    """Экранирование для HTML-разметки Telegram."""
    return html.escape(str(text))


def f_category(cat_id: str) -> dict | None:
    return next((c for c in MENU["categories"] if c["id"] == cat_id), None)


def f_find_item(item_id: str) -> dict | None:
    for c in MENU["categories"]:
        for it in c["items"]:
            if it["id"] == item_id:
                return it
    return None


# ============================ АКЦИИ («Выгода») ===============================
# Заказ по акции — ОТДЕЛЬНЫЙ заказ (своя корзина). Баллы на него не начисляются
# и не списываются («акции не суммируются»). Склад списывается по рецептам
# обычных позиций меню, из которых состоит акция (comps).
# Цены акций и часы — в НАСТРОЙКИ → «Акции».


def _is_promo_id(item_id: str | None) -> bool:
    return bool(item_id) and str(item_id).startswith("pr_")


def _items_in_cats(cat_ids) -> list[dict]:
    out: list[dict] = []
    for cid in cat_ids:
        cat = f_category(cid)
        if cat:
            out.extend(cat.get("items", []))
    return out


def promo_item(pid: str) -> dict | None:
    """Описание акционной позиции (id, name, price, comps, hh, old) или None,
    если такой акции нет / нужных позиций нет в меню."""
    if pid in PROMO_ITEMS:
        p = PROMO_ITEMS[pid]
        old = 0
        for base_id, k in p["comps"]:
            base = f_find_item(base_id)
            if not base:
                return None
            old += int(base["price"]) * k
        return {"id": pid, "name": p["name"], "price": int(p["price"]),
                "comps": list(p["comps"]), "hh": bool(p["hh"]), "old": old}
    if pid.startswith("pr_t:"):  # pr_t:<кальян>:<чай>
        parts = pid.split(":")
        if len(parts) != 3:
            return None
        hid, tid = parts[1], parts[2]
        hookah = next((i for i in _items_in_cats(PROMO_TEA_HOOKAH_CATS) if i["id"] == hid), None)
        tea = next((i for i in _items_in_cats(PROMO_TEA_CATS) if i["id"] == tid), None)
        if not hookah or not tea:
            return None
        return {"id": pid, "name": f"🍵 {hookah['name']} + чай «{tea['name']}» в подарок",
                "price": int(hookah["price"]), "comps": [(hid, 1), (tid, 1)],
                "hh": not PROMO_TEA_ALWAYS,
                "old": int(hookah["price"]) + int(tea["price"])}
    return None


def hh_active(grace_min: int = 0) -> bool:
    """Идут ли счастливые часы (будни 15:00–18:00 МСК; grace — минуты после 18:00)."""
    now = _now_dt()
    if now.weekday() not in HH_WEEKDAYS:
        return False
    start = now.replace(hour=HH_START_H, minute=0, second=0, microsecond=0)
    end = now.replace(hour=HH_END_H, minute=0, second=0, microsecond=0)
    return start <= now < end + timedelta(minutes=grace_min)


def _components(item_id: str) -> list[tuple[str, int]]:
    """Из каких позиций меню состоит позиция заказа (для склада)."""
    if _is_promo_id(item_id):
        p = promo_item(item_id)
        return list(p["comps"]) if p else []
    return [(item_id, 1)]


# ================================ ТЕКСТЫ ====================================

STATUSES: dict[str, str] = {
    "accepted": "🟡 Принят",
    "ready": "🟢 Готов",
    "issued": "🔵 Выдан",
    "paid": "💵 Оплачен",
    "cancelled": "❌ Отменён",
}

# ------------------------------------------------------ СКЛАД ----------------
def _fmt_qty(n: float) -> str:
    """85.0 → «85», 10.5 → «10.5»."""
    return f"{float(n):g}"


def _fmt_price(n) -> str:
    """Себестоимость: 650 → «650 ₽», 6.5 → «6,5 ₽»."""
    v = float(n)
    if v == int(v):
        return _fmt_money(int(v))
    return f"{v:g}".replace(".", ",") + f" {CURRENCY}"


def db_list_ingredients() -> list[dict]:
    with _connect() as conn:
        return [dict(r) for r in conn.execute("SELECT * FROM ingredients ORDER BY name")]


def db_get_ingredient(ing_id: int) -> dict | None:
    with _connect() as conn:
        r = conn.execute("SELECT * FROM ingredients WHERE id = ?", (ing_id,)).fetchone()
        return dict(r) if r else None


def db_add_ingredient(name: str, unit: str, qty: float, cost: float, min_qty: float) -> int:
    now = _now()
    with _connect() as conn:
        cur = conn.execute(
            "INSERT INTO ingredients (name, unit, qty, cost, min_qty, created_at,"
            " updated_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (name.strip(), unit, float(qty), float(cost), float(min_qty), now, now),
        )
        return int(cur.lastrowid)


def db_set_stock_unit(ing_id: int, unit: str) -> None:
    """Смена единицы измерения позиции (шт / г / мл)."""
    with _connect() as conn:
        conn.execute(
            "UPDATE ingredients SET unit = ?, updated_at = ? WHERE id = ?",
            (unit, _now(), int(ing_id)),
        )


def stock_move(ing_id: int, delta: float, reason: str, ref: str | None = None) -> dict | None:
    """Меняет остаток и пишет движение. Возвращает новое состояние или None."""
    with _connect() as conn:
        r = conn.execute("SELECT * FROM ingredients WHERE id = ?", (ing_id,)).fetchone()
        if not r:
            return None
        new_qty = float(r["qty"]) + float(delta)
        conn.execute(
            "UPDATE ingredients SET qty = ?, updated_at = ? WHERE id = ?",
            (new_qty, _now(), ing_id),
        )
        conn.execute(
            "INSERT INTO stock_moves (ingredient_id, delta, reason, ref, created_at)"
            " VALUES (?, ?, ?, ?, ?)",
            (ing_id, float(delta), reason, ref, _now()),
        )
        d = dict(r)
        d["qty"] = new_qty
        return d


def db_set_cost(ing_id: int, cost: float) -> None:
    with _connect() as conn:
        conn.execute(
            "UPDATE ingredients SET cost = ?, updated_at = ? WHERE id = ?",
            (float(cost), _now(), ing_id),
        )


def db_set_min_qty(ing_id: int, min_qty: float) -> None:
    with _connect() as conn:
        conn.execute(
            "UPDATE ingredients SET min_qty = ?, updated_at = ? WHERE id = ?",
            (float(min_qty), _now(), ing_id),
        )


def db_delete_ingredient(ing_id: int) -> None:
    """Удаление позиции: из склада, её рецепты и журнал движений."""
    with _connect() as conn:
        conn.execute("DELETE FROM recipes WHERE ingredient_id = ?", (ing_id,))
        conn.execute("DELETE FROM stock_moves WHERE ingredient_id = ?", (ing_id,))
        conn.execute("DELETE FROM ingredients WHERE id = ?", (ing_id,))


def db_recent_moves(ing_id: int, limit: int = 5) -> list[dict]:
    with _connect() as conn:
        return [
            dict(r)
            for r in conn.execute(
                "SELECT * FROM stock_moves WHERE ingredient_id = ?"
                " ORDER BY id DESC LIMIT ?",
                (ing_id, limit),
            )
        ]


# фильтры «📜 История»: ключ → (подпись кнопки, SQL-условие; без ввода пользователя)
SH_FILTERS = {
    "a": ("Все", "1 = 1"),
    "i": ("📥 Приход", "m.delta > 0 AND m.reason != 'возврат'"),
    "o": ("📤 Списания", "m.delta < 0 AND m.reason != 'списание'"),
    "z": ("🧾 Заказы", "m.reason IN ('списание', 'возврат')"),
}
SH_PERIODS = {"d": ("Сегодня", 0), "w": ("7 дн", 6), "m": ("30 дн", 29), "all": ("Всё", None)}


def _sh_since(period: str) -> str | None:
    """Начало периода: рабочий день начинается в 06:00 (ночная смена — вчерашняя)."""
    back = SH_PERIODS[period][1]
    if back is None:
        return None
    day = (_now_dt() - timedelta(hours=6)).date() - timedelta(days=back)
    return f"{day:%Y-%m-%d} 06:00:00"


def db_stock_history(ing_id: int, flt: str, period: str, page: int) -> dict:
    """Движения склада (новые сверху) + итоги за период по видам.
    ing_id=0 — все позиции. flt/period — ключи SH_FILTERS/SH_PERIODS."""
    since = _sh_since(period)
    where = ["(? = 0 OR m.ingredient_id = ?)", "(? IS NULL OR m.created_at >= ?)"]
    args: list = [ing_id, ing_id, since, since]
    base = " WHERE " + " AND ".join(where)
    with _connect() as conn:
        tot = {}
        for key in ("i", "o", "z"):
            r = conn.execute(
                "SELECT COUNT(*) AS n, COALESCE(SUM(m.delta), 0) AS s FROM stock_moves m"
                + base + " AND " + SH_FILTERS[key][1], args).fetchone()
            tot[key] = (int(r["n"]), float(r["s"]))
        cond = base + " AND " + SH_FILTERS[flt][1]
        total = int(conn.execute("SELECT COUNT(*) FROM stock_moves m" + cond, args).fetchone()[0])
        rows = [dict(r) for r in conn.execute(
            "SELECT m.*, i.name, i.unit FROM stock_moves m"
            " JOIN ingredients i ON i.id = m.ingredient_id" + cond
            + " ORDER BY m.id DESC LIMIT ? OFFSET ?", args + [SH_PAGE, page * SH_PAGE])]
    return {"rows": rows, "total": total, "tot": tot}


def stock_apply_renames() -> list[str]:
    """STOCK_RENAMES → названия позиций склада. Возвращает предупреждения
    (новое название уже занято)."""
    warns: list[str] = []
    if not STOCK_RENAMES:
        return warns
    with _connect() as conn:
        for r in conn.execute("SELECT id, name FROM ingredients").fetchall():
            name = new = r["name"]
            for old, repl in STOCK_RENAMES.items():
                new = re.sub(rf"(?<!\w){re.escape(old)}(?!\w)", repl, new)
            if new == name:
                continue
            taken = conn.execute("SELECT 1 FROM ingredients WHERE name = ? AND id != ?",
                                 (new, r["id"])).fetchone()
            if taken:
                warns.append(f"склад: «{name}» не переименована — «{new}» уже есть")
                continue
            conn.execute("UPDATE ingredients SET name = ?, updated_at = ? WHERE id = ?",
                         (new, _now(), r["id"]))
            logger.info("Склад: «%s» → «%s»", name, new)
    for w in warns:
        logger.warning(w)
    return warns


def db_rename_ingredient(ing_id: int, name: str) -> bool:
    """Переименовать позицию склада. False — такое название уже занято."""
    try:
        with _connect() as conn:
            conn.execute("UPDATE ingredients SET name = ?, updated_at = ? WHERE id = ?",
                         (name, _now(), int(ing_id)))
        return True
    except sqlite3.IntegrityError:
        return False


def db_recipes_for_item(item_id: str) -> list[dict]:
    """Ингредиенты позиции меню с количеством на 1 порцию."""
    with _connect() as conn:
        return [
            dict(r)
            for r in conn.execute(
                "SELECT r.id AS rid, r.qty AS need, i.* FROM recipes r"
                " JOIN ingredients i ON i.id = r.ingredient_id"
                " WHERE r.item_id = ? ORDER BY i.name",
                (item_id,),
            )
        ]


def db_upsert_recipe(item_id: str, ing_id: int, qty: float) -> bool:
    """True — добавлено, False — уже было (обновили количество)."""
    with _connect() as conn:
        try:
            conn.execute(
                "INSERT INTO recipes (item_id, ingredient_id, qty) VALUES (?, ?, ?)",
                (item_id, ing_id, float(qty)),
            )
            return True
        except sqlite3.IntegrityError:
            conn.execute(
                "UPDATE recipes SET qty = ? WHERE item_id = ? AND ingredient_id = ?",
                (float(qty), item_id, ing_id),
            )
            return False


def db_delete_recipe(item_id: str, ing_id: int) -> None:
    with _connect() as conn:
        conn.execute(
            "DELETE FROM recipes WHERE item_id = ? AND ingredient_id = ?",
            (item_id, ing_id),
        )


def _item_needs(conn: sqlite3.Connection, item_id: str) -> dict[int, float]:
    """Сколько каждого ингредиента уходит на 1 шт позиции заказа
    (акции раскладываются на обычные позиции меню)."""
    needs: dict[int, float] = {}
    for base_id, k in _components(item_id):
        for r in conn.execute(
            "SELECT r.ingredient_id, r.qty FROM recipes r"
            " JOIN ingredients i ON i.id = r.ingredient_id WHERE r.item_id = ?",
            (base_id,),
        ):
            iid = int(r["ingredient_id"])
            needs[iid] = needs.get(iid, 0.0) + float(r["qty"]) * k
    return needs


def _reserved_map(conn: sqlite3.Connection) -> dict[int, float]:
    """Ингредиенты, «зарезервированные» открытыми (ещё не оплаченными и не
    отменёнными) заказами — склад списывается только при оплате, поэтому без
    резерва несколько открытых заказов могли бы продать больше, чем есть.
    Заказы старше RESERVE_HOURS (забыли закрыть) резерв не держат; предзаказ
    держит резерв, пока время прихода не ушло в прошлое больше чем на RESERVE_HOURS."""
    since = (_now_dt() - timedelta(hours=RESERVE_HOURS)).strftime("%Y-%m-%d %H:%M:%S")
    rows = conn.execute(
        "SELECT oi.item_id, oi.qty FROM order_items oi JOIN orders o ON o.id = oi.order_id"
        " WHERE o.status IN ('accepted', 'ready', 'issued') AND o.stock_deducted = 0"
        " AND oi.item_id IS NOT NULL AND (o.created_at >= ? OR o.arrive_at >= ?)",
        (since, since),
    ).fetchall()
    cache: dict[str, dict[int, float]] = {}
    res: dict[int, float] = {}
    for it in rows:
        iid = it["item_id"]
        if iid not in cache:
            cache[iid] = _item_needs(conn, iid)
        for ing_id, need in cache[iid].items():
            res[ing_id] = res.get(ing_id, 0.0) + need * int(it["qty"])
    return res


def reserved_snapshot() -> dict[int, float]:
    with _connect() as conn:
        return _reserved_map(conn)


def max_available(item_id: str, reserved: dict[int, float] | None = None) -> int | None:
    """Сколько штук позиции можно продать по остаткам за вычетом резерва
    открытых заказов. None — рецепта нет, ограничения нет."""
    with _connect() as conn:
        needs = _item_needs(conn, item_id)
        if not needs:
            return None
        if reserved is None:
            reserved = _reserved_map(conn)
        limit: int | None = None
        for ing_id, need in needs.items():
            if need <= 0:
                continue
            row = conn.execute("SELECT qty FROM ingredients WHERE id = ?", (ing_id,)).fetchone()
            if not row:
                continue
            free = max(0.0, float(row["qty"]) - reserved.get(ing_id, 0.0))
            avail = int((free + 1e-9) // need)
            limit = avail if limit is None else min(limit, avail)
        return limit


def item_in_stock(item_id: str, reserved: dict[int, float] | None = None) -> bool:
    """False — позицию скрыть из меню (не хватает ингредиентов даже на 1 шт).
    Нет рецепта — True (не списываем — значит, не блокируем)."""
    m = max_available(item_id, reserved)
    return m is None or m >= 1


def cart_stock_problems(cart: list[dict]) -> list[str]:
    """Нехватка ингредиентов под ВСЮ корзину (с учётом общих ингредиентов
    и резерва открытых заказов)."""
    problems: list[str] = []
    with _connect() as conn:
        demands: dict[int, float] = {}
        for it in cart:
            iid = it.get("id")
            if not iid:
                continue
            for ing_id, need in _item_needs(conn, iid).items():
                demands[ing_id] = demands.get(ing_id, 0.0) + need * int(it["qty"])
        reserved = _reserved_map(conn)
        for ing_id, need in demands.items():
            row = conn.execute(
                "SELECT * FROM ingredients WHERE id = ?", (ing_id,)
            ).fetchone()
            if not row:
                continue
            free = max(0.0, float(row["qty"]) - reserved.get(ing_id, 0.0))
            if need > free + 1e-9:
                problems.append(
                    f"• {esc(row['name'])} — нужно {_fmt_qty(need)} {esc(row['unit'])}, "
                    f"а доступно {_fmt_qty(free)} {esc(row['unit'])}"
                )
    return problems


def stock_apply_order(order: dict, sign: int) -> list[dict]:
    """Списание (sign=-1) или возврат (sign=+1) ингредиентов по рецепту позиций
    заказа. Возвращает затронутые ингредиенты с новыми остатками."""
    with _connect() as conn:
        items = conn.execute(
            "SELECT item_id, qty FROM order_items"
            " WHERE order_id = ? AND item_id IS NOT NULL",
            (order["id"],),
        ).fetchall()
        affected: dict[int, dict] = {}
        for it in items:
            for ing_id, need in _item_needs(conn, it["item_id"]).items():
                delta = float(sign) * float(need) * float(it["qty"])
                row = conn.execute(
                    "SELECT * FROM ingredients WHERE id = ?", (ing_id,)
                ).fetchone()
                if not row:
                    continue
                new_qty = float(row["qty"]) + delta
                conn.execute(
                    "UPDATE ingredients SET qty = ?, updated_at = ? WHERE id = ?",
                    (new_qty, _now(), row["id"]),
                )
                conn.execute(
                    "INSERT INTO stock_moves (ingredient_id, delta, reason, ref,"
                    " created_at) VALUES (?, ?, ?, ?, ?)",
                    (
                        row["id"],
                        delta,
                        "списание" if sign < 0 else "возврат",
                        f"заказ#{order['id']}",
                        _now(),
                    ),
                )
                d = dict(row)
                d["qty"] = new_qty
                affected[row["id"]] = d
        return list(affected.values())


def db_stock_value() -> int:
    """Вложено денег в остатки."""
    with _connect() as conn:
        r = conn.execute(
            "SELECT COALESCE(SUM(qty * cost), 0) AS v FROM ingredients"
        ).fetchone()
        return int(round(r["v"]))


def db_cogs_range(d_from: str, d_to: str) -> int:
    """Себестоимость проданного за период (по текущим с/с ингредиентов):
    только ОПЛАЧЕННЫЕ заказы, по рецептам их позиций (акции — по составу)."""
    total = 0.0
    with _connect() as conn:
        rows = conn.execute(
            "SELECT oi.item_id, oi.qty FROM order_items oi"
            " JOIN orders o ON o.id = oi.order_id"
            " WHERE substr(o.created_at, 1, 10) BETWEEN ? AND ?"
            " AND o.status = 'paid' AND oi.item_id IS NOT NULL",
            (d_from, d_to),
        ).fetchall()
        costs = {int(r["id"]): float(r["cost"]) for r in conn.execute("SELECT id, cost FROM ingredients")}
        for it in rows:
            for ing_id, need in _item_needs(conn, it["item_id"]).items():
                total += need * float(it["qty"]) * costs.get(ing_id, 0.0)
    return int(round(total))


async def _stock_alerts(bot, affected: list[dict]) -> None:
    """Личное сообщение админам: закончилось или упало ниже минимума."""
    lines: list[str] = []
    for ing in affected:
        q, mn = float(ing["qty"]), float(ing["min_qty"])
        if q <= 0:
            lines.append(
                f"⛔ Закончилось: «{esc(ing['name'])}» (0 {esc(ing['unit'])}) — "
                "позиции меню с ним скрыты до прихода"
            )
        elif mn > 0 and q <= mn:
            lines.append(
                f"⚠️ Мало: «{esc(ing['name'])}» — {_fmt_qty(q)} {esc(ing['unit'])} "
                f"(мин {_fmt_qty(mn)}), пора заказать"
            )
    if not lines:
        return
    # длинный список — несколькими сообщениями (лимит Telegram 4096 символов)
    chunks = ["📦 <b>Склад</b>"]
    for ln in lines:
        if len(chunks[-1]) + len(ln) + 1 > 3800:
            chunks.append("📦 <b>Склад</b> (продолжение)")
        chunks[-1] += "\n" + ln
    for admin_id in ADMIN_IDS:
        for text in chunks:
            try:
                await bot.send_message(admin_id, text)
            except Exception:
                logger.info("Не удалось прислать склад-алерт админу %s", admin_id)
                break


# ============================== ЛОЯЛЬНОСТЬ ==================================
# Карта гостя = его Telegram-аккаунт + номер телефона (номер карты).
# Найти гостя админ может по телефону или @нику. Баллы: 1 балл = 1 ₽,
# кешбэк LOYALTY_PERCENT % от оплаченной деньгами части обычного заказа.
# Баланс всегда = сумма журнала bonus_moves (никаких «кэшированных» остатков).


def normalize_phone(raw) -> str | None:
    """«8 (900) 123-45-67» / «9001234567» / «+7 900…» → «+79001234567»."""
    digits = re.sub(r"\D", "", str(raw or ""))
    if len(digits) == 11 and digits[0] == "8":
        digits = "7" + digits[1:]
    elif len(digits) == 10 and digits[0] == "9":
        digits = "7" + digits  # российский мобильный без кода страны
    if not 10 <= len(digits) <= 15:
        return None
    return "+" + digits


def db_guest_get(user_id: int | None) -> dict | None:
    if not user_id:
        return None
    with _connect() as conn:
        row = conn.execute("SELECT * FROM guests WHERE user_id = ?", (user_id,)).fetchone()
    return dict(row) if row else None


def guest_card(user_id: int | None) -> dict | None:
    """Карта есть, только если гость оставил телефон."""
    g = db_guest_get(user_id)
    return g if g and g.get("phone") else None


def db_guest_admin_set_phone(uid: int, phone: str) -> tuple[bool, dict | None]:
    """Смена номера карты администратором: баллы и история остаются (ключ карты —
    Telegram id). (False, владелец) — номер уже у другой карты."""
    try:
        with _connect() as conn:
            owner = conn.execute("SELECT * FROM guests WHERE phone = ? AND user_id != ?",
                                 (phone, int(uid))).fetchone()
            if owner:
                return False, dict(owner)
            conn.execute("UPDATE guests SET phone = ?, phone_verified = 0, updated_at = ?"
                         " WHERE user_id = ?", (phone, _now(), int(uid)))
    except sqlite3.IntegrityError:   # номер успели занять параллельно
        return False, None
    return True, None


def db_guest_touch(user) -> None:
    """Обновить @ник и имя у существующей карты (для поиска по нику)."""
    if not user:
        return
    with _connect() as conn:
        conn.execute(
            "UPDATE guests SET username = ?, first_name = ?, updated_at = ? WHERE user_id = ?",
            ((user.username or None), (user.first_name or None), _now(), user.id),
        )


def db_guest_set_phone(user, phone: str, verified: bool) -> tuple[bool, str]:
    """Привязать телефон к карте гостя. (True, "") — ок; (False, "taken") —
    номер уже принадлежит другому аккаунту. Подтверждённый (своим контактом)
    номер забирает себе номер, введённый кем-то вручную без подтверждения."""
    now = _now()
    with _connect() as conn:
        owner = conn.execute(
            "SELECT user_id, phone_verified FROM guests WHERE phone = ?", (phone,)
        ).fetchone()
        if owner and int(owner["user_id"]) != int(user.id):
            if verified and not owner["phone_verified"]:
                conn.execute(
                    "UPDATE guests SET phone = NULL, phone_verified = 0, updated_at = ?"
                    " WHERE user_id = ?",
                    (now, owner["user_id"]),
                )
                logger.warning(
                    "Телефон %s перепривязан: был у %s (не подтверждён) → %s (контакт)",
                    phone, owner["user_id"], user.id,
                )
            else:
                return False, "taken"
        me = conn.execute(
            "SELECT phone, phone_verified FROM guests WHERE user_id = ?", (user.id,)
        ).fetchone()
        if me:
            keep_verified = bool(me["phone_verified"]) and me["phone"] == phone
            conn.execute(
                "UPDATE guests SET phone = ?, phone_verified = ?, username = ?,"
                " first_name = ?, updated_at = ? WHERE user_id = ?",
                (phone, int(verified or keep_verified), user.username or None,
                 user.first_name or None, now, user.id),
            )
        else:
            conn.execute(
                "INSERT INTO guests (user_id, phone, phone_verified, username, first_name,"
                " created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (user.id, phone, int(verified), user.username or None,
                 user.first_name or None, now, now),
            )
    return True, ""


def _bonus_balance(conn: sqlite3.Connection, user_id: int) -> int:
    row = conn.execute(
        "SELECT COALESCE(SUM(delta), 0) AS b FROM bonus_moves WHERE user_id = ?", (user_id,)
    ).fetchone()
    return int(row["b"])


def db_bonus_balance(user_id: int | None) -> int:
    if not user_id:
        return 0
    with _connect() as conn:
        return _bonus_balance(conn, user_id)


def _bonus_add(conn, user_id: int, delta: int, kind: str, amount: int | None = None,
               order_id: int | None = None, admin_id: int | None = None,
               note: str | None = None) -> None:
    conn.execute(
        "INSERT INTO bonus_moves (user_id, delta, kind, amount, order_id, admin_id, note,"
        " created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        (user_id, int(delta), kind, amount, order_id, admin_id, note, _now()),
    )


def bonus_max_for(total: int) -> int:
    """Сколько баллов можно списать в оплату чека на сумму total."""
    return max(0, int(total) * BONUS_MAX_PAY_PCT // 100)


def bonus_available_for_order(order: dict) -> int:
    """Сколько баллов можно списать в этот заказ (0 — нельзя: акция/нет карты)."""
    if order.get("kind") == "promo" or not order.get("user_id"):
        return 0
    if not guest_card(order["user_id"]):
        return 0
    return min(db_bonus_balance(order["user_id"]), bonus_max_for(order["total"]))


def db_set_bonus_pending(order_id: int, points: int) -> None:
    with _connect() as conn:
        conn.execute(
            "UPDATE orders SET bonus_pending = ? WHERE id = ? AND status != 'paid'",
            (max(0, int(points)), order_id),
        )


def db_claim_paid(order_id: int, method: str) -> bool:
    """Атомарно перевести заказ в «Оплачен». True — перевели мы (дальше —
    баллы и склад); False — заказ уже был оплачен (двойное нажатие)."""
    with _connect() as conn:
        cur = conn.execute(
            "UPDATE orders SET status = 'paid', pay_method = ?, updated_at = ?"
            " WHERE id = ? AND status != 'paid'",
            (method, _now(), order_id),
        )
        return cur.rowcount == 1


def loyalty_on_paid(order_id: int) -> dict:
    """После оплаты: списать «отложенные» баллы и начислить кешбэк.
    Вызывать только тому, кто выиграл db_claim_paid (ровно один раз)."""
    res = {"spent": 0, "accrued": 0, "balance": None}
    with _connect() as conn:
        o = conn.execute("SELECT * FROM orders WHERE id = ?", (order_id,)).fetchone()
        if not o:
            return res
        uid = o["user_id"]
        card = None
        if uid and o["kind"] != "promo":
            card = conn.execute(
                "SELECT 1 FROM guests WHERE user_id = ? AND phone IS NOT NULL", (uid,)
            ).fetchone()
        if not card:
            conn.execute("UPDATE orders SET bonus_pending = 0 WHERE id = ?", (order_id,))
            return res
        bal = _bonus_balance(conn, uid)
        spend = max(0, min(int(o["bonus_pending"] or 0), bal, bonus_max_for(o["total"])))
        if spend:
            _bonus_add(conn, uid, -spend, "spend", order_id=order_id,
                       note=f"оплата заказа #{order_id}")
        accrue = (int(o["total"]) - spend) * LOYALTY_PERCENT // 100
        if accrue > 0:
            _bonus_add(conn, uid, accrue, "accrual", amount=int(o["total"]) - spend,
                       order_id=order_id, note=f"кешбэк {LOYALTY_PERCENT}% за заказ #{order_id}")
        conn.execute(
            "UPDATE orders SET bonus_used = ?, bonus_accrued = ?, bonus_pending = 0 WHERE id = ?",
            (spend, max(0, accrue), order_id),
        )
        res.update(spent=spend, accrued=max(0, accrue), balance=_bonus_balance(conn, uid))
    return res


def loyalty_on_refund(order_id: int) -> dict:
    """Возврат оплаченного заказа: вернуть списанные баллы и забрать
    начисленный кешбэк (не ниже нуля). Атомарно и ровно один раз."""
    res = {"returned": 0, "taken": 0, "balance": None, "user_id": None}
    with _connect() as conn:
        o = conn.execute("SELECT * FROM orders WHERE id = ?", (order_id,)).fetchone()
        if not o or not o["user_id"]:
            return res
        cur = conn.execute(
            "UPDATE orders SET bonus_used = 0, bonus_accrued = 0, bonus_pending = 0"
            " WHERE id = ? AND (bonus_used > 0 OR bonus_accrued > 0)",
            (order_id,),
        )
        if cur.rowcount != 1:
            return res
        uid = int(o["user_id"])
        used, acc = int(o["bonus_used"] or 0), int(o["bonus_accrued"] or 0)
        if used:
            _bonus_add(conn, uid, used, "refund", order_id=order_id,
                       note=f"возврат баллов, заказ #{order_id}")
        take = 0
        if acc:
            take = min(acc, max(0, _bonus_balance(conn, uid)))
            if take:
                note = f"отмена кешбэка, заказ #{order_id}"
                if take < acc:
                    note += f" (не хватило {acc - take})"
                _bonus_add(conn, uid, -take, "refund", order_id=order_id, note=note)
        res.update(returned=used, taken=take, balance=_bonus_balance(conn, uid), user_id=uid)
    return res


def loyalty_offline_accrue(user_id: int, check_sum: int, admin_id: int) -> int:
    """Админ начисляет кешбэк за чек, оплаченный мимо бота."""
    pts = int(check_sum) * LOYALTY_PERCENT // 100
    with _connect() as conn:
        _bonus_add(conn, user_id, pts, "offline", amount=int(check_sum),
                   admin_id=admin_id, note=f"чек {int(check_sum)} ₽ (вне бота)")
    return pts


def loyalty_manual_spend(user_id: int, points: int, admin_id: int) -> bool:
    with _connect() as conn:
        if points <= 0 or _bonus_balance(conn, user_id) < points:
            return False
        _bonus_add(conn, user_id, -int(points), "manual", admin_id=admin_id,
                   note="списание администратором")
    return True


def guest_stats(user_id: int) -> dict:
    """Сумма оплат, число заказов, визиты (дни), последний визит, история."""
    with _connect() as conn:
        o = conn.execute(
            "SELECT COUNT(*) AS n, COALESCE(SUM(total - bonus_used), 0) AS s"
            " FROM orders WHERE user_id = ? AND status = 'paid'",
            (user_id,),
        ).fetchone()
        off = conn.execute(
            "SELECT COALESCE(SUM(amount), 0) AS s FROM bonus_moves"
            " WHERE user_id = ? AND kind = 'offline'",
            (user_id,),
        ).fetchone()
        rent = conn.execute(
            "SELECT COALESCE(SUM(price), 0) AS s FROM reservations"
            " WHERE user_id = ? AND status = 'paid'",
            (user_id,),
        ).fetchone()
        # рабочий день 15:00–03:00: заказ в 01:00 относится к предыдущей дате
        days = conn.execute(
            "SELECT COUNT(DISTINCT d) AS n, MAX(d) AS last FROM ("
            " SELECT date(created_at, '-6 hours') AS d FROM orders"
            "  WHERE user_id = ? AND status = 'paid'"
            " UNION SELECT date(created_at, '-6 hours') FROM bonus_moves"
            "  WHERE user_id = ? AND kind = 'offline'"
            " UNION SELECT date(COALESCE(paid_at, created_at), '-6 hours') FROM reservations"
            "  WHERE user_id = ? AND status = 'paid')",
            (user_id, user_id, user_id),
        ).fetchone()
        moves = [
            dict(r) for r in conn.execute(
                "SELECT * FROM bonus_moves WHERE user_id = ? ORDER BY id DESC LIMIT 7",
                (user_id,),
            )
        ]
        bal = _bonus_balance(conn, user_id)
    return {
        "orders": int(o["n"]), "spent": int(o["s"]) + int(off["s"]) + int(rent["s"]),
        "visits": int(days["n"] or 0), "last": days["last"], "balance": bal, "moves": moves,
    }


def db_guest_search(query: str) -> list[dict]:
    """Поиск карты по телефону (полностью или последние 10 цифр) или @нику."""
    q = (query or "").strip()
    with _connect() as conn:
        if q.startswith("@") or re.search(r"[A-Za-z_]", q):
            nick = q.lstrip("@").lower()
            if not nick:
                return []
            rows = conn.execute(
                "SELECT * FROM guests WHERE phone IS NOT NULL AND lower(username) = ?",
                (nick,),
            ).fetchall()
            return [dict(r) for r in rows]
        digits = re.sub(r"\D", "", q)
        if len(digits) < 4:
            return []
        norm = normalize_phone(q)
        rows = []
        if norm:
            rows = conn.execute("SELECT * FROM guests WHERE phone = ?", (norm,)).fetchall()
        if not rows:
            tail = digits[-10:]
            rows = conn.execute(
                "SELECT * FROM guests WHERE phone IS NOT NULL AND phone LIKE ? LIMIT 10",
                ("%" + tail,),
            ).fetchall()
    return [dict(r) for r in rows]


BONUS_KIND_LABELS = {
    "accrual": "кешбэк", "spend": "оплата баллами", "refund": "возврат",
    "offline": "чек вне бота", "manual": "списание админом",
}


def _fmt_pts(n: int) -> str:
    return f"{int(n):,}".replace(",", " ")


def card_text(user_id: int, for_admin: bool = False) -> str:
    g = db_guest_get(user_id) or {}
    st = guest_stats(user_id)
    lines = ["💎 <b>Карта лояльности Zig Zag</b>", ""]
    if for_admin:
        who = esc(g.get("first_name") or "Гость")
        if g.get("username"):
            who += f" (@{esc(g['username'])})"
        lines.append(f"👤 {who} · id <code>{user_id}</code>")
    phone = g.get("phone") or "—"
    mark = " ✅" if g.get("phone_verified") else (" (введён вручную)" if g.get("phone") else "")
    lines += [
        f"📱 Номер карты: <b>{esc(phone)}</b>{mark}",
        f"🎁 Баллы: <b>{_fmt_pts(st['balance'])}</b> (1 балл = 1 {CURRENCY})",
        f"💸 Кешбэк: <b>{LOYALTY_PERCENT}%</b> с оплаченных заказов"
        f" · оплатить баллами можно до {BONUS_MAX_PAY_PCT}% чека",
        f"🧾 Оплачено всего: {_fmt_money(st['spent'])} · заказов: {st['orders']}"
        f" · визитов: {st['visits']}",
    ]
    if st["last"]:
        lines.append(f"🕒 Последний визит: {st['last'][8:10]}.{st['last'][5:7]}.{st['last'][:4]}")
    if st["moves"]:
        lines += ["", "<b>Последние операции:</b>"]
        for m in st["moves"]:
            d = int(m["delta"])
            sign = "+" if d >= 0 else "−"
            when = (m.get("created_at") or "")[5:16]
            label = BONUS_KIND_LABELS.get(m["kind"], m["kind"])
            ref = f" · #{m['order_id']}" if m.get("order_id") else ""
            lines.append(f"  {sign}{_fmt_pts(abs(d))} · {label}{ref} · {when}")
    lines += ["", "ℹ️ На заказы по акциям («Выгода») баллы не начисляются и не списываются."]
    return "\n".join(lines)


# ========================= АНАЛИТИКА ПО КАРТАМ ЛОЯЛЬНОСТИ ======================
# Деньги гостя = заказы в боте (оплачено деньгами, без баллов) + чеки вне бота
# (введённые через /guest) + оплаченная аренда VIP/PlayStation.
# Визит = рабочий день (15:00–03:00), в который была оплата.
_MONTHS = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]
_TIME_BUCKETS = [(15, 18, "15–18"), (18, 21, "18–21"), (21, 24, "21–00"), (0, 3, "00–03")]


def _an_today():
    return (_now_dt() - timedelta(hours=6)).date()


def _an_day(ts: str):
    """Рабочий день оплаты: 01:00 ночи — это ещё предыдущая дата."""
    return (datetime.strptime(str(ts)[:19], "%Y-%m-%d %H:%M:%S") - timedelta(hours=6)).date()


def _an_events(conn: sqlite3.Connection, uid: int | None = None) -> list[dict]:
    """Все оплаты гостей: заказы, чеки вне бота, аренда. uid=None — всех."""
    flt, args = ("AND user_id = ?", (uid,)) if uid else ("AND user_id IS NOT NULL", ())
    ev: list[dict] = []
    for r in conn.execute(
        f"SELECT id, user_id, total - bonus_used AS m, kind, pay_method, created_at"
        f" FROM orders WHERE status = 'paid' {flt}", args,
    ):
        ev.append({"uid": r["user_id"], "src": "promo" if r["kind"] == "promo" else "order",
                   "m": int(r["m"] or 0), "day": _an_day(r["created_at"]),
                   "hour": int(str(r["created_at"])[11:13]), "pay": r["pay_method"]})
    for r in conn.execute(
        f"SELECT user_id, amount, created_at FROM bonus_moves WHERE kind = 'offline' {flt}", args
    ):
        ev.append({"uid": r["user_id"], "src": "offline", "m": int(r["amount"] or 0),
                   "day": _an_day(r["created_at"]), "hour": None, "pay": None})
    for r in conn.execute(
        f"SELECT user_id, price, pay_method, COALESCE(paid_at, created_at) AS paid"
        f" FROM reservations WHERE status = 'paid' {flt}", args
    ):
        try:
            day = _an_day(r["paid"])   # визит — рабочий день оплаты, как у заказов
        except (TypeError, ValueError):
            continue
        ev.append({"uid": r["user_id"], "src": "rent", "m": int(r["price"] or 0),
                   "day": day, "hour": None, "pay": r["pay_method"]})
    return ev


def _item_categories() -> dict[str, str]:
    out: dict[str, str] = {}
    for c in MENU["categories"]:
        label = f"{c.get('emoji', '')} {c['name']}".strip()
        for it in c["items"]:
            out[it["id"]] = label
    return out


def _pct(a: float, b: float) -> int:
    return int(round(100 * a / b)) if b else 0


def _an_date(d) -> str:
    return f"{d:%d.%m.%Y}"


def _guest_label(g: dict, phone: bool = True) -> str:
    who = g.get("first_name") or "Гость"
    if g.get("username"):
        who += f" (@{g['username']})"
    if phone and g.get("phone"):
        who += f" · {g['phone']}"
    return who


def guest_analytics(uid: int) -> dict:
    with _connect() as conn:
        ev = _an_events(conn, uid)
        items = [dict(r) for r in conn.execute(
            "SELECT oi.item_id, MAX(oi.name) AS name, MAX(o.kind) AS okind,"
            " SUM(oi.qty) AS q, SUM(oi.qty * oi.price) AS s"
            " FROM order_items oi JOIN orders o ON o.id = oi.order_id"
            " WHERE o.user_id = ? AND o.status = 'paid'"
            " GROUP BY COALESCE(oi.item_id, oi.name) ORDER BY q DESC, s DESC",
            (uid,),
        )]
        n_cancel = int(conn.execute(
            "SELECT COUNT(*) AS n FROM orders WHERE user_id = ? AND status = 'cancelled'", (uid,)
        ).fetchone()["n"])
        rsv = [dict(r) for r in conn.execute(
            "SELECT kind, status FROM reservations WHERE user_id = ?", (uid,)
        )]
        bal = _bonus_balance(conn, uid)
        spent_pts = int(conn.execute(
            "SELECT COALESCE(SUM(-delta), 0) AS s FROM bonus_moves WHERE user_id = ?"
            " AND delta < 0 AND kind IN ('spend', 'manual')", (uid,)
        ).fetchone()["s"]) - int(conn.execute(
            "SELECT COALESCE(SUM(delta), 0) AS s FROM bonus_moves WHERE user_id = ?"
            " AND delta > 0 AND kind = 'refund'", (uid,)
        ).fetchone()["s"])
    spent_pts = max(0, spent_pts)
    a: dict = {"items": items, "n_cancel": n_cancel, "balance": bal,
               "pts_spent": spent_pts, "pts_earned": bal + spent_pts}
    by = {k: [e for e in ev if e["src"] == k] for k in ("order", "promo", "offline", "rent")}
    a["m_order"] = sum(e["m"] for e in by["order"])
    a["n_order"] = len(by["order"])
    a["m_promo"] = sum(e["m"] for e in by["promo"])
    a["n_promo"] = len(by["promo"])
    a["m_offline"] = sum(e["m"] for e in by["offline"])
    a["n_offline"] = len(by["offline"])
    a["m_rent"] = sum(e["m"] for e in by["rent"])
    a["n_rent"] = len(by["rent"])
    a["total"] = a["m_order"] + a["m_promo"] + a["m_offline"] + a["m_rent"]
    bot_orders = by["order"] + by["promo"]
    a["avg_check"] = (a["m_order"] + a["m_promo"]) // len(bot_orders) if bot_orders else 0
    days = sorted({e["day"] for e in ev})
    today = _an_today()
    a["visits"] = len(days)
    a["first"] = days[0] if days else None
    a["last"] = days[-1] if days else None
    a["since_last"] = (today - days[-1]).days if days else None
    a["visits_30"] = sum(1 for d in days if (today - d).days < 30)
    a["avg_gap"] = round((days[-1] - days[0]).days / (len(days) - 1)) if len(days) > 1 else None
    a["avg_visit"] = a["total"] // len(days) if days else 0
    months: dict[str, int] = {}
    for e in ev:
        k = f"{e['day']:%Y-%m}"
        months[k] = months.get(k, 0) + e["m"]
    a["months"] = sorted(months.items(), reverse=True)[:6]
    wd: dict[int, int] = {}
    for d in days:
        wd[d.weekday()] = wd.get(d.weekday(), 0) + 1
    a["weekdays"] = sorted(wd.items(), key=lambda x: -x[1])
    tb: dict[str, int] = {}
    for e in bot_orders:
        for h0, h1, lbl in _TIME_BUCKETS:
            if h0 <= e["hour"] < h1:
                tb[lbl] = tb.get(lbl, 0) + 1
    a["time"] = sorted(tb.items(), key=lambda x: -x[1])
    a["n_timed"] = sum(tb.values())
    pays: dict[str, int] = {}
    for e in bot_orders + by["rent"]:
        if e["pay"]:
            pays[e["pay"]] = pays.get(e["pay"], 0) + 1
    a["pays"] = sorted(pays.items(), key=lambda x: -x[1])
    cats_map = _item_categories()
    cats: dict[str, int] = {}
    for it in items:
        if it["okind"] == "promo" or str(it["item_id"] or "").startswith("pr_"):
            c = "🎁 Акции"
        else:
            c = cats_map.get(it["item_id"] or "", "Прочее")
        cats[c] = cats.get(c, 0) + int(it["s"] or 0)
    a["cats"] = sorted(cats.items(), key=lambda x: -x[1])
    a["cats_total"] = sum(cats.values())
    a["rsv_total"] = len(rsv)
    a["rsv_kinds"] = {}
    a["rsv_st"] = {}
    for r in rsv:
        a["rsv_kinds"][r["kind"]] = a["rsv_kinds"].get(r["kind"], 0) + 1
        a["rsv_st"][r["status"]] = a["rsv_st"].get(r["status"], 0) + 1
    seg: list[str] = []
    if not days:
        seg.append("⚪ Ещё не было оплат")
    else:
        sleeping = a["since_last"] > SEG_SLEEP_DAYS
        if a["visits"] == 1 and not sleeping:
            seg.append("🆕 Новый")
        if a["total"] >= SEG_VIP_SPENT or a["visits"] >= SEG_VIP_VISITS:
            seg.append("💎 Ценный гость")
        if a["visits_30"] >= SEG_OFTEN_30:
            seg.append("🔥 Ходит часто")
        elif a["visits"] >= 3 and not sleeping and "💎 Ценный гость" not in seg:
            seg.append("🔁 Постоянный")
        if sleeping:
            seg.append(f"😴 Давно не был ({a['since_last']} дн.)")
    if a["rsv_st"].get("no_show", 0) >= 2:
        seg.append("⚠️ Часто не приходит на брони")
    a["segments"] = seg
    return a


def guest_analytics_text(uid: int) -> str:
    g = db_guest_get(uid) or {}
    a = guest_analytics(uid)
    L = ["📊 <b>Аналитика гостя</b>", f"👤 {esc(_guest_label(g))}",
         f"🏷 {' · '.join(a['segments'])}", ""]
    L.append("<b>💰 Деньги</b>")
    L.append(f"Всего потрачено: <b>{_fmt_money(a['total'])}</b>")
    if a["n_order"]:
        L.append(f"  • заказы в боте: {_fmt_money(a['m_order'])} ({a['n_order']} шт.)")
    if a["n_promo"]:
        L.append(f"  • заказы по акции: {_fmt_money(a['m_promo'])} ({a['n_promo']} шт.)")
    if a["n_offline"]:
        L.append(f"  • чеки вне бота: {_fmt_money(a['m_offline'])} ({a['n_offline']} шт.)")
    if a["n_rent"]:
        L.append(f"  • аренда VIP/PS: {_fmt_money(a['m_rent'])} ({a['n_rent']} шт.)")
    if a["avg_check"] or a["avg_visit"]:
        L.append(f"Средний чек в боте: {_fmt_money(a['avg_check'])} · за визит: {_fmt_money(a['avg_visit'])}")
    if a["months"]:
        L.append("По месяцам: " + " · ".join(
            f"{_MONTHS[int(k[5:7]) - 1]} {k[2:4]} — {_fmt_money(v)}" for k, v in a["months"]))
    L += ["", "<b>📅 Визиты</b>"]
    if a["visits"]:
        L.append(f"Визитов: <b>{a['visits']}</b> · за 30 дней: {a['visits_30']}")
        L.append(f"Первый: {_an_date(a['first'])} · последний: {_an_date(a['last'])}"
                 f" ({'сегодня' if a['since_last'] == 0 else str(a['since_last']) + ' дн. назад'})")
        if a["avg_gap"] is not None:
            L.append(f"Приходит в среднем раз в {max(1, a['avg_gap'])} дн.")
        if a["weekdays"]:
            L.append("Любимые дни: " + ", ".join(WDAYS[d] for d, _ in a["weekdays"][:2]))
        if a["time"]:
            t, n = a["time"][0]
            L.append(f"Обычно заказывает: {t} ({_pct(n, a['n_timed'])}% заказов)")
    else:
        L.append("Оплаченных визитов пока нет")
    L += ["", "<b>❤️ Что заказывает</b>"]
    if a["items"]:
        for i, it in enumerate(a["items"][:5], 1):
            L.append(f"{i}. {esc(it['name'])} — {int(it['q'])} шт · {_fmt_money(it['s'] or 0)}")
        if a["cats"]:
            L.append("Категории: " + " · ".join(
                f"{esc(c)} {_pct(v, a['cats_total'])}%" for c, v in a["cats"][:4]))
    else:
        L.append("Оплаченных заказов в боте пока нет")
    L.append("")
    if a["pays"]:
        L.append("💳 Оплата: " + " · ".join(f"{PAY_METHODS.get(k, k)} {n}" for k, n in a["pays"]))
    L.append(f"🎁 Баллы: получено {_fmt_pts(a['pts_earned'])} · потрачено {_fmt_pts(a['pts_spent'])}"
             f" · баланс <b>{_fmt_pts(a['balance'])}</b>")
    if a["rsv_total"]:
        kinds = " · ".join(f"{RSV_KINDS.get(k, k)} {n}" for k, n in a["rsv_kinds"].items())
        st = a["rsv_st"]
        came = st.get("paid", 0) + st.get("confirmed", 0)
        L.append(f"📌 Брони: {a['rsv_total']} ({kinds})")
        L.append(f"   подтверждено/оплачено {came} · отменено {st.get('cancelled', 0)}"
                 f" · не пришёл {st.get('no_show', 0)}")
    if a["n_cancel"]:
        L.append(f"❌ Отменённых заказов: {a['n_cancel']}")
    L += ["", "ℹ️ Учитываются оплаты через бота, чеки, внесённые через /guest, и аренда."]
    return "\n".join(L)


def loyalty_overview(days: int | None) -> dict:
    """Сводка по программе лояльности за последние days дней (None — всё время)."""
    today = _an_today()
    since = today - timedelta(days=days - 1) if days else None
    with _connect() as conn:
        cards = {int(r["user_id"]): dict(r) for r in conn.execute(
            "SELECT * FROM guests WHERE phone IS NOT NULL")}
        ev = _an_events(conn)
        n_all = n_card = 0
        rev_all = rev_card = 0
        for r in conn.execute(
            "SELECT user_id, total - bonus_used AS m, created_at FROM orders WHERE status = 'paid'"
        ):
            if since and _an_day(r["created_at"]) < since:
                continue
            n_all += 1
            rev_all += int(r["m"] or 0)
            if r["user_id"] and int(r["user_id"]) in cards:
                n_card += 1
                rev_card += int(r["m"] or 0)
        moves = [dict(r) for r in conn.execute(
            "SELECT user_id, delta, kind, created_at FROM bonus_moves")]
        outstanding = sum(int(m["delta"]) for m in moves)
    o = {"days": days, "since": since, "today": today, "cards": len(cards),
         "verified": sum(1 for g in cards.values() if g.get("phone_verified")),
         "new_cards": sum(1 for g in cards.values()
                          if not since or _an_day(g["created_at"]) >= since),
         "rev_all": rev_all, "rev_card": rev_card, "n_all": n_all, "n_card": n_card,
         "outstanding": outstanding}
    pm = [m for m in moves if not since or _an_day(m["created_at"]) >= since]
    o["pts_earned"] = sum(int(m["delta"]) for m in pm
                          if int(m["delta"]) > 0 and m["kind"] in ("accrual", "offline", "manual"))
    o["pts_spent"] = sum(-int(m["delta"]) for m in pm
                         if int(m["delta"]) < 0 and m["kind"] in ("spend", "manual"))
    per: dict[int, dict] = {}
    for e in ev:
        uid = int(e["uid"])
        if uid not in cards:
            continue
        p = per.setdefault(uid, {"all_m": 0, "all_days": set(), "m": 0, "days": set()})
        p["all_m"] += e["m"]
        p["all_days"].add(e["day"])
        if not since or e["day"] >= since:
            p["m"] += e["m"]
            p["days"].add(e["day"])
    active = {u: p for u, p in per.items() if p["days"]}
    o["active"] = len(active)
    o["repeat"] = sum(1 for p in active.values() if len(p["days"]) >= 2)
    o["top"] = [(u, p["m"], len(p["days"])) for u, p in
                sorted(active.items(), key=lambda x: -x[1]["m"])[:5]]
    sleepers = [
        (u, p["all_m"], max(p["all_days"])) for u, p in per.items()
        if len(p["all_days"]) >= 2 and (today - max(p["all_days"])).days > SEG_SLEEP_DAYS
    ]
    sleepers.sort(key=lambda x: -x[1])
    o["sleepers_n"] = len(sleepers)
    o["sleepers"] = sleepers[:5]
    o["names"] = {u: cards[u] for u in cards}
    o["per"] = per
    o["card_days"] = [_an_day(g["created_at"]) for g in cards.values() if g.get("created_at")]
    return o


_LA_PERIODS = {"7": 7, "30": 30, "90": 90, "all": None}


def loyalty_text(o: dict) -> str:
    title = "всё время" if not o["days"] else f"{o['days']} дн."
    span = f" ({o['since']:%d.%m}–{o['today']:%d.%m})" if o["since"] else ""
    avg_card = o["rev_card"] // o["n_card"] if o["n_card"] else 0
    n_other = o["n_all"] - o["n_card"]
    avg_other = (o["rev_all"] - o["rev_card"]) // n_other if n_other else 0
    L = [f"📊 <b>Лояльность · {title}</b>{span}", "",
         f"💳 Карт: <b>{o['cards']}</b> (✅ подтверждено {o['verified']}) · новых: {o['new_cards']}",
         f"👥 Гостей с картой, которые платили: <b>{o['active']}</b> · приходили 2+ раз: "
         f"{o['repeat']} ({_pct(o['repeat'], o['active'])}%)",
         f"💰 Выручка заказов в боте: {_fmt_money(o['rev_all'])}, из них по картам "
         f"{_fmt_money(o['rev_card'])} ({_pct(o['rev_card'], o['rev_all'])}%)",
         f"🧾 Средний чек: с картой {_fmt_money(avg_card) if o['n_card'] else '—'}"
         f" · без карты {_fmt_money(avg_other) if n_other else '—'}",
         f"🎁 Баллы: начислено {_fmt_pts(o['pts_earned'])} · списано {_fmt_pts(o['pts_spent'])}",
         f"   на счетах гостей сейчас: <b>{_fmt_pts(o['outstanding'])}</b> (= ₽ будущих скидок)"]
    if o["top"]:
        L += ["", "🏆 <b>Топ гостей за период</b>"]
        for i, (u, m, n) in enumerate(o["top"], 1):
            L.append(f"{i}. {esc(_guest_label(o['names'][u]))} — {_fmt_money(m)} · визитов {n}")
    L += ["", f"😴 <b>Давно не были</b> (>{SEG_SLEEP_DAYS} дн., ходили 2+ раз): {o['sleepers_n']}"]
    for u, m, last in o["sleepers"]:
        L.append(f"• {esc(_guest_label(o['names'][u]))} — всего {_fmt_money(m)}"
                 f" · последний визит {last:%d.%m}")
    L += ["", "Кнопки ниже открывают карту гостя 👇" if (o["top"] or o["sleepers"]) else ""]
    return "\n".join(L).rstrip()


def kb_loyalty(o: dict, cur: str) -> InlineKeyboardMarkup:
    rows = [[_btn(("• " if k == cur else "") + ("Всё время" if k == "all" else f"{k} дн."),
                  f"la:{k}") for k in _LA_PERIODS]]
    btns = [_btn(f"🏆 {i}. {(o['names'][u].get('first_name') or 'Гость')[:18]}", f"ga:o:{u}")
            for i, (u, _, _) in enumerate(o["top"], 1)]
    btns += [_btn(f"😴 {(o['names'][u].get('first_name') or 'Гость')[:18]}", f"ga:o:{u}")
             for u, _, _ in o["sleepers"]]
    rows += [btns[i : i + 2] for i in range(0, len(btns), 2)]
    rows.append([_btn("✖ Закрыть", "ga:close")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


# ======================= ХРАНИЛИЩЕ ДИАЛОГОВ (FSM) ============================
class SQLiteStorage(BaseStorage):
    """FSM в той же SQLite-БД: корзина, стол, шаг брони и т.п. переживают
    перезапуск/деплой. Через FSM_TTL_HOURS бездействия запись сбрасывается
    (вчерашний стол и корзина не «всплывут» на следующий день)."""

    def __init__(self, ttl_hours: int = FSM_TTL_HOURS) -> None:
        self.ttl = float(ttl_hours) * 3600

    @staticmethod
    def _k(key: StorageKey) -> str:
        return (
            f"{key.bot_id}:{key.chat_id}:{key.user_id}:{key.thread_id or ''}:"
            f"{key.business_connection_id or ''}:{key.destiny}"
        )

    def _row(self, k: str):
        with _connect() as conn:
            row = conn.execute(
                "SELECT state, data, updated_at FROM fsm WHERE k = ?", (k,)
            ).fetchone()
            if row and time.time() - float(row["updated_at"]) > self.ttl:
                conn.execute("DELETE FROM fsm WHERE k = ?", (k,))
                return None
        return row

    def _write(self, k: str, **fields) -> None:
        self._row(k)  # просроченную запись сначала удаляем
        now = time.time()
        with _connect() as conn:
            exists = conn.execute("SELECT 1 FROM fsm WHERE k = ?", (k,)).fetchone()
            if exists:
                sets = ", ".join(f"{f} = ?" for f in fields) + ", updated_at = ?"
                conn.execute(
                    f"UPDATE fsm SET {sets} WHERE k = ?", (*fields.values(), now, k)
                )
            else:
                conn.execute(
                    "INSERT INTO fsm (k, state, data, updated_at) VALUES (?, ?, ?, ?)",
                    (k, fields.get("state"), fields.get("data", "{}"), now),
                )

    async def set_state(self, key: StorageKey, state: StateType = None) -> None:
        st = state.state if isinstance(state, State) else state
        self._write(self._k(key), state=st)

    async def get_state(self, key: StorageKey) -> str | None:
        row = self._row(self._k(key))
        return row["state"] if row else None

    async def set_data(self, key: StorageKey, data: Mapping[str, Any]) -> None:
        js = json.dumps(dict(data), ensure_ascii=False, default=str)
        self._write(self._k(key), data=js)

    async def get_data(self, key: StorageKey) -> dict[str, Any]:
        row = self._row(self._k(key))
        if not row:
            return {}
        try:
            d = json.loads(row["data"] or "{}")
        except (TypeError, ValueError):
            return {}
        return d if isinstance(d, dict) else {}

    async def close(self) -> None:
        return None

    def purge_expired(self) -> int:
        with _connect() as conn:
            cur = conn.execute(
                "DELETE FROM fsm WHERE updated_at < ?", (time.time() - self.ttl,)
            )
            return cur.rowcount


# --- Статусы бронирования ---
RSV_STATUSES: dict[str, str] = {
    "new": "🟡 Ожидает",
    "confirmed": "✅ Подтверждена",
    "cancelled": "❌ Отменена",
    "no_show": "🚫 Не пришёл",
    "paid": "💳 Оплачен",
}

# --- Способы оплаты заказа (запрашиваем у админа при статусе «Оплачен») ---
PAY_METHODS: dict[str, str] = {
    "cash": "нал",
    "card": "безнал",
    "transfer": "перевод",
}
PAY_BUTTONS: dict[str, str] = {
    "cash": "💵 Наличные",
    "card": "💳 Безнал",
    "transfer": "🔄 Перевод",
}


def _fmt_money(n) -> str:
    """15520 → «15 520 ₽»."""
    return f"{int(n):,}".replace(",", " ") + f" {CURRENCY}"

WELCOME = (
    "💨 Добро пожаловать!\n\n"
    "Это бот кальянной Zig Zag⚡\n\n"
    "Хотите забронировать стол/vip комнату ? Или сделать заказ ?"
)

# --- сценарий лояльности (первый раз / постоянный + телефон) ---
LOYAL_ASK = "💳 Вы первый раз или постоянный гость?"
LOYAL_FIRST = "🌟 Отлично, что зашли!\n\nОставьте номер телефона — заведём карту лояльности 🎁"
LOYAL_REG = "💎 С возвращением!\n\nВведите номер телефона — это номер вашей карты лояльности:"
PHONE_ASK = "📱 Отправьте номер кнопкой ниже или просто напишите его:"
PHONE_SKIP = "⏭ Пропустить"
PHONE_BACK = "🔙 Назад"

ASK_TABLE = "🪑 Введите номер стола цифрами (от 1 до 8), например <b>5</b>:"
ERR_TABLE = "⚠️ Номер стола — число от 1 до 8 (VIP-комната и барная стойка бронируются отдельно). Попробуйте ещё раз:"
ERR_MENU = "⚠️ Эта позиция больше недоступна."
CART_EMPTY = "🛒 Корзина пуста."
ASK_TABLE_FIRST = "Сначала укажите номер стола — отправьте /start"

BTN_MENU = "📋 Меню"
BTN_CART = "🛒 Корзина"
BTN_TABLE = "🪑 Стол: {table} (сменить)"
BTN_BOOK = "📅 Бронирование"
BTN_CARD = "💎 Моя карта"
BTN_PRE = "📝 Предзаказ"


def main_text(table: str) -> str:
    return (
        f"🪑 Стол: <b>{esc(table)}</b>\n\n"
        "Выберите раздел каталога или откройте корзину.\n"
        "Когда всё выберете — оформите заказ, и он сразу уйдёт персоналу в канал."
    )


def cart_text(cart: list[dict], table: str | None, title: str = "🛒 <b>Корзина</b>") -> str:
    if not cart:
        return CART_EMPTY
    lines = [title]
    if table:
        lines.append(f"🪑 Стол: <b>{esc(table)}</b>")
    lines.append("")
    total = 0
    for n, i in enumerate(cart, 1):
        sum_i = i["price"] * i["qty"]
        total += sum_i
        lines.append(f"{n}. {esc(i['name'])} × {i['qty']} — {_fmt_money(sum_i)}")
    lines += ["", f"💰 <b>Итого: {_fmt_money(total)}</b>"]
    return "\n".join(lines)


def confirm_text(cart: list[dict], table: str | None, promo: bool = False) -> str:
    total = sum(i["price"] * i["qty"] for i in cart)
    head = "✅ <b>Подтверждение заказа по акции</b>" if promo else "✅ <b>Подтверждение заказа</b>"
    lines = [head, ""]
    lines.append(f"🪑 Стол: <b>{esc(table) if table else 'не указан'}</b>")
    lines.append("")
    for i in cart:
        lines.append(f"• {esc(i['name'])} × {i['qty']} — {_fmt_money(i['price'] * i['qty'])}")
    lines += ["", f"💰 <b>К оплате: {_fmt_money(total)}</b>"]
    if promo:
        lines.append("ℹ️ Акции не суммируются: баллы не начисляются и не списываются")
    lines += ["", "Отправить заказ?"]
    return "\n".join(lines)


def sent_text(order_id: int, table: str, total: int) -> str:
    return (
        f"✅ Заказ <b>#{order_id}</b> отправлен!\n"
        f"🪑 Стол: <b>{esc(table)}</b>\n"
        f"💰 Сумма: <b>{_fmt_money(total)}</b>\n\n"
        "Персонал подтвердит заказ — как только он будет готов, вам сообщат.\n"
        "Можете продолжать выбирать: /start"
    )


def table_set_text(table: str) -> str:
    return f"🪑 Стол <b>{esc(table)}</b> установлен.\n\nВыбирайте раздел каталога 👇"


def order_text(order: dict) -> str:
    total = sum(i["price"] * i["qty"] for i in order["items"])
    promo = order.get("kind") == "promo"
    if order.get("kind") == "pre":
        lines = [f"📝 <b>Предзаказ #{order['id']}</b>",
                 f"⏰ К приходу: <b>{_pre_when(order)}</b>"]
        if order.get("rsv_id"):
            lines.append(f"📅 К брони #{order['rsv_id']}")
        lines.append(f"🕒 Оформлен {order['created_at'][8:10]}.{order['created_at'][5:7]}"
                     f" {order['created_at'][11:16]}")
    else:
        head = (f"🔥 <b>Заказ по акции #{order['id']}</b>" if promo
                else f"🧾 <b>Заказ #{order['id']}</b>")
        lines = [
            head,
            f"🪑 Стол: <b>{esc(order['table_no'])}</b>",
            f"🕒 {order['created_at'][11:16]}",
        ]
    who = f"@{esc(order['username'])}" if order.get("username") else ""
    card = guest_card(order.get("user_id"))
    if card:
        bal = db_bonus_balance(order["user_id"])
        lines.append(f"💎 {who} {esc(card['phone'])} · баллов: {_fmt_pts(bal)}".replace("  ", " "))
    elif who:
        lines.append(f"👤 {who} · без карты")
    lines.append("")
    for i in order["items"]:
        lines.append(f"• {esc(i['name'])} × {i['qty']} — {_fmt_money(i['price'] * i['qty'])}")
    lines += ["", f"💰 <b>Итого: {_fmt_money(total)}</b>"]
    used = int(order.get("bonus_used") or 0)
    pending = int(order.get("bonus_pending") or 0)
    if used:
        lines.append(f"🎁 Баллами: −{_fmt_money(used)} · деньгами: <b>{_fmt_money(total - used)}</b>")
    elif pending and order.get("status") != "paid":
        lines.append(f"🎁 Спишем баллами: −{_fmt_money(pending)} · к оплате: <b>{_fmt_money(total - pending)}</b>")
    if int(order.get("bonus_accrued") or 0):
        lines.append(f"💎 Начислено кешбэка: +{_fmt_pts(order['bonus_accrued'])}")
    if promo:
        lines.append("ℹ️ Акция: баллы не начисляются и не списываются")
    lines.append("")
    status = STATUSES.get(order["status"], order["status"])
    if order["status"] == "paid" and order.get("pay_method"):
        status += f" · {PAY_METHODS.get(order['pay_method'], order['pay_method'])}"
    ts = (order.get("updated_at") or "")[11:16]
    lines.append(f"Статус: {status} · {ts}")
    return "\n".join(lines)


def _order_notify_text(o: dict) -> str:
    """Личное сообщение гостю о смене статуса его заказа."""
    pre = o.get("kind") == "pre"
    if pre:
        head = f"📝 Предзаказ <b>#{o['id']}</b> · к <b>{_pre_when(o)}</b>"
    else:
        head = f"🧾 Заказ <b>#{o['id']}</b> · стол <b>{esc(o['table_no'])}</b>"
    s = o["status"]
    if s == "accepted":
        if pre:
            return f"🟡 {head}\nПринят! Подготовим к вашему приходу."
        return f"🟡 {head}\nПринят! Персонал уже занимается вашим заказом."
    if s == "ready":
        if pre:
            return f"🟢 {head}\nГотов!"
        return f"🟢 {head}\nГотов! Скоро Вам его принесут."
    if s == "issued":
        return f"🔵 {head}\nВыдан. Приятного отдыха! 💨"
    if s == "paid":
        m = PAY_METHODS.get(o.get("pay_method") or "", "")
        tail = f" ({m})" if m else ""
        return f"💵 {head}\nОплата получена{tail}. Спасибо!"
    if s == "cancelled":
        return f"❌ {head}\nК сожалению, заказ отменён. Вопросы — администратору."
    return f"📋 {head}\nСтатус: {STATUSES.get(s, s)}"


# ============================== КЛАВИАТУРЫ ==================================


def kb_main(table: str | None) -> ReplyKeyboardMarkup:
    rows = [[KeyboardButton(text=BTN_MENU), KeyboardButton(text=BTN_CART)]]
    rows.append([KeyboardButton(text=BTN_BOOK), KeyboardButton(text=BTN_CARD)])
    last = [KeyboardButton(text=BTN_PRE)]
    if table:
        last.append(KeyboardButton(text=BTN_TABLE.format(table=table)))
    rows.append(last)
    return ReplyKeyboardMarkup(keyboard=rows, resize_keyboard=True)


def _btn(text: str, data: str) -> InlineKeyboardButton:
    return InlineKeyboardButton(text=text, callback_data=data)


# --- фото «🔥 Выгода» и «📜 Правила» ---
# Файлы promo.jpg и rules.jpg лежат рядом с bot.py (в корне репозитория).
# Заменить картинку — загрузить новый файл с тем же именем и перезапустить бота.
PROMO_PHOTO = Path(__file__).with_name("promo.jpg")   # «🔥 Выгода»
RULES_PHOTO = Path(__file__).with_name("rules.jpg")   # «📜 Правила»


def kb_categories(cart_count: int = 0) -> InlineKeyboardMarkup:
    rows = [[_btn(f"{c['emoji']} {c['name']}", f"c:{c['id']}")] for c in MENU["categories"]]
    cart_label = BTN_CART + (f" ({cart_count})" if cart_count else "")
    rows.append([_btn("🔥 Выгода", "info:benefit"), _btn("🎁 Заказ по акции", "p:home")])
    rows.append([_btn("📜 Правила", "info:rules")])
    rows.append([_btn(cart_label, "cv")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_info_photo(promo: bool = False) -> InlineKeyboardMarkup:
    """Под фотоакцией/правилами — «Назад» (удаляет фото, меню остаётся);
    под «Выгодой» ещё и «Заказать по акции»."""
    rows = []
    if promo:
        rows.append([_btn("🎁 Заказать по акции", "p:new")])
    rows.append([_btn("🔙 Назад", "info:close")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_category(cat_id: str, cart_count: int = 0) -> InlineKeyboardMarkup:
    cat = f_category(cat_id) or {"items": []}
    reserved = reserved_snapshot()
    rows = [
        [_btn(f"＋ {it['name']} — {_fmt_money(it['price'])}", f"a:{it['id']}")]
        for it in cat["items"]
        if item_in_stock(it["id"], reserved)  # распроданное скрываем
    ]
    cart_label = BTN_CART + (f" ({cart_count})" if cart_count else "")
    rows.append([_btn(cart_label, "cv")])
    rows.append([_btn("🔙 Назад", "m")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_cart(cart: list[dict]) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    for idx in range(len(cart)):
        rows.append(
            [
                _btn("➖", f"cd:{idx}"),
                _btn(f"{cart[idx]['qty']} шт.", "noop"),
                _btn("➕", f"ci:{idx}"),
                _btn("🗑", f"cr:{idx}"),
            ]
        )
    rows.append([_btn("✅ Оформить заказ", "co")])
    rows.append([_btn("🗑 Очистить корзину", "cx")])
    rows.append([_btn("🔙 В меню", "m")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_confirm() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [_btn("✅ Отправить заказ", "so")],
            [_btn("🔙 Назад к корзине", "cb")],
        ]
    )


def kb_status(order_id: int) -> InlineKeyboardMarkup:
    order = db_get_order(order_id)
    if order and order["status"] == "paid":
        # Оплачено — рабочие кнопки статусов скрыты, доступен только возврат
        return InlineKeyboardMarkup(
            inline_keyboard=[
                [_btn("↩️ Оформить возврат", f"st:{order_id}:cancelled")]
            ]
        )
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                _btn(STATUSES["accepted"], f"st:{order_id}:accepted"),
                _btn(STATUSES["ready"], f"st:{order_id}:ready"),
            ],
            [
                _btn(STATUSES["issued"], f"st:{order_id}:issued"),
                _btn(STATUSES["paid"], f"st:{order_id}:paid"),
            ],
            [_btn(STATUSES["cancelled"], f"st:{order_id}:cancelled")],
        ]
    )


def kb_pay(order_id: int) -> InlineKeyboardMarkup:
    """Кнопки уточнения способа оплаты (после нажатия «Оплачен»).
    Если у гостя есть баллы — можно отложить их списание в этот чек."""
    rows = [
        [
            _btn(PAY_BUTTONS["cash"], f"pay:{order_id}:cash"),
            _btn(PAY_BUTTONS["card"], f"pay:{order_id}:card"),
        ],
        [_btn(PAY_BUTTONS["transfer"], f"pay:{order_id}:transfer")],
    ]
    order = db_get_order(order_id)
    if order and order["status"] != "paid":
        pending = int(order.get("bonus_pending") or 0)
        if pending:
            rows.append([_btn(f"↩️ Не списывать баллы ({_fmt_pts(pending)})", f"pay:{order_id}:nobonus")])
        else:
            avail = bonus_available_for_order(order)
            if avail > 0:
                rows.append([_btn(f"🎁 Списать баллы: −{_fmt_pts(avail)}", f"pay:{order_id}:bonus")])
    rows.append([_btn("🔙 Назад", f"pay:{order_id}:back")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


# =============================== ХЕНДЕЛЕРЫ ==================================

router = Router()
# Бот работает только в личных чатах: в группах молчит — не отвечает на каждое
# сообщение и не публикует админ-данные (/stats, /guest…) посторонним.
# Кнопки персонала в канале — это нажатия (callback), их фильтр не касается.
router.message.filter(F.chat.type == "private")


class OrderFSM(StatesGroup):
    table = State()     # гость вводит номер стола
    phone = State()     # гость вводит телефон карты лояльности
    rsv_name = State()  # гость вводит имя, на кого бронь


# ------------------------------------------------ утилиты гостя --------------
async def _cart(state: FSMContext) -> list[dict]:
    return (await state.get_data()).get("cart", [])


async def _table(state: FSMContext) -> str | None:
    return (await state.get_data()).get("table")


def _cart_count(cart: list[dict]) -> int:
    return sum(i["qty"] for i in cart)


def _cart_total(cart: list[dict]) -> int:
    return sum(i["price"] * i["qty"] for i in cart)


async def _show_menu(target, state: FSMContext, edit: bool = False) -> None:
    cart = await _cart(state)
    text = "📋 Выберите раздел:"
    kb = kb_categories(_cart_count(cart))
    if edit:
        try:
            await target.edit_text(text, reply_markup=kb)
        except TelegramBadRequest:
            pass
    else:
        await target.answer(text, reply_markup=kb)


# ============================ ГРАФИКИ (PNG) =================================
# Картинки к /stats и /loyalty. matplotlib импортируется лениво при первом
# графике (не тратит память на старте). Нет библиотеки или CHARTS=0 —
# бот просто присылает текст, как раньше. Рисуем без pyplot (Figure + Agg):
# потокобезопасно, без глобального состояния; рендер — в отдельном потоке.
_MPL: Any = None                      # None — ещё не пробовали; False — недоступно
_CHART_LOCK = asyncio.Lock()          # один рендер за раз (экономим память/CPU)
_CHART_MSGS: dict[tuple[int, int], int] = {}  # (chat, id текста) → id фото с графиком
_CH_BG, _CH_PANEL, _CH_FG, _CH_MUTED = "#15151c", "#20202b", "#ECECF1", "#9a9aae"
_CH_GRID = "#2c2c3a"
_CH_COLORS = ("#ff9f43", "#4ecdc4", "#a29bfe", "#ff6b81", "#feca57", "#54a0ff",
              "#1dd1a1", "#c8d6e5")


def _mpl():
    """(Figure, FigureCanvasAgg) или False, если графики недоступны."""
    global _MPL
    if _MPL is None:
        if not CHARTS_ENABLED:
            _MPL = False
            return _MPL
        try:
            from matplotlib.backends.backend_agg import FigureCanvasAgg
            from matplotlib.figure import Figure

            _MPL = (Figure, FigureCanvasAgg)
        except Exception as e:  # noqa: BLE001 — нет пакета/битая установка
            logger.warning("Графики выключены: matplotlib недоступен (%s). "
                           "Добавьте matplotlib в requirements.txt", e)
            _MPL = False
    return _MPL


def _plain(s: Any, n: int = 24) -> str:
    """Подпись для картинки: без эмодзи (в шрифте их нет) и не длиннее n."""
    t = re.sub(r"[\U00010000-\U0010FFFF\u2600-\u27BF\uFE0F\u200D]", "", str(s or ""))
    t = " ".join(t.split()) or "—"
    if len(t) <= n:
        return t
    # объём/вес в конце сохраняем: «Вода с газом… 0,65 л» ≠ «… 0,5 л»
    m = re.search(r"\s(\d+(?:[.,]\d+)?\s?(?:л|мл|г|кг|шт\.?))$", t)
    if m and len(m.group(1)) < n - 8:
        tail = m.group(1)
        return t[: n - len(tail) - 2].rstrip(" /,") + "… " + tail
    return t[: n - 1].rstrip() + "…"


def _money_short(n: float) -> str:
    n = int(round(n or 0))
    if abs(n) >= 1_000_000:
        return f"{n / 1_000_000:.1f}".replace(".", ",").replace(",0", "") + " млн"
    if abs(n) >= 10_000:
        return f"{n / 1000:.0f} тыс"
    return f"{n:,}".replace(",", " ")


def _ch_new_fig(height: float = 11.6):
    Figure, Canvas = _MPL
    fig = Figure(figsize=(10, height), dpi=110, facecolor=_CH_BG)
    Canvas(fig)
    return fig


def _ch_png(fig) -> bytes:
    import io

    buf = io.BytesIO()
    fig.canvas.print_png(buf)
    return buf.getvalue()


def _ch_title(fig, title: str, sub: str) -> None:
    fig.text(0.04, 0.962, title, color=_CH_FG, fontsize=21, fontweight="bold", va="center")
    fig.text(0.96, 0.962, sub, color=_CH_MUTED, fontsize=13, ha="right", va="center")


def _ch_tiles(fig, tiles: list[tuple[str, str, str]], top: float = 0.925,
              row_h: float = 0.085, gap: float = 0.012) -> float:
    """Плитки KPI по 3 в ряд: (подпись, значение, мелкая приписка). → нижняя граница."""
    from matplotlib.patches import FancyBboxPatch

    cols, x0, w_all = 3, 0.04, 0.92
    w = (w_all - gap * (cols - 1)) / cols
    y = top
    for i, (label, value, note) in enumerate(tiles):
        c = i % cols
        if c == 0 and i:
            y -= row_h + gap
        x = x0 + c * (w + gap)
        fig.add_artist(FancyBboxPatch(
            (x, y - row_h), w, row_h, boxstyle="round,pad=0,rounding_size=0.012",
            transform=fig.transFigure, facecolor=_CH_PANEL, edgecolor="none"))
        fig.add_artist(FancyBboxPatch(
            (x, y - row_h), 0.006, row_h, boxstyle="square,pad=0",
            transform=fig.transFigure, facecolor=_CH_COLORS[i % len(_CH_COLORS)],
            edgecolor="none"))
        fig.text(x + 0.022, y - 0.022, label, color=_CH_MUTED, fontsize=11, va="center")
        fig.text(x + 0.022, y - 0.052, value, color=_CH_FG, fontsize=19,
                 fontweight="bold", va="center")
        if note:
            fig.text(x + 0.022, y - 0.074, note, color=_CH_MUTED, fontsize=9.5, va="center")
    return y - row_h


def _ch_axes(fig, rect, title: str):
    ax = fig.add_axes(rect, facecolor=_CH_BG)
    ax.set_title(title, color=_CH_FG, fontsize=13.5, fontweight="bold", loc="left", pad=10)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(_CH_GRID)
    ax.tick_params(colors=_CH_MUTED, labelsize=9.5)
    return ax


def _ch_empty(ax, text: str = "нет данных за период") -> None:
    ax.set_xticks([])
    ax.set_yticks([])
    for s in ax.spines.values():
        s.set_visible(False)
    ax.text(0.5, 0.5, text, color=_CH_MUTED, fontsize=12, ha="center", va="center",
            transform=ax.transAxes)


def _ch_bars(ax, labels: list[str], values: list[float], color: str, money: bool = True) -> None:
    if not values or not any(values):
        _ch_empty(ax)
        return
    xs = list(range(len(values)))
    ax.bar(xs, values, color=color, width=0.72, zorder=3)
    ax.grid(axis="y", color=_CH_GRID, linewidth=0.8, zorder=0)
    ax.set_xlim(-0.6, len(values) - 0.4)
    step = max(1, (len(labels) + 13) // 14)        # не больше ~14 подписей
    ax.set_xticks(xs[::step])
    ax.set_xticklabels(labels[::step])
    ax.yaxis.set_major_formatter(
        __import__("matplotlib.ticker", fromlist=["FuncFormatter"]).FuncFormatter(
            lambda v, _p: _money_short(v) if money else f"{v:g}"))
    if len(values) <= 16:
        top = max(values)
        for x, v in zip(xs, values):
            if v:
                ax.text(x, v + top * 0.015, _money_short(v) if money else f"{v:g}",
                        color=_CH_FG, fontsize=8.5, ha="center", va="bottom")
    ax.set_ylim(0, max(values) * 1.14)


def _ch_hbars(ax, labels: list[str], values: list[float], color: str, fmt) -> None:
    if not values:
        _ch_empty(ax)
        return
    ys = list(range(len(values)))[::-1]
    ax.barh(ys, values, color=color, height=0.62, zorder=3)
    ax.set_yticks(ys)
    ax.set_yticklabels(labels, color=_CH_FG, fontsize=10)
    ax.set_xticks([])
    ax.spines["bottom"].set_visible(False)
    top = max(values) or 1
    for y, v in zip(ys, values):
        ax.text(v + top * 0.02, y, fmt(v), color=_CH_FG, fontsize=9.5, va="center")
    ax.set_xlim(0, top * 1.32)


def _ch_donut(ax, labels: list[str], values: list[float], center: str, sub: str = "") -> None:
    pairs = [(lbl, v) for lbl, v in zip(labels, values) if v > 0]
    if not pairs:
        _ch_empty(ax)
        return
    vals = [v for _, v in pairs]
    total = sum(vals)
    ax.pie(vals, colors=[_CH_COLORS[i % len(_CH_COLORS)] for i in range(len(vals))],
           startangle=90, counterclock=False,
           wedgeprops={"width": 0.36, "edgecolor": _CH_BG, "linewidth": 2})
    ax.set_aspect("equal")
    ax.text(0, 0.08, center, color=_CH_FG, fontsize=15, fontweight="bold", ha="center",
            va="center")
    if sub:
        ax.text(0, -0.16, sub, color=_CH_MUTED, fontsize=9.5, ha="center", va="center")
    leg = ax.legend([f"{lbl} — {round(100 * v / total)}%" for lbl, v in pairs],
                    loc="upper center", bbox_to_anchor=(0.5, 0.02), ncol=2,
                    frameon=False, fontsize=9.5, labelcolor=_CH_FG, handlelength=1.0,
                    columnspacing=1.2)
    for h in leg.legend_handles if hasattr(leg, "legend_handles") else leg.legendHandles:
        h.set_edgecolor("none")


def db_stats_series(d_from: str, d_to: str) -> tuple[list[str], list[int], str]:
    """Выручка оплаченных заказов по часам (один день) / дням / неделям / месяцам."""
    with _connect() as conn:
        rows = conn.execute(
            "SELECT created_at, total - bonus_used AS m FROM orders"
            " WHERE status = 'paid' AND substr(created_at, 1, 10) BETWEEN ? AND ?",
            (d_from, d_to),
        ).fetchall()
    if d_from == d_to:
        by_h: dict[int, int] = {}
        for r in rows:
            h = int(str(r["created_at"])[11:13])
            by_h[h] = by_h.get(h, 0) + int(r["m"] or 0)
        h0 = min([*by_h, 15])
        h1 = max([*by_h, 23])
        hours = list(range(h0, h1 + 1))
        return [f"{h:02d}" for h in hours], [by_h.get(h, 0) for h in hours], "по часам"
    d0 = datetime.strptime(d_from, "%Y-%m-%d").date()
    d1 = datetime.strptime(d_to, "%Y-%m-%d").date()
    n_days = (d1 - d0).days + 1
    by_d: dict[str, int] = {}
    for r in rows:
        k = str(r["created_at"])[:10]
        by_d[k] = by_d.get(k, 0) + int(r["m"] or 0)
    if n_days <= 45:
        days = [d0 + timedelta(days=i) for i in range(n_days)]
        return ([f"{d:%d.%m}" for d in days],
                [by_d.get(f"{d:%Y-%m-%d}", 0) for d in days], "по дням")
    if n_days <= 200:
        start = d0 - timedelta(days=d0.weekday())
        labels, vals = [], []
        while start <= d1:
            end = start + timedelta(days=6)
            labels.append(f"{max(start, d0):%d.%m}")
            vals.append(sum(v for k, v in by_d.items()
                            if f"{max(start, d0):%Y-%m-%d}" <= k <= f"{min(end, d1):%Y-%m-%d}"))
            start = end + timedelta(days=1)
        return labels, vals, "по неделям"
    months: list[str] = []
    y, m = d0.year, d0.month
    while (y, m) <= (d1.year, d1.month):
        months.append(f"{y:04d}-{m:02d}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return ([f"{_STATS_MONTHS[int(k[5:7]) - 1][:3]} {k[2:4]}" for k in months],
            [sum(v for kk, v in by_d.items() if kk[:7] == k) for k in months], "по месяцам")


def stats_chart_png(d_from: str, d_to: str) -> bytes:
    s = db_stats_range(d_from, d_to)
    cogs = db_cogs_range(d_from, d_to)
    labels, vals, by = db_stats_series(d_from, d_to)
    paid_n = int(s["by_status"].get("paid", 0))
    rent = sum(int(r["s"] or 0) for r in s.get("rsv_pay") or [])
    fig = _ch_new_fig()
    period = _ddmm(d_from) if d_from == d_to else f"{_ddmm(d_from)} — {_ddmm(d_to)}"
    _ch_title(fig, "Zig Zag · Статистика", period)
    bottom = _ch_tiles(fig, [
        ("Выручка", _fmt_money(s["revenue"]), "оплачено деньгами"),
        ("Прибыль", _fmt_money(s["revenue"] - cogs), f"себестоимость {_fmt_money(cogs)}"),
        ("Средний чек", _fmt_money(s["revenue"] // paid_n) if paid_n else "—",
         f"оплачено заказов: {paid_n}"),
        ("Заказов", str(s["orders"]),
         f"отменено {s['cancelled']}" if s["cancelled"] else "за период"),
        ("Не оплачено", _fmt_money(s["open_sum"]), f"в работе: {s['open_n']} шт"),
        ("Брони", str(s["rsv_total"]), f"аренда оплачено {_fmt_money(rent)}" if rent
         else "создано за период"),
    ])
    ax = _ch_axes(fig, [0.09, bottom - 0.36, 0.87, 0.27], f"Выручка {by}")
    _ch_bars(ax, labels, vals, _CH_COLORS[0])
    ax2 = _ch_axes(fig, [0.33, 0.06, 0.22, bottom - 0.53], "Топ позиций, шт")
    _ch_hbars(ax2, [_plain(r["name"], 26) for r in s["top"]],
              [int(r["q"]) for r in s["top"]], _CH_COLORS[1], lambda v: f"{int(v)}")
    ax3 = fig.add_axes([0.62, 0.10, 0.34, bottom - 0.55], facecolor=_CH_BG)
    ax3.set_title("Способы оплаты", color=_CH_FG, fontsize=13.5, fontweight="bold",
                  loc="left", pad=10)
    names = {"cash": "Наличные", "card": "Безнал", "transfer": "Перевод"}
    pay: dict[str, int] = {}
    for r in s["pay"]:
        k = names.get(r["m"] or "", "Не указан")
        pay[k] = pay.get(k, 0) + int(r["s"] or 0)
    _ch_donut(ax3, list(pay), list(pay.values()), _money_short(sum(pay.values())),
              "заказы, ₽")
    return _ch_png(fig)


def _card_segment(p: dict | None, today) -> str:
    """Главный сегмент гостя с картой (для кольцевой диаграммы)."""
    if not p or not p["all_days"]:
        return "Без оплат"
    days = p["all_days"]
    if (today - max(days)).days > SEG_SLEEP_DAYS:
        return "Давно не были"
    if p["all_m"] >= SEG_VIP_SPENT or len(days) >= SEG_VIP_VISITS:
        return "Ценные"
    if sum(1 for d in days if (today - d).days < 30) >= SEG_OFTEN_30:
        return "Ходят часто"
    if len(days) >= 3:
        return "Постоянные"
    return "Новые" if len(days) == 1 else "Пришли 2-й раз"


_SEG_ORDER = ("Новые", "Пришли 2-й раз", "Постоянные", "Ходят часто", "Ценные",
              "Давно не были", "Без оплат")


def loyalty_chart_png(o: dict) -> bytes:
    fig = _ch_new_fig()
    title = "всё время" if not o["days"] else f"{o['days']} дн."
    span = f"{o['since']:%d.%m} — {o['today']:%d.%m}" if o["since"] else title
    _ch_title(fig, "Zig Zag · Лояльность", span)
    avg_card = o["rev_card"] // o["n_card"] if o["n_card"] else 0
    n_other = o["n_all"] - o["n_card"]
    avg_other = (o["rev_all"] - o["rev_card"]) // n_other if n_other else 0
    bottom = _ch_tiles(fig, [
        ("Карт всего", str(o["cards"]), f"подтверждено {o['verified']}"),
        ("Новых карт", str(o["new_cards"]), f"за {title}"),
        ("Платили с картой", str(o["active"]),
         f"гостей · {_pct(o['rev_card'], o['rev_all'])}% выручки заказов"),
        ("Вернулись 2+ раз", f"{_pct(o['repeat'], o['active'])}%",
         f"{o['repeat']} из {o['active']}"),
        ("Баллы начислено", _fmt_pts(o["pts_earned"]), f"списано {_fmt_pts(o['pts_spent'])}"),
        ("Баллы на счетах", _fmt_pts(o["outstanding"]), "= ₽ будущих скидок"),
    ])
    # новые карты по времени
    days = [d for d in o.get("card_days", []) if not o["since"] or d >= o["since"]]
    if o["days"] and o["days"] <= 31:
        rng = [o["since"] + timedelta(days=i) for i in range(o["days"])]
        labels = [f"{d:%d.%m}" for d in rng]
        vals = [sum(1 for x in days if x == d) for d in rng]
        by = "по дням"
    elif o["days"]:
        start = o["since"] - timedelta(days=o["since"].weekday())
        labels, vals = [], []
        while start <= o["today"]:
            labels.append(f"{max(start, o['since']):%d.%m}")
            vals.append(sum(1 for x in days if start <= x <= start + timedelta(days=6)))
            start += timedelta(days=7)
        by = "по неделям"
    else:
        months = sorted({f"{d:%Y-%m}" for d in days})[-12:] or [f"{o['today']:%Y-%m}"]
        labels = [f"{_STATS_MONTHS[int(k[5:7]) - 1][:3]} {k[2:4]}" for k in months]
        vals = [sum(1 for x in days if f"{x:%Y-%m}" == k) for k in months]
        by = "по месяцам"
    ax = _ch_axes(fig, [0.08, bottom - 0.33, 0.50, 0.25], f"Новые карты {by}")
    _ch_bars(ax, labels, vals, _CH_COLORS[1], money=False)
    if any(vals):
        ax.yaxis.get_major_locator().set_params(integer=True)
    # средний чек с картой / без
    ax2 = _ch_axes(fig, [0.66, bottom - 0.33, 0.30, 0.25], "Средний чек, ₽")
    if o["n_card"] or n_other:
        _ch_bars(ax2, ["с картой", "без карты"], [avg_card, avg_other], _CH_COLORS[0])
        if len(ax2.patches) > 1:
            ax2.patches[1].set_facecolor(_CH_COLORS[7])
    else:
        _ch_empty(ax2)
    # топ гостей
    ax3 = _ch_axes(fig, [0.25, 0.06, 0.27, bottom - 0.47], "Топ гостей, ₽")
    _ch_hbars(ax3, [_plain(o["names"][u].get("first_name") or "Гость", 16) + f" ({n})"
                    for u, _, n in o["top"]],
              [m for _, m, _ in o["top"]], _CH_COLORS[2], _money_short)
    # сегменты карт
    seg: dict[str, int] = {}
    for u in o["names"]:
        k = _card_segment(o.get("per", {}).get(u), o["today"])
        seg[k] = seg.get(k, 0) + 1
    ax4 = fig.add_axes([0.60, 0.10, 0.36, bottom - 0.50], facecolor=_CH_BG)
    ax4.set_title("Кто держит карты", color=_CH_FG, fontsize=13.5, fontweight="bold",
                  loc="left", pad=10)
    keys = [k for k in _SEG_ORDER if seg.get(k)]
    _ch_donut(ax4, keys, [seg[k] for k in keys], str(o["cards"]), "карт")
    return _ch_png(fig)


async def _render_chart(fn, *args) -> bytes | None:
    """PNG или None (графики выключены / ошибка рисования — не мешаем тексту)."""
    if not _mpl():
        return None
    try:
        async with _CHART_LOCK:
            return await asyncio.to_thread(fn, *args)
    except Exception:
        logger.exception("Не удалось нарисовать график %s", getattr(fn, "__name__", fn))
        return None


def _remember_chart(chat_id: int, text_id: int, photo_id: int) -> None:
    _CHART_MSGS[(chat_id, text_id)] = photo_id
    while len(_CHART_MSGS) > 300:          # не копим бесконечно
        _CHART_MSGS.pop(next(iter(_CHART_MSGS)))


async def _answer_chart_and_text(message: Message, png: bytes | None, caption: str,
                                 text: str, kb) -> None:
    """Сначала картинка, под ней — текст с кнопками (кнопки остаются внизу)."""
    photo_id = None
    if png:
        try:
            pm = await message.answer_photo(BufferedInputFile(png, "zigzag.png"),
                                            caption=caption)
            photo_id = pm.message_id
        except Exception:
            logger.exception("Не удалось отправить график")
    tm = await message.answer(text, reply_markup=kb)
    if photo_id and tm:
        _remember_chart(tm.chat.id, tm.message_id, photo_id)


async def _update_chart(bot: Bot, chat_id: int, text_id: int, png: bytes | None,
                        caption: str) -> None:
    """Заменить картинку над текстом; если её нет (перезапуск) — прислать новую."""
    if not png:
        return
    pid = _CHART_MSGS.get((chat_id, text_id))
    if pid:
        try:
            await bot.edit_message_media(
                media=InputMediaPhoto(media=BufferedInputFile(png, "zigzag.png"),
                                      caption=caption),
                chat_id=chat_id, message_id=pid)
            return
        except TelegramBadRequest as e:
            if "not modified" in str(e).lower():
                return
        except Exception:
            logger.exception("Не удалось обновить график")
    try:
        pm = await bot.send_photo(chat_id, BufferedInputFile(png, "zigzag.png"),
                                  caption=caption)
        _remember_chart(chat_id, text_id, pm.message_id)
    except Exception:
        logger.exception("Не удалось отправить график")


# ------------------------------------------- админ-команды -------------------
_STATS_MONTHS = (
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
)


def _ddmm(date_s: str) -> str:
    return f"{date_s[8:10]}.{date_s[5:7]}"


def db_stats_range(d_from: str, d_to: str) -> dict:
    """Сводка по заказам, оплатам, продажам и броням за произвольный период."""
    with _connect() as conn:
        # выручка — только ОПЛАЧЕННЫЕ заказы и только деньгами (без баллов)
        o = conn.execute(
            "SELECT COUNT(*) AS n,"
            " COALESCE(SUM(CASE WHEN status = 'paid' THEN total - bonus_used END), 0) AS revenue,"
            " COALESCE(SUM(CASE WHEN status = 'paid' THEN bonus_used END), 0) AS bonus_used,"
            " COALESCE(SUM(CASE WHEN status IN ('accepted', 'ready', 'issued')"
            "   THEN total END), 0) AS open_sum,"
            " COALESCE(SUM(CASE WHEN status IN ('accepted', 'ready', 'issued')"
            "   THEN 1 ELSE 0 END), 0) AS open_n,"
            " COALESCE(SUM(CASE WHEN status = 'paid' AND kind = 'promo' THEN 1 ELSE 0 END), 0)"
            "   AS promo_n,"
            " COALESCE(SUM(CASE WHEN status = 'cancelled' THEN 1 ELSE 0 END), 0) AS cancelled,"
            " COALESCE(SUM(CASE WHEN kind = 'pre' THEN 1 ELSE 0 END), 0) AS pre_n,"
            " COALESCE(SUM(CASE WHEN kind = 'pre' AND status = 'cancelled' THEN 1 ELSE 0 END), 0)"
            "   AS pre_cancel"
            " FROM orders WHERE substr(created_at, 1, 10) BETWEEN ? AND ?",
            (d_from, d_to),
        ).fetchone()
        accrued = conn.execute(
            "SELECT COALESCE(SUM(delta), 0) AS a FROM bonus_moves"
            " WHERE kind IN ('accrual', 'offline') AND substr(created_at, 1, 10) BETWEEN ? AND ?",
            (d_from, d_to),
        ).fetchone()["a"]
        by_status = {
            r["status"]: r["n"]
            for r in conn.execute(
                "SELECT status, COUNT(*) AS n FROM orders"
                " WHERE substr(created_at, 1, 10) BETWEEN ? AND ? GROUP BY status",
                (d_from, d_to),
            )
        }
        pay = [
            dict(r)
            for r in conn.execute(
                "SELECT COALESCE(pay_method, '') AS m, COUNT(*) AS n,"
                " COALESCE(SUM(total - bonus_used), 0) AS s FROM orders"
                " WHERE status = 'paid' AND substr(created_at, 1, 10) BETWEEN ? AND ?"
                " GROUP BY pay_method",
                (d_from, d_to),
            )
        ]
        top = [
            dict(r)
            for r in conn.execute(
                "SELECT oi.name AS name, SUM(oi.qty) AS q FROM order_items oi"
                " JOIN orders o ON o.id = oi.order_id"
                " WHERE o.status = 'paid'"
                " AND substr(o.created_at, 1, 10) BETWEEN ? AND ?"
                " GROUP BY oi.name ORDER BY q DESC, oi.name LIMIT 7",
                (d_from, d_to),
            )
        ]
        rsv = {
            r["status"]: r["n"]
            for r in conn.execute(
                "SELECT status, COUNT(*) AS n FROM reservations"
                " WHERE substr(created_at, 1, 10) BETWEEN ? AND ? GROUP BY status",
                (d_from, d_to),
            )
        }
        rsv_pay = [
            dict(r)
            for r in conn.execute(
                "SELECT COALESCE(pay_method, '') AS m, COUNT(*) AS n,"
                " COALESCE(SUM(price), 0) AS s FROM reservations"
                " WHERE status = 'paid'"
                " AND substr(COALESCE(paid_at, created_at), 1, 10) BETWEEN ? AND ?"
                " GROUP BY pay_method",
                (d_from, d_to),
            )
        ]
    return {
        "orders": o["n"],
        "revenue": o["revenue"],
        "bonus_used": o["bonus_used"],
        "bonus_accrued": int(accrued or 0),
        "open_sum": o["open_sum"],
        "open_n": o["open_n"],
        "promo_n": o["promo_n"],
        "cancelled": o["cancelled"],
        "pre_n": o["pre_n"],
        "pre_cancel": o["pre_cancel"],
        "by_status": by_status,
        "pay": pay,
        "top": top,
        "rsv": rsv,
        "rsv_total": sum(rsv.values()),
        "rsv_pay": rsv_pay,
    }


def stats_text(d_from: str, d_to: str) -> str:
    s = db_stats_range(d_from, d_to)
    cogs = db_cogs_range(d_from, d_to)
    period = _ddmm(d_from) if d_from == d_to else f"{_ddmm(d_from)} — {_ddmm(d_to)}"
    lines = [f"📊 <b>Аналитика</b> · {period}", ""]

    lines.append(
        f"📦 Заказов: <b>{s['orders']}</b>"
        + (f" · отменено {s['cancelled']}" if s["cancelled"] else "")
    )
    st_parts = [
        f"{STATUSES[c]}: {n}"
        for c, n in s["by_status"].items()
        if c in STATUSES and n
    ]
    if st_parts:
        lines.append("Статусы: " + " · ".join(st_parts))
    lines += [
        f"💰 Выручка (оплачено деньгами): <b>{_fmt_money(s['revenue'])}</b>",
        f"  себестоимость: {_fmt_money(cogs)}",
        f"  прибыль: <b>{_fmt_money(s['revenue'] - cogs)}</b>",
        f"  вложение на складе (сейчас): {_fmt_money(db_stock_value())}",
    ]
    if s["open_n"]:
        lines.append(f"⏳ Не оплачено (в работе): {s['open_n']} шт · {_fmt_money(s['open_sum'])}")
    if s["promo_n"]:
        lines.append(f"🔥 Оплачено заказов по акции: {s['promo_n']}")
    if s.get("pre_n"):
        lines.append(f"📝 Предзаказов: {s['pre_n']}"
                     + (f" · отменено {s['pre_cancel']}" if s.get("pre_cancel") else ""))
    if s["bonus_used"] or s["bonus_accrued"]:
        lines.append(
            f"💎 Баллы: списано в оплату {_fmt_pts(s['bonus_used'])}"
            f" · начислено {_fmt_pts(s['bonus_accrued'])}"
        )
    lines.append("")

    lines.append("💵 <b>Оплачено по способам:</b>")
    labels = {"cash": "Нал", "card": "Безнал", "transfer": "Перевод"}
    if s["pay"]:
        total_n = total_s = 0
        for row in s["pay"]:
            label = labels.get(row["m"] or "", "Не указан")
            lines.append(f" • {label} — {row['n']} шт · {_fmt_money(row['s'])}")
            total_n += row["n"]
            total_s += row["s"]
        lines.append(f" Итого: <b>{total_n} шт · {_fmt_money(total_s)}</b>")
    else:
        lines.append(" — оплаченных заказов за период нет")
    for row in s.get("rsv_pay") or []:
        rlabel = labels.get(row["m"] or "", "Не указан")
        lines.append(
            f" • Аренда ({rlabel}) — {row['n']} шт · {_fmt_money(row['s'])}"
        )
    lines.append("")

    lines.append("🔥 <b>Топ позиций (оплачено):</b>")
    if s["top"]:
        for i, row in enumerate(s["top"], 1):
            lines.append(f" {i}. {esc(row['name'])} — {row['q']} шт")
    else:
        lines.append(" — продаж за период нет")
    lines.append("")

    rsv_parts = []
    if s["rsv"].get("confirmed"):
        rsv_parts.append(f"✅ {s['rsv']['confirmed']}")
    if s["rsv"].get("new"):
        rsv_parts.append(f"🟡 {s['rsv']['new']}")
    if s["rsv"].get("cancelled"):
        rsv_parts.append(f"❌ {s['rsv']['cancelled']}")
    if s["rsv"].get("no_show"):
        rsv_parts.append(f"🚫 {s['rsv']['no_show']}")
    if s["rsv"].get("paid"):
        rsv_parts.append(f"💳 {s['rsv']['paid']}")
    rsv_line = f"📅 Броней создано: <b>{s['rsv_total']}</b>"
    if rsv_parts:
        rsv_line += " (" + " · ".join(rsv_parts) + ")"
    lines.append(rsv_line)
    return "\n".join(lines)


def kb_stats_period() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [_btn("📅 Сегодня", "sg:day"), _btn("⏪ Вчера", "sg:yday")],
            [_btn("📊 Неделя", "sg:week"), _btn("📆 Месяц", "sg:month")],
            [_btn("🗓 Свой период", "sg:range")],
        ]
    )


def _range_prompt(mode: str, start: str | None = None) -> str:
    if mode == "end" and start:
        return (
            "🗓 <b>Период</b>\n\n"
            f"Начало: <b>{_ddmm(start)}</b> — выберите дату <b>ОКОНЧАНИЯ</b>:"
        )
    return "🗓 <b>Период</b>\n\nВыберите дату <b>НАЧАЛА</b>:"


def kb_month_grid(month: str, hi: str | None = None) -> InlineKeyboardMarkup:
    """Сетка дней месяца; будущие даты недоступны. month = YYYY-MM."""
    year, mon = int(month[:4]), int(month[5:7])
    today_iso = _now_dt().strftime("%Y-%m-%d")
    rows: list[list[InlineKeyboardButton]] = [
        [_btn(d, "noop") for d in ("пн", "вт", "ср", "чт", "пт", "сб", "вс")]
    ]
    for week in calendar.Calendar(firstweekday=0).monthdayscalendar(year, mon):
        row: list[InlineKeyboardButton] = []
        for day in week:
            if not day:
                row.append(_btn("·", "noop"))
                continue
            iso = f"{year:04d}-{mon:02d}-{day:02d}"
            if iso > today_iso:
                row.append(_btn("·", "noop"))
            elif iso == hi:
                row.append(_btn(f"✅{day}", f"sg:d:{iso}"))
            else:
                row.append(_btn(str(day), f"sg:d:{iso}"))
        rows.append(row)

    def _shift(m: str, delta: int) -> str:
        y, mm = int(m[:4]), int(m[5:7]) + delta
        y += (mm - 1) // 12
        mm = (mm - 1) % 12 + 1
        return f"{y:04d}-{mm:02d}"

    prev, nxt = _shift(month, -1), _shift(month, 1)
    rows.append(
        [
            _btn("◀️", f"sg:mv:{prev}") if prev >= "2024-01" else _btn(" ", "noop"),
            _btn(f"{_STATS_MONTHS[mon - 1]} {year}", "noop"),
            _btn("▶️", f"sg:mv:{nxt}") if nxt <= today_iso[:7] else _btn(" ", "noop"),
        ]
    )
    rows.append([_btn("🔙 Назад", "sg:back")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


@router.message(Command("stats"))
async def cmd_stats(message: Message, state: FSMContext):
    """CRM-аналитика: сегодня + кнопки периодов (день/неделя/месяц/свой)."""
    if message.from_user.id not in ADMIN_IDS:
        return
    await state.clear()
    today = _now_dt().strftime("%Y-%m-%d")
    png = await _render_chart(stats_chart_png, today, today)
    await _answer_chart_and_text(message, png, _stats_caption(today, today),
                                 stats_text(today, today), kb_stats_period())


def _stats_caption(d_from: str, d_to: str) -> str:
    period = _ddmm(d_from) if d_from == d_to else f"{_ddmm(d_from)} — {_ddmm(d_to)}"
    return f"📊 Статистика · {period}"


async def _stats_show(cb: CallbackQuery, d_from: str, d_to: str) -> None:
    """Обновить текст сводки и картинку над ним."""
    try:
        await cb.message.edit_text(stats_text(d_from, d_to), reply_markup=kb_stats_period())
    except TelegramBadRequest as e:
        if "not modified" in str(e).lower():
            return  # тот же период ещё раз — всё уже на экране
        raise
    png = await _render_chart(stats_chart_png, d_from, d_to)
    await _update_chart(cb.bot, cb.message.chat.id, cb.message.message_id, png,
                        _stats_caption(d_from, d_to))


@router.callback_query(F.data.startswith("sg:"))
async def cb_stats(cb: CallbackQuery, state: FSMContext):
    if not _is_admin(cb.from_user.id):
        await cb.answer("⛔ Только для администратора", show_alert=True)
        return
    parts = cb.data.split(":")
    today = _now_dt()
    today_iso = today.strftime("%Y-%m-%d")
    try:
        # --- быстрые периоды ---
        if parts[1] in ("day", "yday", "week", "month"):
            await state.clear()
            d_to = today_iso
            if parts[1] == "day":
                d_from = today_iso
            elif parts[1] == "yday":  # календарное вчера (как «Сегодня» — по датам)
                d_from = d_to = (today - timedelta(days=1)).strftime("%Y-%m-%d")
            elif parts[1] == "week":
                d_from = (today - timedelta(days=6)).strftime("%Y-%m-%d")
            else:
                d_from = (today - timedelta(days=29)).strftime("%Y-%m-%d")
            await cb.answer()
            await _stats_show(cb, d_from, d_to)
            return
        # --- свой период: календарь ---
        if parts[1] == "range":
            await cb.answer()
            await state.set_state(StatsFSM.picking)
            await state.update_data(mode="start", ed_msg=cb.message.message_id)
            await cb.message.edit_text(
                _range_prompt("start"), reply_markup=kb_month_grid(today_iso[:7])
            )
            return
        if parts[1] == "back":
            await state.clear()
            await cb.answer()
            await _stats_show(cb, today_iso, today_iso)
            return
        if parts[1] == "mv":  # листание месяцев
            if await state.get_state() != StatsFSM.picking:
                await state.set_state(StatsFSM.picking)
                await state.update_data(mode="start", ed_msg=cb.message.message_id)
            data = await state.get_data()
            mode = data.get("mode") or "start"
            start = data.get("start")
            await cb.answer()
            await cb.message.edit_text(
                _range_prompt(mode, start),
                reply_markup=kb_month_grid(
                    parts[2], hi=start if mode == "end" else None
                ),
            )
            return
        if parts[1] == "d":  # клик по дню
            day = parts[2]
            if await state.get_state() != StatsFSM.picking:
                await state.set_state(StatsFSM.picking)
                await state.update_data(mode="start", ed_msg=cb.message.message_id)
            data = await state.get_data()
            mode = data.get("mode") or "start"
            if mode == "start":
                await state.update_data(mode="end", start=day)
                await cb.answer()
                await cb.message.edit_text(
                    _range_prompt("end", day),
                    reply_markup=kb_month_grid(day[:7], hi=day),
                )
                return
            start = data.get("start") or day
            d_from, d_to = sorted((start, day))
            await state.clear()
            await cb.answer("Готово ✅")
            await _stats_show(cb, d_from, d_to)
            return
    except Exception:
        logger.exception("Ошибка в колбэке статистики: %s", cb.data)
        try:
            await cb.answer("Ошибка", show_alert=True)
        except Exception:
            pass


@router.message(Command("orders"))
async def cmd_orders(message: Message):
    if message.from_user.id not in ADMIN_IDS:
        return
    orders = db_recent_orders(10)
    if not orders:
        await message.answer("Заказов пока нет.")
        return
    lines = ["🗂 Последние заказы:", ""]
    for o in orders:
        status = STATUSES.get(o["status"], o["status"])
        lines.append(
            f"#{o['id']} · {_place_short(o['table_no'])} · {_fmt_money(o['total'])} · "
            f"{status} · {o['created_at'][11:16]}"
        )
    await message.answer("\n".join(lines))


@router.message(Command("timing"))
async def cmd_timing(message: Message):
    """Статистика: сколько времени проходит от принятия заказа до выдачи."""
    if message.from_user.id not in ADMIN_IDS:
        return
    now = _now_dt()
    issued, active = db_timing_today()
    lines = ["⏱ <b>Сроки выдачи</b> (сегодня)", ""]
    if issued:
        durations: list[int] = []
        for row in issued:
            c = datetime.strptime(row["created_at"], "%Y-%m-%d %H:%M:%S")
            i = datetime.strptime(row["issued_at"], "%Y-%m-%d %H:%M:%S")
            mins = max(0, int((i - c).total_seconds() // 60))
            durations.append(mins)
            lines.append(f"#{row['id']} · {_place_short(row['table_no'])} → выдан за <b>{mins} мин</b>")
        avg = round(sum(durations) / len(durations))
        lines += [
            "",
            f"Среднее время выдачи: <b>{avg} мин</b>",
            f"Выдано заказов: <b>{len(durations)}</b>",
        ]
    else:
        lines.append("Сегодня ничего ещё не выдано.")
    if active:
        lines += ["", "🟢 <b>В работе сейчас:</b>"]
        for row in active:
            c = datetime.strptime(row["created_at"], "%Y-%m-%d %H:%M:%S")
            mins = max(0, int((now - c).total_seconds() // 60))
            st = STATUSES.get(row["status"], row["status"])
            lines.append(
                f"#{row['id']} · {_place_short(row['table_no'])} · {st} · в работе {mins} мин"
            )
    await message.answer("\n".join(lines))


# ------------------------------------- /stock и /recipe (админ) -------------
# ------------------------------------------------ /backup -------------------
# Резервная копия всей базы (заказы, брони, карты и баллы гостей, склад,
# рецепты) — файлом админу в личку. Копия делается штатным backup SQLite:
# согласованная, можно в любой момент, даже во время заказов.
BACKUP_MAX_BYTES = 49 * 1024 * 1024  # бот может отправить файл до 50 МБ


def db_backup_bytes() -> tuple[bytes, str, dict]:
    """Копия базы → (байты файла, расширение «.db» или «.db.gz», сводка)."""
    import gzip
    import tempfile

    fd, tmp = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        src = sqlite3.connect(DB_PATH)
        dst = sqlite3.connect(tmp)
        try:
            src.backup(dst)

            def cnt(sql: str) -> int:
                try:
                    return int(dst.execute(sql).fetchone()[0] or 0)
                except sqlite3.Error:
                    return 0

            info = {
                "integrity": str(dst.execute("PRAGMA integrity_check").fetchone()[0]),
                "orders": cnt("SELECT COUNT(*) FROM orders"),
                "rsv": cnt("SELECT COUNT(*) FROM reservations"),
                "cards": cnt("SELECT COUNT(*) FROM guests WHERE phone IS NOT NULL"),
                "points": cnt("SELECT COALESCE(SUM(delta), 0) FROM bonus_moves"),
                "ings": cnt("SELECT COUNT(*) FROM ingredients"),
            }
        finally:
            dst.close()
            src.close()
        data = Path(tmp).read_bytes()
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass
    ext = ".db"
    if len(data) > BACKUP_MAX_BYTES:  # очень большая база — сжимаем
        data = gzip.compress(data)
        ext = ".db.gz"
    return data, ext, info


@router.message(Command("backup"))
async def cmd_backup(message: Message):
    """Резервная копия базы файлом (только админам из ADMIN_IDS, только в личке)."""
    if message.from_user.id not in ADMIN_IDS:
        return
    wait = await message.answer("💾 Готовлю резервную копию базы…")
    try:
        data, ext, info = await asyncio.to_thread(db_backup_bytes)
    except Exception:
        logger.exception("Резервная копия базы не удалась")
        await message.answer("❌ Не удалось сделать копию базы — подробности в логах Bothost.")
        return
    if len(data) > BACKUP_MAX_BYTES:
        await message.answer(
            f"❌ База слишком большая для отправки в Telegram "
            f"({len(data) // (1024 * 1024)} МБ даже в сжатом виде). "
            "Скачайте файл orders.db с сервера вручную."
        )
        return
    now = _now_dt()
    ok = info["integrity"] == "ok"
    unpack = " (сначала распакуйте .gz)" if ext.endswith(".gz") else ""
    caption = (
        f"💾 <b>Резервная копия базы</b> · {now:%d.%m.%Y %H:%M}\n"
        f"Заказов: {info['orders']} · броней: {info['rsv']} · карт гостей: {info['cards']}\n"
        f"Баллов на картах: {info['points']} · позиций склада: {info['ings']}\n"
        f"Проверка целостности: {'✅ в порядке' if ok else '⚠️ ' + esc(info['integrity'][:80])}\n\n"
        "🔒 В файле телефоны гостей — храните у себя, не пересылайте посторонним.\n"
        "♻️ Восстановление: остановите бота → положите этот файл в папку данных "
        f"под именем <code>orders.db</code>{unpack} (сейчас база: "
        f"<code>{esc(DB_PATH)}</code>) → запустите бота."
    )
    name = f"zigzag_backup_{now:%Y-%m-%d_%H%M}{ext}"
    try:
        await message.answer_document(BufferedInputFile(data, name), caption=caption[:1024])
    except Exception:
        logger.exception("Не удалось отправить резервную копию")
        await message.answer("❌ Копия сделана, но Telegram не принял файл — попробуйте ещё раз.")
        return
    logger.info("Резервная копия базы (%d байт) отправлена админу %s",
                len(data), message.from_user.id)
    try:
        await message.bot.delete_message(message.chat.id, wait.message_id)
    except Exception:
        pass


class StockFSM(StatesGroup):
    qty = State()      # приход/расход: количество
    money = State()    # приход: сумма закупки (расчёт средней с/с)
    cost = State()     # себестоимость за единицу
    minq = State()     # минимальный остаток
    new_name = State() # создание: название
    new_qty = State()  # создание: начальный остаток
    new_cost = State() # создание: себестоимость
    new_minq = State() # создание: мин. остаток
    rename = State()   # новое название позиции


class RecipeFSM(StatesGroup):
    qty = State()      # сколько ингредиента на 1 порцию


class StatsFSM(StatesGroup):
    picking = State()  # выбор периода в /stats: mode=start|end, start, ed_msg


async def _admin_msg_guard(message: Message, state: FSMContext) -> bool:
    """Ввод в шагах склада/рецептов — только админам; чужой — сбрасываем."""
    if _is_admin(message.from_user.id if message.from_user else None):
        return True
    await state.clear()
    return False


def _parse_num(t: str) -> float | None:
    """«500» / «10,5» → число; мусор → None."""
    try:
        v = float(t.replace(",", ".").replace(" ", ""))
    except (ValueError, AttributeError):
        return None
    return v if v >= 0 else None


def _menu_item(item_id: str) -> dict | None:
    for c in MENU["categories"]:
        for it in c.get("items", []):
            if it.get("id") == item_id:
                return it
    return None


def _cancel_kb(cb_data: str) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[[_btn("✖ Отмена", cb_data)]])


async def _edit_or_send(message: Message, data: dict, text: str, kb) -> None:
    """Правит сохранённое сообщение (data['ed_msg']); если не вышло — шлёт новое."""
    mid = data.get("ed_msg")
    if mid:
        try:
            await message.bot.edit_message_text(
                text=text, chat_id=message.chat.id, message_id=mid, reply_markup=kb
            )
            return
        except Exception:
            logger.info("Не удалось отредактировать сообщение %s, шлём новое", mid)
    try:
        await message.answer(text, reply_markup=kb)
    except Exception:
        logger.exception("Не удалось отправить сообщение")


SK_PAGE = 30  # позиций на странице «🗂 По позициям» (лимит Telegram — 100 кнопок)


def _sk_page(ings: list[dict], page: int) -> tuple[list[dict], int, int]:
    """Срез позиций для страницы → (позиции, номер страницы, всего страниц)."""
    pages = max(1, -(-len(ings) // SK_PAGE))
    page = min(max(0, int(page)), pages - 1)
    return ings[page * SK_PAGE:(page + 1) * SK_PAGE], page, pages


def _sk_page_of(ing_id: int) -> int:
    """На какой странице «По позициям» находится позиция (для «🔙 Назад»)."""
    ids = [int(i["id"]) for i in db_list_ingredients()]
    return ids.index(int(ing_id)) // SK_PAGE if int(ing_id) in ids else 0


def stock_list_text(page: int = 0) -> str:
    ings = db_list_ingredients()
    head = (
        "📦 <b>Склад</b>\n"
        f"Вложение на складе: <b>{_fmt_money(db_stock_value())}</b>"
    )
    if not ings:
        return head + "\n\nПока пусто — добавьте первую позицию кнопкой ниже."
    chunk, page, pages = _sk_page(ings, page)
    lines = [head]
    if pages > 1:
        first = page * SK_PAGE + 1
        lines.append(f"Позиции {first}–{first + len(chunk) - 1} из {len(ings)} · "
                     f"стр. {page + 1}/{pages}")
    lines.append("")
    for ing in chunk:
        mark = ""
        if float(ing["qty"]) <= 0:
            mark = " ⛔"
        elif float(ing["min_qty"]) > 0 and float(ing["qty"]) <= float(ing["min_qty"]):
            mark = " ⚠️"
        lines.append(
            f"• {esc(ing['name'])} — {_fmt_qty(ing['qty'])} {esc(ing['unit'])}{mark}"
        )
    return _fit_lines("\n".join(lines))


def _fit_lines(text: str, limit: int = 3900) -> str:
    """Укоротить текст до лимита Telegram по границе строки (не рвёт HTML-теги)."""
    if len(text) <= limit:
        return text
    return text[:limit - 40].rsplit("\n", 1)[0] + "\n…(список сокращён)"


def stock_detail_text(ing: dict) -> str:
    mark = ""
    if float(ing["qty"]) <= 0:
        mark = " ⛔"
    elif float(ing["min_qty"]) > 0 and float(ing["qty"]) <= float(ing["min_qty"]):
        mark = " ⚠️"
    min_note = (
        f"{_fmt_qty(ing['min_qty'])} {esc(ing['unit'])}"
        if float(ing["min_qty"]) > 0
        else "не контролируется"
    )
    lines = [
        f"📦 <b>{esc(ing['name'])}</b>",
        "",
        f"Остаток: <b>{_fmt_qty(ing['qty'])} {esc(ing['unit'])}</b>{mark}",
        f"Себестоимость: {_fmt_price(ing['cost'])} за {esc(ing['unit'])}",
        f"Вложено: <b>{_fmt_money(round(float(ing['qty']) * float(ing['cost'])))}</b>",
        f"Мин. остаток: {min_note}",
    ]
    moves = db_recent_moves(ing["id"])
    if moves:
        lines += ["", "<b>Последние движения:</b>"]
        for m in moves:
            d = float(m["delta"])
            sign = "+" if d >= 0 else "−"
            when = (m.get("created_at") or "")[5:16]
            ref = f" · {m['ref']}" if m.get("ref") else ""
            lines.append(
                f"  {sign}{_fmt_qty(abs(d))} {esc(ing['unit'])} · {m['reason']}"
                f"{ref} · {when}"
            )
    return "\n".join(lines)


def kb_stock_list(page: int = 0) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    chunk, page, pages = _sk_page(db_list_ingredients(), page)
    for ing in chunk:
        rows.append(
            [
                _btn(
                    f"{ing['name']} · {_fmt_qty(ing['qty'])} {ing['unit']}",
                    f"sk:{ing['id']}",
                )
            ]
        )
    if pages > 1:
        nav = []
        if page > 0:
            nav.append(_btn("◀️", f"sk:pg:{page - 1}"))
        nav.append(_btn(f"{page + 1}/{pages}", "noop"))
        if page < pages - 1:
            nav.append(_btn("▶️", f"sk:pg:{page + 1}"))
        rows.append(nav)
    rows.append([_btn("➕ Новая позиция", "skn"), _btn("📋 Таблица", "sx:table")])
    rows.append([_btn("📜 История", "sh:0:a:w:0"), _btn("✖ Закрыть", "sk:close")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_stock_detail(ing_id: int) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                _btn("➕ Приход", f"sk:{ing_id}:in"),
                _btn("➖ Расход", f"sk:{ing_id}:out"),
            ],
            [
                _btn("✏️ Себестоимость", f"sk:{ing_id}:cost"),
                _btn("📉 Мин.остаток", f"sk:{ing_id}:min"),
            ],
            [_btn("📏 Единица", f"sk:{ing_id}:unit"), _btn("✏️ Название", f"sk:{ing_id}:name")],
            [_btn("📜 История", f"sh:{ing_id}:a:m:0")],
            [_btn("🗑 Удалить позицию", f"sk:{ing_id}:rm")],
            [_btn("🔙 Назад", f"sk:pg:{_sk_page_of(ing_id)}")],
        ]
    )


@router.message(Command("stock"))
async def cmd_stock(message: Message, state: FSMContext):
    """Склад: остатки, себестоимость, приход/расход."""
    if message.from_user.id not in ADMIN_IDS:
        return
    await state.clear()
    await _send_stock_table(message)


@router.callback_query(F.data.startswith("sk:"))
async def cb_stock(cb: CallbackQuery, state: FSMContext):
    if not _is_admin(cb.from_user.id):
        await cb.answer("⛔ Только для администратора", show_alert=True)
        return
    parts = cb.data.split(":")
    try:
        # --- приход введён, но сумму закупки не меняем ---
        if len(parts) == 2 and parts[1] == "costskip":
            data = await state.get_data()
            ing = db_get_ingredient(data.get("sk_ing") or 0)
            pending = float(data.get("sk_pending_qty") or 0)
            await state.clear()
            if not ing or pending <= 0:
                await cb.answer("Нечего приходовать", show_alert=True)
                return
            updated = stock_move(ing["id"], +pending, "приход")
            await cb.answer("Принято ✅ (с/с не менялась)")
            if updated:
                await _stock_alerts(cb.bot, [updated])
            try:
                await cb.message.edit_text(
                    f"✅ Приход: {_fmt_qty(pending)} {esc(updated['unit'])} "
                    f"(без изменения с/с)\n\n" + stock_detail_text(updated),
                    reply_markup=kb_stock_detail(ing["id"]),
                )
            except TelegramBadRequest:
                pass
            return
        # --- «🗂 По позициям»: страница списка ---
        if len(parts) == 3 and parts[1] == "pg":
            await state.clear()
            await cb.answer()
            page = int(parts[2]) if parts[2].isdigit() and len(parts[2]) < 5 else 0
            try:
                await cb.message.edit_text(stock_list_text(page), reply_markup=kb_stock_list(page))
            except TelegramBadRequest:
                pass
            return
        # --- служебные ---
        if len(parts) == 2 and parts[1] in ("list", "close", "cancel"):
            await state.clear()
            await cb.answer("Отменили" if parts[1] == "cancel" else "Готово")
            text = stock_list_text() if parts[1] != "close" else "✅ Готово"
            kb = kb_stock_list() if parts[1] != "close" else None
            try:
                await cb.message.edit_text(text, reply_markup=kb)
            except TelegramBadRequest:
                pass
            return
        # --- смена единицы измерения: показать выбор ---
        if len(parts) == 3 and parts[2] == "unit":
            ing = db_get_ingredient(int(parts[1]))
            if not ing:
                await cb.answer("Позиция не найдена", show_alert=True)
                return
            await cb.answer()
            try:
                await cb.message.edit_text(
                    f"📏 <b>{esc(ing['name'])}</b>\n\n"
                    "Единица измерения для остатков и рецептов.\n"
                    "Число остатка не изменится — после смены поправьте его, "
                    "если нужно.",
                    reply_markup=InlineKeyboardMarkup(
                        inline_keyboard=[
                            [
                                _btn(u, f"sk:{ing['id']}:u:{u}")
                                for u in ("шт", "г", "мл")
                            ],
                            [_btn("🔙 Назад", f"sk:{ing['id']}")],
                        ]
                    ),
                )
            except TelegramBadRequest:
                pass
            return
        # --- карточка позиции ---
        if len(parts) == 2:
            ing = db_get_ingredient(int(parts[1]))
            if not ing:
                await cb.answer("Позиция не найдена", show_alert=True)
                return
            await cb.answer()
            try:
                await cb.message.edit_text(
                    stock_detail_text(ing), reply_markup=kb_stock_detail(ing["id"])
                )
            except TelegramBadRequest:
                pass
            return
        # --- приход/расход/себестоимость/минимум/удаление ---
        if len(parts) == 3:
            ing = db_get_ingredient(int(parts[1]))
            act = parts[2]
            if not ing:
                await cb.answer("Позиция не найдена", show_alert=True)
                return
            if act == "rm":  # экран подтверждения удаления
                await cb.answer()
                await cb.message.edit_text(
                    f"🗑 <b>Удалить позицию?</b>\n\n"
                    f"{esc(ing['name'])} — остаток "
                    f"{_fmt_qty(ing['qty'])} {esc(ing['unit'])}\n\n"
                    "Позиция исчезнет из склада и из всех рецептов. "
                    "Действие необратимо.",
                    reply_markup=InlineKeyboardMarkup(
                        inline_keyboard=[
                            [_btn("🗑 Да, удалить", f"sk:{ing['id']}:yes")],
                            [_btn("🔙 Назад", f"sk:{ing['id']}")],
                        ]
                    ),
                )
                return
            if act == "yes":  # подтверждено — удаляем
                db_delete_ingredient(ing["id"])
                await cb.answer(f"«{ing['name']}» удалена")
                try:
                    await cb.message.edit_text(
                        stock_list_text(), reply_markup=kb_stock_list()
                    )
                except TelegramBadRequest:
                    pass
                return
            prompts = {
                "in": (
                    f"➕ Приход: <b>{esc(ing['name'])}</b>\n\n"
                    f"Сколько поступило ({esc(ing['unit'])})? "
                    "Отправьте числом, например 500"
                ),
                "out": (
                    f"➖ Расход: <b>{esc(ing['name'])}</b>\n\n"
                    f"Сколько списать ({esc(ing['unit'])})? "
                    f"На складе сейчас: {_fmt_qty(ing['qty'])}"
                ),
                "cost": (
                    f"✏️ Себестоимость: <b>{esc(ing['name'])}</b>\n\n"
                    f"Цена за 1 {esc(ing['unit'])} в ₽ (можно с копейками, например 6,5)"
                ),
                "min": (
                    f"📉 Мин. остаток: <b>{esc(ing['name'])}</b>\n\n"
                    f"Порог ({esc(ing['unit'])}), 0 — не контролировать"
                ),
                "name": (
                    f"✏️ Название: <b>{esc(ing['name'])}</b>\n\n"
                    "Отправьте новое название (до 60 символов). Остаток, рецепты и "
                    "история сохранятся. Название в меню гостя меняется в menu.json."
                ),
            }
            if act not in prompts:
                await cb.answer()
                return
            states = {
                "in": StockFSM.qty,
                "out": StockFSM.qty,
                "cost": StockFSM.cost,
                "min": StockFSM.minq,
                "name": StockFSM.rename,
            }
            await cb.answer()
            await state.set_state(states[act])
            await state.update_data(
                sk_ing=ing["id"], sk_action=act, ed_msg=cb.message.message_id
            )
            await cb.message.edit_text(
                prompts[act], reply_markup=_cancel_kb("sk:cancel")
            )
            return
        # --- применение новой единицы ---
        if len(parts) == 4 and parts[2] == "u":
            ing = db_get_ingredient(int(parts[1]))
            if not ing:
                await cb.answer("Позиция не найдена", show_alert=True)
                return
            if parts[3] not in ("шт", "г", "мл"):
                await cb.answer("Неизвестная единица", show_alert=True)
                return
            db_set_stock_unit(ing["id"], parts[3])
            updated = db_get_ingredient(ing["id"])
            await cb.answer(f"Единица: {parts[3]}")
            try:
                await cb.message.edit_text(
                    stock_detail_text(updated),
                    reply_markup=kb_stock_detail(updated["id"]),
                )
            except TelegramBadRequest:
                pass
            return
    except Exception:
        logger.exception("Ошибка в колбэке склада: %s", cb.data)
        try:
            await cb.answer("Ошибка", show_alert=True)
        except Exception:
            pass


@router.callback_query(F.data.startswith("skn"))
async def cb_stock_new(cb: CallbackQuery, state: FSMContext):
    if not _is_admin(cb.from_user.id):
        await cb.answer("⛔ Только для администратора", show_alert=True)
        return
    parts = cb.data.split(":")
    try:
        if cb.data == "skn":
            await cb.answer()
            await state.set_state(StockFSM.new_name)
            await state.update_data(ed_msg=cb.message.message_id)
            await cb.message.edit_text(
                "➕ <b>Новая позиция склада</b>\n\n"
                "Введите название (например «Табак Musthave»):",
                reply_markup=_cancel_kb("skn:cancel"),
            )
            return
        if cb.data == "skn:cancel":
            await state.clear()
            await cb.answer("Отменили")
            try:
                await cb.message.edit_text(
                    stock_list_text(), reply_markup=kb_stock_list()
                )
            except TelegramBadRequest:
                pass
            return
        if len(parts) == 3 and parts[1] == "unit":
            data = await state.get_data()
            if "sk_new_name" not in data:
                await cb.answer("Начните заново через /stock", show_alert=True)
                return
            await cb.answer()
            await state.update_data(unit=parts[2])
            await state.set_state(StockFSM.new_qty)
            await cb.message.edit_text(
                f"«{esc(data['sk_new_name'])}», единица: <b>{esc(parts[2])}</b>\n\n"
                "Введите начальный остаток (число, можно 0):",
                reply_markup=_cancel_kb("skn:cancel"),
            )
            return
    except Exception:
        logger.exception("Ошибка в колбэке новой позиции: %s", cb.data)
        try:
            await cb.answer("Ошибка", show_alert=True)
        except Exception:
            pass


# --- ввод чисел для склада ---
@router.message(StockFSM.qty, F.text)
async def stock_qty_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    data = await state.get_data()
    ing = db_get_ingredient(data.get("sk_ing") or 0)
    v = _parse_num((message.text or "").strip())
    if ing is None:
        await state.clear()
        await message.answer("Позиция не найдена. Откройте /stock заново.")
        return
    if v is None or v <= 0 or v > 1_000_000:
        await message.answer(
            "Введите число больше 0 (например 500), либо ✖ Отмена кнопкой выше."
        )
        return
    action = data.get("sk_action")
    if action != "out":
        # приход: сперва сумма закупки — по ней посчитаем среднюю себестоимость
        await state.set_state(StockFSM.money)
        await state.update_data(sk_pending_qty=v)
        await _edit_or_send(
            message,
            data,
            f"➕ Принято: {_fmt_qty(v)} {esc(ing['unit'])} «{esc(ing['name'])}».\n\n"
            "Сколько заплатили за это поступление, ₽ — сумма за всё сразу?\n"
            "Например: 10 упаковок × 650 ₽ = 6500\n"
            "(себестоимость посчитается автоматически)",
            InlineKeyboardMarkup(
                inline_keyboard=[
                    [_btn("⏭ Не менять с/с", "sk:costskip")],
                    [_btn("✖ Отмена", "sk:cancel")],
                ]
            ),
        )
        return
    if v > float(ing["qty"]):
        await message.answer(
            f"Недостаточно: на складе {_fmt_qty(ing['qty'])} {ing['unit']}. "
            "Введите число не больше этого."
        )
        return
    updated = stock_move(ing["id"], -v, "расход")
    if updated is None:
        await state.clear()
        await message.answer("Позиция не найдена.")
        return
    await state.clear()
    await _stock_alerts(message.bot, [updated])
    await _edit_or_send(
        message,
        data,
        f"✅ Списано: {_fmt_qty(v)} {esc(updated['unit'])} → "
        f"остаток {_fmt_qty(updated['qty'])} {esc(updated['unit'])}\n\n"
        + stock_detail_text(updated),
        kb_stock_detail(ing["id"]),
    )


@router.message(StockFSM.money, F.text)
async def stock_money_input(message: Message, state: FSMContext):
    """Сумма закупки прихода → средневзвешенная себестоимость."""
    if not await _admin_msg_guard(message, state):
        return
    data = await state.get_data()
    ing = db_get_ingredient(data.get("sk_ing") or 0)
    pending = float(data.get("sk_pending_qty") or 0)
    v = _parse_num((message.text or "").strip())
    if ing is None:
        await state.clear()
        await message.answer("Позиция не найдена. Откройте /stock заново.")
        return
    if v is None or v < 0 or v > 100_000_000 or pending <= 0:
        await message.answer(
            "Введите сумму в ₽ числом (например 6500), либо ✖ Отмена."
        )
        return
    # средневзвешенная: (старый остаток × старая с/с + сумма закупки) / новый остаток
    old_q, old_c = float(ing["qty"]), float(ing["cost"])
    if old_q + pending > 0:
        new_cost = (old_q * old_c + v) / (old_q + pending)
    else:
        new_cost = old_c
    db_set_cost(ing["id"], new_cost)
    updated = stock_move(ing["id"], +pending, "приход", ref=_fmt_price(v))
    await state.clear()
    if updated is None:
        await message.answer("Позиция не найдена.")
        return
    await _stock_alerts(message.bot, [updated])
    await _edit_or_send(
        message,
        data,
        f"✅ Приход: {_fmt_qty(pending)} {esc(updated['unit'])} за {_fmt_price(v)}\n"
        f"Себестоимость: {_fmt_price(new_cost)} за {esc(updated['unit'])}\n\n"
        + stock_detail_text(updated),
        kb_stock_detail(ing["id"]),
    )


@router.message(StockFSM.cost, F.text)
async def stock_cost_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    data = await state.get_data()
    ing = db_get_ingredient(data.get("sk_ing") or 0)
    v = _parse_num((message.text or "").strip())
    if ing is None:
        await state.clear()
        await message.answer("Позиция не найдена. Откройте /stock заново.")
        return
    if v is None or v < 0 or v > 10_000_000:
        await message.answer("Введите цену в ₽ числом (можно с копейками), либо ✖ Отмена.")
        return
    db_set_cost(ing["id"], v)
    await state.clear()
    updated = db_get_ingredient(ing["id"])
    await _edit_or_send(
        message,
        data,
        f"✅ Себестоимость обновлена\n\n{stock_detail_text(updated)}",
        kb_stock_detail(ing["id"]),
    )


@router.message(StockFSM.minq, F.text)
async def stock_min_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    data = await state.get_data()
    ing = db_get_ingredient(data.get("sk_ing") or 0)
    v = _parse_num((message.text or "").strip())
    if ing is None:
        await state.clear()
        await message.answer("Позиция не найдена. Откройте /stock заново.")
        return
    if v is None or v < 0 or v > 10_000_000:
        await message.answer("Введите число (можно 0), либо ✖ Отмена.")
        return
    db_set_min_qty(ing["id"], v)
    await state.clear()
    updated = db_get_ingredient(ing["id"])
    await _edit_or_send(
        message,
        data,
        f"✅ Минимальный остаток обновлён\n\n{stock_detail_text(updated)}",
        kb_stock_detail(ing["id"]),
    )


# --- создание новой позиции ---
@router.message(StockFSM.new_name, F.text)
async def stock_new_name_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    t = (message.text or "").strip()
    if not t or len(t) > 60:
        await message.answer("Введите название до 60 символов сообщением ✍️")
        return
    if any(i["name"].lower() == t.lower() for i in db_list_ingredients()):
        await message.answer("Позиция с таким названием уже есть. Введите другое.")
        return
    await state.update_data(sk_new_name=t)
    data = await state.get_data()
    await _edit_or_send(
        message,
        data,
        f"«{esc(t)}» — единица измерения:",
        InlineKeyboardMarkup(
            inline_keyboard=[
                [
                    _btn("шт", "skn:unit:шт"),
                    _btn("г", "skn:unit:г"),
                    _btn("мл", "skn:unit:мл"),
                ],
                [_btn("✖ Отмена", "skn:cancel")],
            ]
        ),
    )


@router.message(StockFSM.rename, F.text)
async def stock_rename_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    data = await state.get_data()
    ing = db_get_ingredient(data.get("sk_ing") or 0)
    if ing is None:
        await state.clear()
        await message.answer("Позиция не найдена. Откройте /stock заново.")
        return
    t = " ".join((message.text or "").split())
    if not t or len(t) > 60:
        await message.answer("Введите название до 60 символов сообщением ✍️")
        return
    if any(i["name"].lower() == t.lower() and i["id"] != ing["id"] for i in db_list_ingredients()):
        await message.answer("Позиция с таким названием уже есть. Введите другое.")
        return
    if not db_rename_ingredient(ing["id"], t):
        await message.answer("Позиция с таким названием уже есть. Введите другое.")
        return
    await state.clear()
    updated = db_get_ingredient(ing["id"])
    await _edit_or_send(message, data,
                        f"✅ «{esc(ing['name'])}» → «{esc(t)}»\n\n" + stock_detail_text(updated),
                        kb_stock_detail(ing["id"]))


@router.message(StockFSM.new_qty, F.text)
async def stock_new_qty_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    v = _parse_num((message.text or "").strip())
    if v is None or v < 0 or v > 10_000_000:
        await message.answer("Введите число (можно 0), либо ✖ Отмена.")
        return
    await state.update_data(sk_new_qty=v)
    await state.set_state(StockFSM.new_cost)
    data = await state.get_data()
    await _edit_or_send(
        message,
        data,
        "Себестоимость за 1 единицу, ₽ (можно с копейками, например 6,5):",
        _cancel_kb("skn:cancel"),
    )


@router.message(StockFSM.new_cost, F.text)
async def stock_new_cost_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    v = _parse_num((message.text or "").strip())
    if v is None or v < 0 or v > 10_000_000:
        await message.answer("Введите цену в ₽ числом (можно 0), либо ✖ Отмена.")
        return
    await state.update_data(sk_new_cost=float(v))
    await state.set_state(StockFSM.new_minq)
    data = await state.get_data()
    await _edit_or_send(
        message,
        data,
        "Минимальный остаток для контроля (число, 0 — не контролировать):",
        _cancel_kb("skn:cancel"),
    )


@router.message(StockFSM.new_minq, F.text)
async def stock_new_minq_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    data = await state.get_data()
    v = _parse_num((message.text or "").strip())
    if v is None or v < 0 or v > 10_000_000:
        await message.answer("Введите число (можно 0), либо ✖ Отмена.")
        return
    try:
        ing_id = db_add_ingredient(
            data["sk_new_name"],
            data.get("unit", "шт"),
            data.get("sk_new_qty", 0),
            data.get("sk_new_cost", 0),
            v,
        )
    except Exception:
        logger.exception("Не удалось создать позицию склада")
        await state.clear()
        await message.answer("Не удалось создать позицию. Попробуйте через /stock.")
        return
    await state.clear()
    ing = db_get_ingredient(ing_id)
    await _edit_or_send(
        message,
        data,
        f"✅ Позиция создана\n\n{stock_detail_text(ing)}",
        kb_stock_detail(ing_id),
    )


# ------------------------------- /stock: таблица ------------------------------
# Склад одним экраном + массовые правки: текстом (строка на позицию) или
# Excel-файлом (скачать → заполнить жёлтые колонки → прислать обратно).
# Оба пути: разбор → предпросмотр «было → стало» → ✅ Применить (атомарно,
# с записью в журнал движений) → алерты «мало/закончилось».
class StockBulkFSM(StatesGroup):
    text = State()     # ждём сообщение с правками
    confirm = State()  # показан предпросмотр: sx_ops, sx_src


SX_MARK = "zigzag-stock-v1"
_SX_NUM = r"\d+(?:[.,]\d+)?"


def _q(v: float) -> str:
    """Количество без «1e+06» и хвостовых нулей: 10.500 → «10.5»."""
    t = f"{float(v):.3f}".rstrip("0").rstrip(".")
    return "0" if t in ("-0", "") else t


def _stock_flag(ing: dict) -> str:
    q, mn = float(ing["qty"]), float(ing["min_qty"])
    if q <= 0:
        return " ⛔"
    if mn > 0 and q <= mn:
        return " ⚠️"
    return ""


def stock_table_texts() -> list[str]:
    """Таблица склада моноширинным шрифтом; длинная — несколькими сообщениями."""
    ings = db_list_ingredients()
    out_n = sum(1 for i in ings if float(i["qty"]) <= 0)
    low_n = sum(1 for i in ings if float(i["qty"]) > 0 and float(i["min_qty"]) > 0
                and float(i["qty"]) <= float(i["min_qty"]))
    head = [f"📦 <b>Склад</b> · вложено <b>{_fmt_money(db_stock_value())}</b>"]
    if out_n or low_n:
        head.append(f"⛔ закончилось: {out_n} · ⚠️ мало: {low_n}")
    if not ings:
        return ["\n".join(head) + "\n\nПока пусто — добавьте первую позицию кнопкой ниже."]
    # ~36 символов в строке — помещается на экран телефона без переноса
    rows = [f"{ing['id']:>3} {_plain(ing['name'], 19):<19} "
            f"{_q(ing['qty']) + _plain(ing['unit'], 3):>7} "
            f"{_q(ing['min_qty']) if float(ing['min_qty']) > 0 else '—':>4}"
            f"{_stock_flag(ing)}" for ing in ings]
    hdr = f"{'№':>3} {'Позиция':<19} {'Ост.':>7} {'Мин':>4}"
    foot = ("\n\n№ — номер для правок текстом. С/с и суммы — в Excel или "
            "«🗂 По позициям».")
    chunks: list[list[str]] = [[]]
    size = 0
    for r in rows:
        if size + len(r) > 3300 and chunks[-1]:
            chunks.append([])
            size = 0
        chunks[-1].append(esc(r))
        size += len(r) + 1
    texts = []
    for i, ch in enumerate(chunks):
        t = ("\n".join(head) + "\n\n" if i == 0 else f"📦 Склад · продолжение {i + 1}\n\n")
        t += f"<pre>{esc(hdr)}\n" + "\n".join(ch) + "</pre>"
        if i == len(chunks) - 1:
            t += foot
        texts.append(t)
    return texts


def kb_stock_table() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [_btn("📥 Скачать Excel", "sx:xlsx"), _btn("✏️ Правки текстом", "sx:bulk")],
        [_btn("🗂 По позициям", "sk:list"), _btn("➕ Новая позиция", "skn")],
        [_btn("📜 История прихода и списаний", "sh:0:a:w:0")],
        [_btn("✖ Закрыть", "sk:close")],
    ])


def _sx_help(example_id: int | None = None) -> str:
    n = example_id or 12
    return (
        "✏️ <b>Массовые правки склада</b>\n\n"
        "Пришлите одним сообщением, по строке на изменение. "
        "Первое число — № из таблицы /stock:\n"
        f"<code>{n} +5</code> — приход 5\n"
        f"<code>{n} +5 3250</code> — приход 5 на сумму 3250 ₽ (с/с пересчитается)\n"
        f"<code>{n} -2</code> — расход (списание) 2\n"
        f"<code>{n} =40</code> — остаток стал 40 (инвентаризация)\n"
        f"<code>{n} =40 650</code> — остаток 40 и с/с 650 ₽ за ед.\n"
        f"<code>{n} сс 650</code> — себестоимость 650 ₽ за ед.\n"
        f"<code>{n} мин 10</code> — минимальный остаток 10 (0 — не следить)\n\n"
        "Числа без пробелов внутри, дробные — через точку или запятую. "
        "Сначала покажу, что изменится, и только после ✅ применю."
    )


def _sx_num(t: str) -> float:
    return float(t.replace(",", "."))


def sx_parse_text(text: str) -> tuple[list[dict], list[str]]:
    """Строки «№ операция» → (операции, ошибки)."""
    ops: list[dict] = []
    errs: list[str] = []
    lines = [ln.strip() for ln in (text or "").splitlines() if ln.strip()]
    if len(lines) > SX_MAX_OPS:
        return [], [f"Слишком много строк ({len(lines)}), максимум {SX_MAX_OPS} за раз."]
    N = _SX_NUM
    for k, ln in enumerate(lines, 1):
        t = ln.lower().replace("−", "-").replace("–", "-").replace("—", "-")
        t = re.sub(r"\s*(₽|руб\.?|р\.?)\s*$", "", t)
        m = re.fullmatch(r"#?(\d{1,7})[.):]?\s+(.+)", t)
        if not m:
            errs.append(f"стр. {k} «{esc(ln[:40])}»: начните с № позиции")
            continue
        iid, rest = int(m.group(1)), m.group(2).strip()
        op = None
        if mm := re.fullmatch(rf"\+\s*({N})(?:\s+(?:за\s+)?({N}))?", rest):
            op = {"op": "in", "v": _sx_num(mm.group(1)),
                  "sum": _sx_num(mm.group(2)) if mm.group(2) else None}
        elif mm := re.fullmatch(rf"-\s*({N})", rest):
            op = {"op": "out", "v": _sx_num(mm.group(1))}
        elif mm := re.fullmatch(rf"=\s*({N})(?:\s+(?:сс\s*|с/с\s*)?({N}))?", rest):
            op = {"op": "set", "v": _sx_num(mm.group(1)),
                  "cost": _sx_num(mm.group(2)) if mm.group(2) else None}
        elif mm := re.fullmatch(rf"(?:сс|с/с|cost)\s*=?\s*({N})", rest):
            op = {"op": "cost", "v": _sx_num(mm.group(1))}
        elif mm := re.fullmatch(rf"(?:мин|min)\.?\s*=?\s*({N})", rest):
            op = {"op": "min", "v": _sx_num(mm.group(1))}
        if op is None:
            errs.append(f"стр. {k} «{esc(ln[:40])}»: не понял операцию")
            continue
        op["id"] = iid
        op["src"] = f"стр. {k}"
        ops.append(op)
    return ops, errs


def _sx_check_bounds(op: dict) -> str | None:
    v = op["v"]
    if op["op"] in ("in", "out") and v <= 0:
        return "количество должно быть больше 0"
    if op["op"] in ("in", "out", "set", "min") and v > SX_MAX_QTY:
        return "слишком большое количество"
    if op["op"] == "cost" and v > SX_MAX_COST:
        return "слишком большая себестоимость"
    if op.get("sum") is not None and op["sum"] > SX_MAX_SUM:
        return "слишком большая сумма закупки"
    if op.get("cost") is not None and op["cost"] > SX_MAX_COST:
        return "слишком большая себестоимость"
    return None


def _sx_step(cur: dict, op: dict) -> str | None:
    """Применяет операцию к словарю {qty, cost, min_qty}. Ошибка → текст."""
    q, c = float(cur["qty"]), float(cur["cost"])
    v = float(op["v"])
    if op["op"] == "in":
        if op.get("sum") is not None:
            # средневзвешенная, как при ручном приходе в карточке позиции
            cur["cost"] = (q * c + float(op["sum"])) / (q + v) if q + v > 0 else c
        cur["qty"] = q + v
    elif op["op"] == "out":
        if v > q + 1e-9:
            return f"списать {_q(v)}, а на складе {_q(q)}"
        cur["qty"] = q - v
    elif op["op"] == "set":
        cur["qty"] = v
        if op.get("cost") is not None:
            cur["cost"] = float(op["cost"])
    elif op["op"] == "cost":
        cur["cost"] = v
    elif op["op"] == "min":
        cur["min_qty"] = v
    return None


def sx_preview(ops: list[dict], errs: list[str], src: str) -> tuple[str, list[dict]]:
    """Текст предпросмотра и список операций, которые пройдут проверку."""
    ings = {int(i["id"]): i for i in db_list_ingredients()}
    errs = list(errs)
    sim: dict[int, dict] = {}
    good: list[dict] = []
    for op in ops:
        ing = ings.get(int(op["id"]))
        if not ing:
            errs.append(f"{op['src']}: позиции №{op['id']} нет на складе")
            continue
        bad = _sx_check_bounds(op)
        if bad:
            errs.append(f"{op['src']} (№{op['id']}): {bad}")
            continue
        cur = sim.setdefault(ing["id"], {k: float(ing[k]) for k in ("qty", "cost", "min_qty")})
        before = dict(cur)
        bad = _sx_step(cur, op)
        if bad:
            cur.update(before)
            errs.append(f"{op['src']} «{esc(ing['name'])}»: {bad}")
            continue
        good.append(op)
    L = [f"🔎 <b>Проверка правок ({esc(src)})</b>", ""]
    if sim:
        for iid, cur in sim.items():
            ing = ings[iid]
            u = esc(ing["unit"])
            parts = []
            if abs(cur["qty"] - float(ing["qty"])) > 1e-9:
                parts.append(f"остаток {_q(ing['qty'])} → <b>{_q(cur['qty'])}</b> {u}")
            if abs(cur["cost"] - float(ing["cost"])) > 1e-6:
                parts.append(f"с/с {_fmt_price(round(float(ing['cost']), 2))} → "
                             f"<b>{_fmt_price(round(cur['cost'], 2))}</b>")
            if abs(cur["min_qty"] - float(ing["min_qty"])) > 1e-9:
                parts.append(f"мин {_q(ing['min_qty'])} → <b>{_q(cur['min_qty'])}</b>")
            if parts:
                L.append(f"• №{iid} {esc(ing['name'])}: " + "; ".join(parts))
            else:
                L.append(f"• №{iid} {esc(ing['name'])}: без изменений")
    if not good:
        L.append("Изменений нет.")
    if errs:
        L += ["", f"⚠️ <b>Пропущу ({len(errs)}):</b>"] + [f"– {e}" for e in errs[:25]]
        if len(errs) > 25:
            L.append(f"…и ещё {len(errs) - 25}")
    text = "\n".join(L)
    if len(text) > 3900:  # лимит Telegram 4096
        text = text[:3850].rsplit("\n", 1)[0] + "\n…(список сокращён)"
    return text, good


def stock_apply_ops(ops: list[dict], src: str) -> tuple[list[str], list[str], list[dict]]:
    """Атомарно применяет операции к АКТУАЛЬНЫМ остаткам (за время между
    предпросмотром и ✅ могли пройти продажи). → (сделано, пропущено, затронутые)."""
    done: list[str] = []
    skipped: list[str] = []
    touched: dict[int, dict] = {}
    alert_ids: set[int] = set()   # алерт «мало» — только где менялся остаток/минимум
    now = _now()
    with _connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        for op in ops:
            r = conn.execute("SELECT * FROM ingredients WHERE id = ?", (int(op["id"]),)).fetchone()
            if not r:
                skipped.append(f"№{op['id']}: позиция удалена")
                continue
            ing = touched.get(int(r["id"])) or dict(r)
            if _sx_check_bounds(op):
                skipped.append(f"№{op['id']}: неверное значение")
                continue
            cur = {k: float(ing[k]) for k in ("qty", "cost", "min_qty")}
            bad = _sx_step(cur, op)
            if bad:
                skipped.append(f"«{esc(ing['name'])}»: {bad}")
                continue
            delta = cur["qty"] - float(ing["qty"])
            if abs(delta) > 1e-9 or abs(cur["min_qty"] - float(ing["min_qty"])) > 1e-9:
                alert_ids.add(int(ing["id"]))
            conn.execute(
                "UPDATE ingredients SET qty = ?, cost = ?, min_qty = ?, updated_at = ?"
                " WHERE id = ?", (cur["qty"], cur["cost"], cur["min_qty"], now, ing["id"]))
            if abs(delta) > 1e-9:
                reason = {"in": "приход", "out": "расход"}.get(op["op"], "инвентаризация")
                ref = src if op.get("sum") is None else f"{_fmt_price(op['sum'])} · {src}"
                conn.execute(
                    "INSERT INTO stock_moves (ingredient_id, delta, reason, ref, created_at)"
                    " VALUES (?, ?, ?, ?, ?)", (ing["id"], delta, reason, ref, now))
            label = {"in": f"+{_q(op['v'])}", "out": f"−{_q(op['v'])}",
                     "set": f"={_q(op['v'])}", "cost": f"с/с {_fmt_price(op['v'])}",
                     "min": f"мин {_q(op['v'])}"}[op["op"]]
            done.append(f"№{ing['id']} {esc(ing['name'])}: {label}")
            ing.update(cur)
            touched[int(ing["id"])] = ing
    return done, skipped, [i for k, i in touched.items() if k in alert_ids]


# ----------------------------- Excel (openpyxl) --------------------------------
_SX_COLS = ("ID", "Позиция", "Ед.", "Остаток", "С/с за ед., ₽", "Мин. остаток",
            "Вложено, ₽", "➕ Приход", "Сумма закупки, ₽", "➖ Расход",
            "Новый остаток", "Новая с/с, ₽", "Новый мин.")
# ключ → как узнать колонку при разборе (по заголовку, а не по букве)
_SX_INPUTS = (("in", "приход"), ("sum", "сумма закупки"), ("out", "расход"),
              ("set", "новый остаток"), ("cost", "новая с/с"), ("min", "новый мин"))


def _openpyxl():
    try:
        import openpyxl

        return openpyxl
    except Exception:  # noqa: BLE001
        logger.warning("Excel недоступен: добавьте openpyxl в requirements.txt")
        return None


def stock_xlsx_bytes() -> bytes:
    import io

    opx = _openpyxl()
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = opx.Workbook()
    ws = wb.active
    ws.title = "Склад"
    ws.append(list(_SX_COLS))
    head_fill = PatternFill("solid", fgColor="2B2B3A")
    in_fill = PatternFill("solid", fgColor="FFF2CC")
    ro_fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="BF8F00") if c.column >= 8 else head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ings = db_list_ingredients()
    for ing in ings:
        q, cst = float(ing["qty"]), float(ing["cost"])
        ws.append([int(ing["id"]), ing["name"], ing["unit"], q, round(cst, 2),
                   float(ing["min_qty"]), round(q * cst), None, None, None, None, None, None])
    last = len(ings) + 1
    for row in ws.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.border = border
            c.fill = in_fill if c.column >= 8 else ro_fill
            if c.column in (4, 6, 8, 10, 11, 13):
                c.number_format = "0.###"
            elif c.column in (5, 12):
                c.number_format = "#,##0.00"
            elif c.column in (7, 9):
                c.number_format = "#,##0"
    if last >= 2:
        dv = DataValidation(type="decimal", operator="between", formula1="0",
                            formula2=str(SX_MAX_SUM), allow_blank=True,
                            showErrorMessage=True, errorTitle="Только число",
                            error="Введите число ≥ 0 или оставьте пусто")
        ws.add_data_validation(dv)
        dv.add(f"H2:M{last}")
    for col, w in zip("ABCDEFGHIJKLM", (6, 34, 6, 10, 12, 12, 12, 11, 15, 11, 14, 13, 12)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:M{last}"
    info = wb.create_sheet("Как заполнять")
    for line in (
        "Как обновить склад файлом",
        "",
        "1. Заполняйте только ЖЁЛТЫЕ колонки (H–M). Пустая ячейка — без изменений, 0 — это значение.",
        "2. «➕ Приход» — сколько пришло; «Сумма закупки, ₽» — за весь приход (с/с пересчитается средневзвешенно).",
        "3. «➖ Расход» — списание (бой, порча, угощение).",
        "4. «Новый остаток» — после пересчёта (инвентаризация). Не заполняйте вместе с приходом/расходом.",
        "5. «Новая с/с, ₽» — себестоимость за единицу; «Новый мин.» — порог «мало» (0 — не следить).",
        "6. Колонки ID и названия не меняйте. Строки можно сортировать и фильтровать.",
        "7. Сохраните файл (.xlsx) и пришлите его боту в личные сообщения.",
        "   Бот покажет, что изменится, и применит только после ✅.",
        "",
        "Новые позиции файлом не создаются — используйте «➕ Новая позиция» в /stock.",
    ):
        info.append([line])
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 110
    meta = wb.create_sheet("_zz")
    meta["A1"] = SX_MARK
    meta["A2"] = _now()
    meta.sheet_state = "hidden"
    wb.properties.title = "Zig Zag — склад"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sx_cell_num(v) -> tuple[float | None, bool]:
    """Ячейка → (число|None, ошибка?). Пусто — (None, False)."""
    if v is None:
        return None, False
    if isinstance(v, bool):
        return None, True
    if isinstance(v, (int, float)):
        return (float(v), False) if v >= 0 else (None, True)
    s = str(v).strip().replace("\u00a0", "").replace(" ", "")
    if not s or s in ("-", "—"):
        return None, False
    s = re.sub(r"(₽|руб\.?|р\.?)$", "", s)
    n = _parse_num(s)
    return (n, False) if n is not None else (None, True)


def sx_parse_xlsx(data: bytes) -> tuple[list[dict], list[str]]:
    """Excel-файл склада → (операции, ошибки)."""
    import io

    opx = _openpyxl()
    try:
        wb = opx.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — не xlsx / повреждён
        return [], ["Не удалось открыть файл. Нужен .xlsx (Excel / Google Таблицы → .xlsx)."]
    try:
        ws = wb["Склад"] if "Склад" in wb.sheetnames else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None) or ()
        norm = [re.sub(r"[^\wа-яё/ ]", "", str(h or "").lower()).strip() for h in header]
        col: dict[str, int] = {}
        for i, h in enumerate(norm):
            if h == "id":
                col["id"] = i
            elif h == "позиция":
                col["name"] = i
            for key, needle in _SX_INPUTS:
                # «Расход» ≠ «приход»: сравниваем начало, «Новый остаток» ≠ «Остаток»
                if h.startswith(needle) and key not in col:
                    col[key] = i
        if "id" not in col or not any(k in col for k, _ in _SX_INPUTS):
            return [], ["Это не таблица склада: нет колонок «ID» и «➕ Приход»… "
                        "Скачайте шаблон кнопкой «📥 Скачать Excel» в /stock."]
        ings = {int(i["id"]): i for i in db_list_ingredients()}
        ops: list[dict] = []
        errs: list[str] = []
        for rn, row in enumerate(rows, start=2):
            if rn > 5000:
                errs.append("Файл длиннее 5000 строк — остальное не читал.")
                break
            row = tuple(row or ())

            def get(key):
                i = col.get(key)
                return row[i] if i is not None and i < len(row) else None

            raw_id = get("id")
            vals: dict[str, float] = {}
            bad_cells = []
            for key, _ in _SX_INPUTS:
                v, bad = _sx_cell_num(get(key))
                if bad:
                    bad_cells.append(_SX_COLS[7 + [k for k, _ in _SX_INPUTS].index(key)])
                elif v is not None:
                    vals[key] = v
            if raw_id is None and not vals and not bad_cells:
                continue  # пустая строка
            src = f"строка {rn}"
            try:
                iid = int(float(raw_id))
            except (TypeError, ValueError):
                if vals or bad_cells:
                    errs.append(f"{src}: нет ID позиции")
                continue
            if bad_cells:
                errs.append(f"{src} (ID {iid}): не число в «{', '.join(bad_cells)}»")
                continue
            if not vals:
                continue
            ing = ings.get(iid)
            fname = str(get("name") or "").strip()
            if ing and fname and fname != ing["name"]:
                errs.append(f"{src}: ID {iid} в базе «{esc(ing['name'])}», в файле "
                            f"«{esc(fname[:40])}» — строку пропустил (сверьте ID)")
                continue
            if "set" in vals and ("in" in vals or "out" in vals):
                errs.append(f"{src} (ID {iid}): «Новый остаток» вместе с приходом/расходом — "
                            "оставьте что-то одно")
                continue
            if "sum" in vals and "in" not in vals:
                errs.append(f"{src} (ID {iid}): сумма закупки без количества прихода")
                continue
            if "in" in vals:
                ops.append({"id": iid, "op": "in", "v": vals["in"], "sum": vals.get("sum"),
                            "src": src})
            if "out" in vals:
                ops.append({"id": iid, "op": "out", "v": vals["out"], "src": src})
            if "set" in vals:
                ops.append({"id": iid, "op": "set", "v": vals["set"], "cost": None, "src": src})
            if "cost" in vals:
                ops.append({"id": iid, "op": "cost", "v": vals["cost"], "src": src})
            if "min" in vals:
                ops.append({"id": iid, "op": "min", "v": vals["min"], "src": src})
            if len(ops) > SX_MAX_OPS:
                return [], [f"Слишком много изменений (>{SX_MAX_OPS}) — разбейте на части."]
        return ops, errs
    finally:
        wb.close()


def kb_sx_confirm(n: int) -> InlineKeyboardMarkup:
    rows = []
    if n:
        rows.append([_btn(f"✅ Применить ({n})", "sx:apply")])
    rows.append([_btn("✖ Отмена", "sx:cancel")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


async def _sx_show_preview(message: Message, state: FSMContext, ops: list[dict],
                           errs: list[str], src: str) -> None:
    text, good = sx_preview(ops, errs, src)
    if good:
        await state.set_state(StockBulkFSM.confirm)
        await state.update_data(sx_ops=good, sx_src=src)
    else:
        await state.clear()
    await message.answer(text, reply_markup=kb_sx_confirm(len(good)) if good else
                         InlineKeyboardMarkup(inline_keyboard=[
                             [_btn("✏️ Ещё раз", "sx:bulk"), _btn("📋 Таблица", "sx:table")]]))


async def _send_stock_table(message: Message) -> None:
    texts = stock_table_texts()
    for i, t in enumerate(texts):
        await message.answer(t, reply_markup=kb_stock_table() if i == len(texts) - 1 else None)


@router.callback_query(F.data.startswith("sx:"))
async def cb_stock_table(cb: CallbackQuery, state: FSMContext):
    if not _is_admin(cb.from_user.id):
        await cb.answer("⛔ Только для администратора", show_alert=True)
        return
    act = cb.data.split(":", 1)[1]
    try:
        if act == "table":
            await state.clear()
            await cb.answer()
            texts = stock_table_texts()
            if len(texts) == 1:
                try:
                    await cb.message.edit_text(texts[0], reply_markup=kb_stock_table())
                    return
                except TelegramBadRequest:
                    pass
            await _send_stock_table(cb.message)
            return
        if act == "bulk":
            await state.set_state(StockBulkFSM.text)
            ings = db_list_ingredients()
            await cb.answer()
            await cb.message.answer(_sx_help(ings[0]["id"] if ings else None),
                                    reply_markup=_cancel_kb("sx:cancel"))
            return
        if act == "xlsx":
            if not _openpyxl():
                await cb.answer("Excel недоступен: на сервере нет openpyxl "
                                "(добавьте в requirements.txt)", show_alert=True)
                return
            await cb.answer("Готовлю файл…")
            data = await asyncio.to_thread(stock_xlsx_bytes)
            await state.clear()
            await cb.message.answer_document(
                BufferedInputFile(data, f"zigzag_stock_{_now_dt():%Y-%m-%d}.xlsx"),
                caption="📥 Склад в Excel. Заполните <b>жёлтые</b> колонки (пусто — без "
                        "изменений), сохраните и пришлите файл сюда — покажу, что "
                        "изменится, и применю после ✅. Подробно — на листе «Как заполнять».",
            )
            return
        if act == "cancel":
            await state.clear()
            await cb.answer("Отменили")
            try:
                await cb.message.edit_text("✖ Правки склада отменены.",
                                           reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                                               [_btn("📋 Таблица", "sx:table")]]))
            except TelegramBadRequest:
                pass
            return
        if act == "apply":
            if await state.get_state() != StockBulkFSM.confirm:
                await cb.answer("Предпросмотр устарел — пришлите правки заново", show_alert=True)
                return
            data = await state.get_data()
            ops = data.get("sx_ops") or []
            src = data.get("sx_src") or "таблица"
            await state.clear()  # повторное нажатие ничего не применит дважды
            done, skipped, touched = stock_apply_ops(ops, src)
            await cb.answer("Готово ✅" if done else "Ничего не применено")
            L = [f"✅ <b>Склад обновлён ({esc(src)})</b>: {len(done)} изм.", ""]
            L += [f"• {d}" for d in done[:40]]
            if len(done) > 40:
                L.append(f"…и ещё {len(done) - 40}")
            if skipped:
                L += ["", "⚠️ <b>Не применено</b> (остатки изменились после проверки):"]
                L += [f"– {s}" for s in skipped[:20]]
            try:
                await cb.message.edit_text(
                    _fit_lines("\n".join(L)),
                    reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                        [_btn("📋 Таблица", "sx:table"), _btn("✏️ Ещё правки", "sx:bulk")]]))
            except TelegramBadRequest:
                pass
            if touched:
                await _stock_alerts(cb.bot, touched)
            return
        await cb.answer()
    except Exception:
        logger.exception("Ошибка в склад-таблице: %s", cb.data)
        try:
            await cb.answer("Ошибка", show_alert=True)
        except Exception:
            pass


@router.message(StockBulkFSM.text, F.text, ~F.text.startswith("/"))
async def stock_bulk_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    ops, errs = sx_parse_text(message.text or "")
    if not ops and not errs:
        await message.answer("Пусто. Пришлите правки или ✖ Отмена.")
        return
    await _sx_show_preview(message, state, ops, errs, "текстом")


@router.message(F.document, ~F.forward_origin)
async def stock_xlsx_upload(message: Message, state: FSMContext):
    """Админ прислал заполненный Excel склада."""
    if not message.from_user or not _is_admin(message.from_user.id):
        return
    doc = message.document
    name = (doc.file_name or "").lower()
    if not name.endswith(".xlsx"):
        await message.answer("📎 Склад принимаю только файлом .xlsx — шаблон: /stock → "
                             "«📥 Скачать Excel».")
        return
    if (doc.file_size or 0) > SX_MAX_FILE:
        await message.answer("Файл больше 2 МБ — это точно таблица склада?")
        return
    if not _openpyxl():
        await message.answer("Excel недоступен: на сервере нет openpyxl "
                             "(добавьте в requirements.txt).")
        return
    try:
        buf = await message.bot.download(doc)
        data = buf.read() if buf else b""
    except Exception:
        logger.exception("Не удалось скачать файл склада")
        await message.answer("Не удалось скачать файл, пришлите ещё раз.")
        return
    if len(data) > SX_MAX_FILE:
        await message.answer("Файл больше 2 МБ — это точно таблица склада?")
        return
    ops, errs = await asyncio.to_thread(sx_parse_xlsx, data)
    if not ops and not errs:
        await state.clear()
        await message.answer("В файле нет изменений: жёлтые колонки пустые.")
        return
    await _sx_show_preview(message, state, ops, errs, "Excel")


# ------------------------------- /stock: история -------------------------------
# «📜 История»: кто/что/когда изменило остаток. Все позиции (из таблицы) или одна
# (из карточки). Фильтры: приход · списания вручную · по заказам; период; страницы.
_SH_REASON = {
    "приход": "📥 приход",
    "расход": "📤 списание",
    "инвентаризация": "🔢 инвентаризация",
    "списание": "🧾 заказ",
    "возврат": "↩️ возврат",
}


def _sh_move_line(m: dict, with_name: bool) -> str:
    d = float(m["delta"])
    qty = f"{'+' if d >= 0 else '−'}{_fmt_qty(abs(d))} {esc(m['unit'])}"
    what = _SH_REASON.get(m["reason"], esc(m["reason"]))
    ref = (m.get("ref") or "").strip()
    if ref.startswith("заказ#"):
        what += f" #{esc(ref[6:])}"
    elif ref:
        what += f" · {esc(ref)}"
    name = f" · {esc(m['name'])}" if with_name else ""
    return f"<code>{(m.get('created_at') or '')[11:16]}</code> <b>{qty}</b>{name} · {what}"


def stock_history_text(ing: dict | None, flt: str, period: str, page: int) -> tuple[str, int]:
    """Текст экрана истории и число страниц."""
    h = db_stock_history(ing["id"] if ing else 0, flt, period, page)
    pages = max(1, -(-h["total"] // SH_PAGE))
    title = f"📜 <b>История: {esc(ing['name'])}</b>" if ing else "📜 <b>История склада</b> · все позиции"
    per = {"d": "сегодня (с 06:00)", "w": "7 дней", "m": "30 дней", "all": "всё время"}[period]
    L = [title, f"Период: {per}", ""]
    (ni, si), (no, so), (nz, sz) = h["tot"]["i"], h["tot"]["o"], h["tot"]["z"]
    if ing:
        u = esc(ing["unit"])
        L.append(f"📥 Приход: <b>{'+' if si else ''}{_fmt_qty(si)} {u}</b> ({ni})")
        L.append(f"📤 Списано вручную: <b>{'−' if so else ''}{_fmt_qty(abs(so))} {u}</b> ({no})")
        sign = "−" if sz < 0 else ("+" if sz > 0 else "")
        L.append(f"🧾 Ушло по заказам: <b>{sign}{_fmt_qty(abs(sz))} {u}</b> ({nz})")
    else:
        L.append(f"📥 Приходов: <b>{ni}</b> · 📤 списаний: <b>{no}</b> · 🧾 по заказам: <b>{nz}</b>")
    if not h["rows"]:
        L.append("\nДвижений нет — выберите другой период или фильтр.")
    else:
        if flt != "a":
            L.append(f"\nПоказаны: {SH_FILTERS[flt][0]}")
        day = None
        for m in h["rows"]:
            d = (m.get("created_at") or "")[:10]
            if d != day:
                day = d
                try:
                    dd = datetime.strptime(d, "%Y-%m-%d")
                    L.append(f"\n<b>{dd.day} {_MONTHS[dd.month - 1]}, {WDAYS[dd.weekday()]}</b>")
                except ValueError:
                    L.append(f"\n<b>{esc(d)}</b>")
            L.append(_sh_move_line(m, with_name=ing is None))
        if pages > 1:
            L += ["", f"Стр. {page + 1} из {pages} · всего {h['total']}"]
    text = "\n".join(L)
    if len(text) > 4000:
        text = text[:3990] + "\n…"
    return text, pages


def kb_stock_history(ing_id: int, flt: str, period: str, page: int, pages: int) -> InlineKeyboardMarkup:
    def cb(f=flt, p=period, pg=0):
        return f"sh:{ing_id}:{f}:{p}:{pg}"

    def mark(on: bool, text: str) -> str:
        return f"• {text} •" if on else text

    rows = [
        [_btn(mark(flt == k, SH_FILTERS[k][0]), cb(f=k)) for k in ("a", "i")],
        [_btn(mark(flt == k, SH_FILTERS[k][0]), cb(f=k)) for k in ("o", "z")],
        [_btn(mark(period == k, v[0]), cb(p=k)) for k, v in SH_PERIODS.items()],
    ]
    nav = []
    if page > 0:
        nav.append(_btn("◀️ Новее", cb(pg=page - 1)))
    if page + 1 < pages:
        nav.append(_btn("Старее ▶️", cb(pg=page + 1)))
    if nav:
        rows.append(nav)
    rows.append([_btn("🔙 К позиции", f"sk:{ing_id}") if ing_id else _btn("📋 Таблица", "sx:table"),
                 _btn("✖ Закрыть", "sk:close")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


@router.callback_query(F.data.startswith("sh:"))
async def cb_stock_history(cb: CallbackQuery, state: FSMContext):
    if not _is_admin(cb.from_user.id):
        await cb.answer("⛔ Только для администратора", show_alert=True)
        return
    try:
        _, ing_s, flt, period, page_s = cb.data.split(":")
        ing_id, page = int(ing_s), int(page_s)
        if flt not in SH_FILTERS or period not in SH_PERIODS or ing_id < 0 or not 0 <= page < 100_000:
            raise ValueError
    except ValueError:
        await cb.answer("Кнопка устарела — откройте /stock заново", show_alert=True)
        return
    try:
        ing = db_get_ingredient(ing_id) if ing_id else None
        if ing_id and not ing:
            await cb.answer("Позиция не найдена", show_alert=True)
            return
        await state.clear()
        text, pages = stock_history_text(ing, flt, period, page)
        if page >= pages:   # данных стало меньше — на последнюю страницу
            page = pages - 1
            text, pages = stock_history_text(ing, flt, period, page)
        await cb.answer()
        kb = kb_stock_history(ing_id, flt, period, page, pages)
        try:
            await cb.message.edit_text(text, reply_markup=kb)
        except TelegramBadRequest as e:
            if "not modified" not in str(e):
                await cb.message.answer(text, reply_markup=kb)
    except Exception:
        logger.exception("Ошибка в истории склада: %s", cb.data)
        try:
            await cb.answer("Ошибка", show_alert=True)
        except Exception:
            pass


# ------------------------------------------- /recipe -------------------------
def kb_rct_categories() -> InlineKeyboardMarkup:
    rows = [
        [_btn(f"{c['emoji']} {c['name']}", f"rct:{c['id']}")]
        for c in MENU["categories"]
    ]
    rows.append([_btn("✖ Закрыть", "rct:close")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_rct_items(cat_id: str) -> InlineKeyboardMarkup:
    cat = f_category(cat_id) or {"items": []}
    rows = []
    for it in cat["items"]:
        mark = "" if item_in_stock(it["id"]) else " ⛔"
        rows.append([_btn(f"{it['name']}{mark}", f"rct:i:{it['id']}")])
    rows.append([_btn("🔙 Разделы", "rct")])
    rows.append([_btn("✖ Закрыть", "rct:close")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def rct_view_text(item_id: str) -> str:
    it = _menu_item(item_id)
    name = it["name"] if it else item_id
    recs = db_recipes_for_item(item_id)
    lines = [
        f"📖 <b>Рецепт: {esc(name)}</b>",
        "Норма на <b>1 порцию</b>:",
        "",
    ]
    if not recs:
        lines.append("Пока пусто — при оплате заказа этот товар не спишется.")
    else:
        for r in recs:
            lines.append(f"• {esc(r['name'])} — {_fmt_qty(r['need'])} {esc(r['unit'])}")
    return "\n".join(lines)


def kb_rct_view(item_id: str) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    for r in db_recipes_for_item(item_id):
        rows.append(
            [
                _btn(f"{r['name']} · {_fmt_qty(r['need'])} {r['unit']}", "noop"),
                _btn("🗑", f"rct:d:{item_id}:{r['id']}"),
            ]
        )
    rows.append([_btn("➕ Ингредиент", f"rct:a:{item_id}")])
    cat = _category_of_item(item_id)
    if cat:
        rows.append([_btn("🔙 К позициям", f"rct:{cat}")])
    rows.append([_btn("✖ Закрыть", "rct:close")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


@router.message(Command("recipe"))
async def cmd_recipe(message: Message, state: FSMContext):
    """Рецепты: из чего списывать позиции меню."""
    if message.from_user.id not in ADMIN_IDS:
        return
    await state.clear()
    await message.answer(
        "📖 <b>Рецепты</b>\n\nВыберите раздел меню:", reply_markup=kb_rct_categories()
    )


@router.callback_query(F.data.startswith("rct"))
async def cb_recipe(cb: CallbackQuery, state: FSMContext):
    if not _is_admin(cb.from_user.id):
        await cb.answer("⛔ Только для администратора", show_alert=True)
        return
    parts = cb.data.split(":")
    try:
        if cb.data == "rct":  # разделы
            await cb.answer()
            await cb.message.edit_text(
                "📖 <b>Рецепты</b>\n\nВыберите раздел меню:",
                reply_markup=kb_rct_categories(),
            )
            return
        if parts[1] == "cancel":
            data = await state.get_data()
            await state.clear()
            await cb.answer("Отменили")
            item_id = data.get("rct_item")
            if item_id:
                try:
                    await cb.message.edit_text(
                        rct_view_text(item_id), reply_markup=kb_rct_view(item_id)
                    )
                except TelegramBadRequest:
                    pass
            return
        if parts[1] == "close":
            await state.clear()
            await cb.answer("Закрыли")
            try:
                await cb.message.edit_text("✅ Готово", reply_markup=None)
            except TelegramBadRequest:
                pass
            return
        if len(parts) == 2:  # раздел → позиции
            cat = f_category(parts[1])
            if not cat:
                await cb.answer("Раздел не найден", show_alert=True)
                return
            await cb.answer()
            await cb.message.edit_text(
                f"{cat['emoji']} <b>{esc(cat['name'])}</b> — выберите позицию:",
                reply_markup=kb_rct_items(cat["id"]),
            )
            return
        if parts[1] == "i":  # карточка рецепта
            await cb.answer()
            await cb.message.edit_text(
                rct_view_text(parts[2]), reply_markup=kb_rct_view(parts[2])
            )
            return
        if parts[1] == "a":  # выбор ингредиента (по страницам)
            item_id = parts[2]
            page = int(parts[3]) if len(parts) > 3 and parts[3].isdigit() and len(parts[3]) < 5 else 0
            ings = db_list_ingredients()
            if not ings:
                await cb.answer()
                await cb.message.edit_text(
                    "На складе пока нет ингредиентов.\n"
                    "Сначала добавьте их в /stock — потом вернитесь сюда.",
                    reply_markup=InlineKeyboardMarkup(
                        inline_keyboard=[[_btn("🔙 Назад", f"rct:i:{item_id}")]]
                    ),
                )
                return
            chunk, page, pages = _sk_page(ings, page)
            rows = [
                [_btn(f"{i['name']} · {i['unit']}", f"rct:p:{item_id}:{i['id']}")]
                for i in chunk
            ]
            if pages > 1:
                nav = []
                if page > 0:
                    nav.append(_btn("◀️", f"rct:a:{item_id}:{page - 1}"))
                nav.append(_btn(f"{page + 1}/{pages}", "noop"))
                if page < pages - 1:
                    nav.append(_btn("▶️", f"rct:a:{item_id}:{page + 1}"))
                rows.append(nav)
            rows.append([_btn("🔙 Назад", f"rct:i:{item_id}")])
            await cb.answer()
            note = (f"\nСтр. {page + 1}/{pages} · позиций на складе: {len(ings)} (по алфавиту)"
                    if pages > 1 else "")
            await cb.message.edit_text(
                "➕ Какой ингредиент добавить?" + note,
                reply_markup=InlineKeyboardMarkup(inline_keyboard=rows),
            )
            return
        if parts[1] == "p":  # ввод количества на порцию
            item_id = parts[2]
            ing = db_get_ingredient(int(parts[3]))
            if not ing:
                await cb.answer("Ингредиент не найден", show_alert=True)
                return
            it = _menu_item(item_id)
            await cb.answer()
            await state.set_state(RecipeFSM.qty)
            await state.update_data(
                rct_item=item_id,
                rct_ing=ing["id"],
                ed_msg=cb.message.message_id,
            )
            await cb.message.edit_text(
                f"«{esc(it['name'] if it else item_id)}» + <b>{esc(ing['name'])}</b>\n\n"
                f"Сколько {esc(ing['unit'])} на 1 порцию? Отправьте числом:",
                reply_markup=_cancel_kb("rct:cancel"),
            )
            return
        if parts[1] == "d":  # убрать ингредиент из рецепта
            item_id = parts[2]
            db_delete_recipe(item_id, int(parts[3]))
            await cb.answer("Убрали из рецепта")
            await cb.message.edit_text(
                rct_view_text(item_id), reply_markup=kb_rct_view(item_id)
            )
            return
    except Exception:
        logger.exception("Ошибка в колбэке рецепта: %s", cb.data)
        try:
            await cb.answer("Ошибка", show_alert=True)
        except Exception:
            pass


@router.message(RecipeFSM.qty, F.text)
async def recipe_qty_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    data = await state.get_data()
    item_id = data.get("rct_item")
    ing = db_get_ingredient(data.get("rct_ing") or 0)
    v = _parse_num((message.text or "").strip())
    if v is None or v <= 0 or v > 1_000_000:
        await message.answer(
            "Введите положительное число (например 10), либо ✖ Отмена кнопкой выше."
        )
        return
    await state.clear()
    if not item_id or not ing:
        await message.answer("Не удалось сохранить — откройте /recipe заново.")
        return
    db_upsert_recipe(item_id, ing["id"], v)
    await _edit_or_send(
        message,
        data,
        f"✅ {esc(ing['name'])} — {_fmt_qty(v)} {esc(ing['unit'])} на порцию\n\n"
        + rct_view_text(item_id),
        kb_rct_view(item_id),
    )


@router.message(Command("bookings"))
async def cmd_bookings(message: Message):
    if message.from_user.id not in ADMIN_IDS:
        return
    res = db_upcoming_reservations(10)
    if not res:
        await message.answer("Ближайших броней нет.")
        return
    lines = ["📅 Ближайшие брони:", ""]
    for r in res:
        kind = RSV_KINDS.get(r["kind"], r["kind"])
        status = RSV_STATUSES.get(r["status"], r["status"])
        d = r["date"][8:10] + "." + r["date"][5:7]
        lines.append(
            f"#{r['id']} · {kind} · {d} {r['time']} · "
            f"{r['guests']} {_guests_word(r['guests'])} · {status}"
        )
    await message.answer("\n".join(lines))


# -------------------------------- кнопки статусов в канале -------------------
_admin_cache: dict[int, tuple[bool, float]] = {}
_ADMIN_CACHE_TTL = 300  # сек


async def _is_channel_admin(bot, user_id: int) -> bool:
    """Админ канала персонала или админ из ADMIN_IDS.
    При ЛЮБОЙ ошибке проверки — отказ (fail-closed), результат ошибки не кэшируем."""
    if _is_admin(user_id):
        return True
    if not CHANNEL_ID:
        return False
    now = time.monotonic()
    cached = _admin_cache.get(user_id)
    if cached and now - cached[1] < _ADMIN_CACHE_TTL:
        return cached[0]
    try:
        member = await bot.get_chat_member(CHANNEL_ID, user_id)
        ok = member.status in ("creator", "administrator")
    except Exception as e:
        logger.warning("Проверка прав %s в канале не удалась (%s) — отказ", user_id, e)
        return False
    _admin_cache[user_id] = (ok, now)
    return ok


def _from_staff_chat(cb: CallbackQuery) -> bool:
    """Нажатие пришло из канала персонала (а не подделано в личке с ботом)."""
    chat = getattr(cb.message, "chat", None) if cb.message else None
    if chat is None or not CHANNEL_ID:
        return False
    if isinstance(CHANNEL_ID, int):
        return chat.id == CHANNEL_ID
    uname = (getattr(chat, "username", None) or "").lower()
    return bool(uname) and uname == str(CHANNEL_ID).lstrip("@").lower()


def _cb_id(raw: str) -> int:
    """Номер заказа/брони из кнопки; мусор и огромные числа → ValueError."""
    v = int(raw)
    if not 0 < v < 2**62:
        raise ValueError(raw)
    return v


async def _staff_guard(cb: CallbackQuery, deny: str) -> bool:
    """True — можно выполнять действие персонала. Иначе отвечает отказом."""
    if not _from_staff_chat(cb):
        await cb.answer("⛔ Действие доступно только в канале персонала", show_alert=True)
        return False
    if not await _is_channel_admin(cb.bot, cb.from_user.id):
        await cb.answer(deny, show_alert=True)
        return False
    return True


@router.callback_query(F.data.startswith("st:"))
async def cb_status(cb: CallbackQuery):
    try:
        _, oid_raw, code = cb.data.split(":")
        order_id = _cb_id(oid_raw)
    except (ValueError, AttributeError):
        await cb.answer("Некорректный запрос", show_alert=True)
        return

    if code not in STATUSES:
        await cb.answer("Неизвестный статус", show_alert=True)
        return

    if not await _staff_guard(cb, "⚠️ Статус меняют только администраторы канала"):
        return

    order = db_get_order(order_id)
    if not order:
        await cb.answer("Заказ не найден (возможно, удалён)", show_alert=True)
        return

    label = STATUSES[code]
    if order["status"] == code:
        await cb.answer(f"Уже: {label}")
        return

    # Заказ уже оплачен — прочие статусы не ставим.
    # Выход только один: ❌ Отменён (возврат / перерасчёт).
    if order["status"] == "paid" and code != "cancelled":
        await cb.answer(
            "🔒 Заказ уже оплачен — статус изменить нельзя. "
            "Возврат/перерасчёт: ❌ Отменён",
            show_alert=True,
        )
        return

    # предзаказ на завтра — гость ещё не пришёл (отмена и «Принят» — можно)
    if code in ("ready", "issued", "paid") and _pre_too_early(order):
        await cb.answer(_pre_early_msg(order), show_alert=True)
        return

    # «Оплачен» — сначала спрашиваем у админа, как была произведена оплата
    if code == "paid":
        try:
            await cb.message.edit_reply_markup(reply_markup=kb_pay(order_id))
        except TelegramBadRequest:
            pass
        await cb.answer("💵 Уточните, как была произведена оплата")
        return

    db_set_status(order_id, code)
    order = db_get_order(order_id)

    refund = None
    if code == "cancelled":
        # возврат ингредиентов на склад (ровно один раз — атомарный флаг)
        if db_claim_stock_flag(order_id, 0):
            affected = stock_apply_order(order, +1)
            await _stock_alerts(cb.bot, affected)
        # возврат баллов / отмена кешбэка (ровно один раз)
        refund = loyalty_on_refund(order_id)
        db_set_bonus_pending(order_id, 0)
        order = db_get_order(order_id)

    try:
        await cb.message.edit_text(order_text(order), reply_markup=kb_status(order_id))
    except Exception:
        logger.exception("Не удалось отредактировать сообщение заказа #%s", order_id)

    # --- уведомление гостю о смене статуса заказа ---
    if order.get("user_id"):
        note = _order_notify_text(order)
        if refund and (refund["returned"] or refund["taken"]):
            note += (
                f"\n💎 Баллы: возвращено {_fmt_pts(refund['returned'])},"
                f" списан кешбэк {_fmt_pts(refund['taken'])}."
                f" Баланс: {_fmt_pts(refund['balance'] or 0)}"
            )
        try:
            await cb.bot.send_message(order["user_id"], note)
        except Exception:
            logger.info(
                "Не удалось уведомить гостя заказа #%s (user_id=%s)",
                order_id,
                order.get("user_id"),
            )

    who = cb.from_user.username or cb.from_user.full_name
    suffix = f" ({who})" if who and not str(cb.from_user.id).startswith("-") else ""
    await cb.answer(f"{label}{suffix}")


@router.callback_query(F.data.startswith("pay:"))
async def cb_pay(cb: CallbackQuery):
    """Выбор способа оплаты: нал / безнал / перевод, баллы (или «Назад»)."""
    try:
        _, oid_raw, action = cb.data.split(":")
        order_id = _cb_id(oid_raw)
    except (ValueError, AttributeError):
        await cb.answer("Некорректный запрос", show_alert=True)
        return

    if not await _staff_guard(cb, "⚠️ Оплату подтверждают только администраторы канала"):
        return

    order = db_get_order(order_id)
    if not order:
        await cb.answer("Заказ не найден", show_alert=True)
        return

    if action == "back":
        db_set_bonus_pending(order_id, 0)
        order = db_get_order(order_id)
        try:
            await cb.message.edit_text(order_text(order), reply_markup=kb_status(order_id))
        except TelegramBadRequest:
            pass
        await cb.answer("Отменили")
        return

    if order["status"] != "paid" and _pre_too_early(order):
        await cb.answer(_pre_early_msg(order), show_alert=True)
        return

    if action in ("bonus", "nobonus"):
        if order["status"] == "paid":
            await cb.answer("Заказ уже оплачен", show_alert=True)
            return
        if action == "bonus":
            pts = bonus_available_for_order(order)
            if pts <= 0:
                await cb.answer("Баллов для списания нет", show_alert=True)
                return
            db_set_bonus_pending(order_id, pts)
            msg = f"🎁 Спишем {_fmt_pts(pts)} баллов · к оплате {_fmt_money(order['total'] - pts)}. Выберите способ оплаты"
        else:
            db_set_bonus_pending(order_id, 0)
            msg = "Баллы не списываем — выберите способ оплаты"
        order = db_get_order(order_id)
        try:
            await cb.message.edit_text(order_text(order), reply_markup=kb_pay(order_id))
        except TelegramBadRequest:
            pass
        await cb.answer(msg, show_alert=action == "bonus")
        return

    if action not in PAY_METHODS:
        await cb.answer("Неизвестный способ оплаты", show_alert=True)
        return

    if not db_claim_paid(order_id, action):
        # уже оплачен — только поправить способ, без повторного списания
        db_set_pay_method(order_id, action)
        order = db_get_order(order_id)
        try:
            await cb.message.edit_text(order_text(order), reply_markup=kb_status(order_id))
        except Exception:
            logger.exception("Не удалось обновить заказ #%s", order_id)
        await cb.answer(f"Способ оплаты: {PAY_METHODS[action]} (уже оплачен)")
        return

    # мы перевели заказ в «Оплачен» — баллы и склад ровно один раз
    loyal = loyalty_on_paid(order_id)
    if db_claim_stock_flag(order_id, 1):
        affected = stock_apply_order(order, -1)
        await _stock_alerts(cb.bot, affected)
    order = db_get_order(order_id)

    try:
        await cb.message.edit_text(order_text(order), reply_markup=kb_status(order_id))
    except Exception:
        logger.exception("Не удалось обновить заказ #%s после оплаты", order_id)

    # уведомляем гостя
    if order.get("user_id"):
        note = _order_notify_text(order)
        if loyal["spent"] or loyal["accrued"]:
            parts = []
            if loyal["spent"]:
                parts.append(f"списано {_fmt_pts(loyal['spent'])}")
            if loyal["accrued"]:
                parts.append(f"начислено +{_fmt_pts(loyal['accrued'])}")
            note += f"\n💎 Баллы: {', '.join(parts)}. Баланс: {_fmt_pts(loyal['balance'] or 0)}"
        try:
            await cb.bot.send_message(order["user_id"], note)
        except Exception:
            logger.info("Не удалось уведомить гостя заказа #%s", order_id)

    tail = f" · −{_fmt_pts(loyal['spent'])} баллами" if loyal["spent"] else ""
    await cb.answer(f"Оплата: {PAY_METHODS[action]} ✅{tail}")


# ------------------------------------------------------ /start ---------------
@router.message(CommandStart())
async def cmd_start(message: Message, command: CommandObject, state: FSMContext):
    # Поддержка QR-ссылки t.me/bot?start=5 или start=table5 — стол запоминаем сразу
    if command.args:
        m = re.search(r"\d+", command.args)
        if m and 1 <= int(m.group()) <= 8:
            await state.update_data(table=str(int(m.group())))
    await state.set_state(None)
    await message.answer(WELCOME, reply_markup=kb_welcome())


@router.message(Command("table"))
async def cmd_table(message: Message, state: FSMContext):
    """Сменить стол."""
    await state.set_state(OrderFSM.table)
    await message.answer(ASK_TABLE)


# ----------------------------------------------------- ввод стола ------------
@router.message(OrderFSM.table)
async def on_table_input(message: Message, state: FSMContext):
    raw = (message.text or "").strip()
    if not raw.isdigit() or not 1 <= int(raw) <= 8:
        await message.answer(ERR_TABLE)
        return
    table = str(int(raw))
    await state.update_data(table=table)
    await state.set_state(None)
    await message.answer(table_set_text(table), reply_markup=kb_main(table))


# ====================== СЦЕНАРИЙ: ВЫБОР → ЛОЯЛЬНОСТЬ ========================
# /start → 2 кнопки (бронь / заказ) → первый раз или постоянный → телефон →
# дальше по выбранному сценарию.


def kb_welcome() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [_btn("📅 Забронировать стол / VIP", "go:rsv")],
            [_btn("🛒 Сделать заказ", "go:order")],
            [_btn("📝 Предзаказ к приходу", "go:pre")],
        ]
    )


def kb_loyal() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [_btn("🌟 Первый раз", "loyal:first")],
            [_btn("💎 Постоянный гость", "loyal:reg")],
            [_btn("🔙 Назад", "go:home")],
        ]
    )


def kb_phone() -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="📱 Отправить мой номер", request_contact=True)],
            [KeyboardButton(text=PHONE_SKIP)],
            [KeyboardButton(text=PHONE_BACK)],
        ],
        resize_keyboard=True,
        input_field_placeholder="+7 900 123-45-67",
    )


async def _back_to_welcome(message: Message, state: FSMContext) -> None:
    """Вернуть гостя на экран выбора (бронь/заказ), убрав клавиатуру телефона."""
    await state.set_state(None)
    await message.answer(WELCOME, reply_markup=ReplyKeyboardRemove())
    await message.answer("Выберите действие 👇", reply_markup=kb_welcome())


async def _proceed(target: str, message: Message, state: FSMContext) -> None:
    """Продолжение сценария после лояльности: 'rsv' — бронь, иначе — заказ."""
    data = await state.get_data()
    if target == "card":
        await _send_card(message, state, message.chat.id)
        return
    if target == "pre":
        await _pre_home(message, state, message.chat.id)
        return
    if target == "rsv":
        await state.update_data(rsv={})
        await message.answer(
            "📅 <b>Бронирование</b>\n\nЧто будем бронировать?",
            reply_markup=kb_rsv_kind(),
        )
        return
    # --- заказ ---
    table = data.get("table")
    if table:
        await message.answer(main_text(table), reply_markup=kb_main(table))
        await message.answer(
            "📋 Выберите раздел:",
            reply_markup=kb_categories(_cart_count(data.get("cart") or [])),
        )
    else:
        await state.set_state(OrderFSM.table)
        await message.answer(ASK_TABLE, reply_markup=ReplyKeyboardRemove())


async def _begin_flow(
    target: str,
    message: Message,
    state: FSMContext,
    cb: CallbackQuery | None = None,
    strip: bool = True,
) -> None:
    """Старт сценария (бронь/заказ): сначала лояльность, если карты ещё нет
    и гость не пропускал вопрос в этой сессии."""
    user = cb.from_user if cb is not None else message.from_user
    if cb is not None:
        if strip:
            try:
                await cb.message.edit_reply_markup(reply_markup=None)  # убрать кнопки приветствия
            except TelegramBadRequest:
                pass
        await cb.answer()
    await state.update_data(next=target)
    data = await state.get_data()
    card = guest_card(user.id if user else None)
    if card:
        db_guest_touch(user)
        if not data.get("card_greeted"):
            await state.update_data(card_greeted=True)
            await message.answer(
                f"💎 Карта <b>{esc(card['phone'])}</b> · баллов: "
                f"<b>{_fmt_pts(db_bonus_balance(user.id))}</b>"
            )
        await _proceed(target, message, state)
    elif data.get("loyal_done"):
        await _proceed(target, message, state)
    else:
        await message.answer(LOYAL_ASK, reply_markup=kb_loyal())


async def _finish_phone(
    phone: str | None, message: Message, state: FSMContext, verified: bool = False
) -> None:
    data = await state.get_data()
    card = guest_card(message.from_user.id if message.from_user else None)
    if card and phone != card["phone"]:
        await state.update_data(loyal_done=True, card_greeted=True)
        await state.set_state(None)
        await message.answer(
            f"💎 Ваша карта: <b>{esc(card['phone'])}</b>. {CARD_PHONE_LOCKED}",
            reply_markup=ReplyKeyboardRemove(),
        )
        await _proceed(data.get("next") or "order", message, state)
        return
    if phone:
        ok, why = db_guest_set_phone(message.from_user, phone, verified)
        if not ok:
            await message.answer(
                "⚠️ Этот номер уже привязан к другой карте.\n\n"
                "Если это ваш номер — отправьте его кнопкой «📱 Отправить мой номер» "
                "(подтверждение через Telegram) или обратитесь к администратору. "
                "Можно также «⏭ Пропустить».",
                reply_markup=kb_phone(),
            )
            return  # остаёмся на шаге телефона
    await state.update_data(phone=phone, loyal_done=True, card_greeted=bool(phone))
    await state.set_state(None)
    if phone:
        bal = db_bonus_balance(message.from_user.id)
        await message.answer(
            f"✅ Карта лояльности привязана! Номер карты: <b>{esc(phone)}</b>\n"
            f"🎁 Баллов: <b>{_fmt_pts(bal)}</b> · кешбэк {LOYALTY_PERCENT}% с каждого "
            f"оплаченного заказа. Карта — кнопка «{BTN_CARD}».",
            reply_markup=ReplyKeyboardRemove(),
        )
    else:
        await message.answer(
            "Хорошо, пропустили 👌 Без номера баллы не начисляются — "
            f"подключить карту можно позже кнопкой «{BTN_CARD}».",
            reply_markup=ReplyKeyboardRemove(),
        )
    await _proceed(data.get("next") or "order", message, state)


@router.callback_query(F.data.startswith("go:"))
async def go_choice(cb: CallbackQuery, state: FSMContext):
    try:
        target = cb.data.split(":", 1)[1]  # "go:rsv" → "rsv"
        if target == "home":
            # «Назад» — вернуться на экран выбора бронь/заказ
            await cb.answer()
            await cb.message.edit_text(WELCOME, reply_markup=kb_welcome())
            return
        if target not in ("rsv", "order", "pre"):
            await cb.answer("Ошибка выбора", show_alert=True)
            return
        await _begin_flow(target, cb.message, state, cb=cb)
    except Exception:
        logger.exception("Ошибка в go_choice (data=%s)", cb.data)
        try:
            await cb.answer("⚠️ Ошибка — посмотрите логи бота на хостинге", show_alert=True)
        except Exception:
            pass


@router.callback_query(F.data.startswith("loyal:"))
async def loyal_pick(cb: CallbackQuery, state: FSMContext):
    try:
        kind = cb.data.split(":", 1)[1]  # "loyal:first" → "first"
        if kind not in ("first", "reg"):
            await cb.answer("Ошибка", show_alert=True)
            return
        card = guest_card(cb.from_user.id)
        if card:   # карта уже есть — номер не спрашиваем
            await cb.answer()
            try:
                await cb.message.edit_text(f"💎 Карта <b>{esc(card['phone'])}</b> уже привязана ✅")
            except TelegramBadRequest:
                pass
            await state.update_data(loyal_done=True, card_greeted=True)
            await _proceed((await state.get_data()).get("next") or "order", cb.message, state)
            return
        await state.update_data(loyal_type=kind)
        await cb.answer()
        prompt = LOYAL_FIRST if kind == "first" else LOYAL_REG
        try:
            await cb.message.edit_text(prompt)  # убираем кнопки выбора
        except TelegramBadRequest:
            pass
        await cb.message.answer(PHONE_ASK, reply_markup=kb_phone())
        await state.set_state(OrderFSM.phone)
    except Exception:
        logger.exception("Ошибка в loyal_pick (data=%s)", cb.data)
        try:
            await cb.answer("⚠️ Ошибка — посмотрите логи бота на хостинге", show_alert=True)
        except Exception:
            pass


@router.message(OrderFSM.phone, F.contact)
async def phone_contact(message: Message, state: FSMContext):
    if not message.contact:
        return
    # принимаем как подтверждённый только СВОЙ контакт (кнопка «Отправить мой номер»)
    if not message.from_user or message.contact.user_id != message.from_user.id:
        await message.answer(
            "⚠️ Это контакт другого человека. Отправьте свой номер кнопкой "
            "«📱 Отправить мой номер» или напишите его текстом.",
            reply_markup=kb_phone(),
        )
        return
    phone = normalize_phone(message.contact.phone_number)
    if not phone:
        await message.answer("Не удалось распознать номер — напишите его текстом.")
        return
    await _finish_phone(phone, message, state, verified=True)


@router.message(OrderFSM.phone, F.text)
async def phone_text(message: Message, state: FSMContext):
    t = (message.text or "").strip()
    if t == PHONE_BACK:
        await _back_to_welcome(message, state)
        return
    if t == PHONE_SKIP:
        await _finish_phone(None, message, state)
        return
    phone = normalize_phone(t)
    if phone:
        await _finish_phone(phone, message, state, verified=False)
        return
    await message.answer(
        "Похоже, это не номер 😅\n\nПришлите его кнопкой «📱 Отправить мой номер», "
        "напишите в чат (например +7 900 123-45-67), нажмите «⏭ Пропустить» "
        "или «🔙 Назад», чтобы вернуться к выбору."
    )


@router.message(OrderFSM.phone)
async def phone_other(message: Message, state: FSMContext):
    await message.answer(
        "📱 Пришлите номер: кнопкой «📱 Отправить мой номер», текстом "
        "или нажмите «⏭ Пропустить»."
    )


# --- ввод имени, на кого бронь (шаг между гостями и длительностью) ---
@router.message(OrderFSM.rsv_name, F.text)
async def rsv_name_input(message: Message, state: FSMContext):
    t = (message.text or "").strip()
    if t in (BTN_MENU, BTN_CART, BTN_BOOK, BTN_CARD, BTN_PRE):
        await message.answer("✍️ Введите имя для брони обычным сообщением")
        return
    if not t or len(t) > 120:
        await message.answer("Введите имя (до 120 символов) сообщением ✍️")
        return
    data = await state.get_data()
    rsv = dict(data.get("rsv") or {})
    rsv["name"] = t
    await state.update_data(rsv=rsv)
    await state.set_state(None)
    # отметим экран ввода имени
    msg_id = data.get("rsv_name_msg")
    if msg_id:
        try:
            await message.bot.edit_message_text(
                chat_id=message.chat.id,
                message_id=msg_id,
                text=f"✍️ Имя: <b>{esc(t)}</b> ✅",
            )
        except TelegramBadRequest:
            pass
    # 🍸 барная стойка — бронь места без часов: сразу к подтверждению
    if rsv.get("kind") == "bar":
        nm = rsv["name"]
        if message.from_user and message.from_user.username and "(@" not in nm:
            nm += f" (@{message.from_user.username})"
        await message.answer(
            _rsv_confirm_text(rsv, nm), reply_markup=kb_rsv_confirm(rsv)
        )
        return
    text, kb = _rsv_dur_screen(rsv)
    await message.answer(text, reply_markup=kb)


@router.message(OrderFSM.rsv_name)
async def rsv_name_other(message: Message, state: FSMContext):
    await message.answer("✍️ Пришлите имя обычным текстовым сообщением")


# ------------------------------------------- кнопки reply-меню ---------------
@router.message(F.text == BTN_MENU)
async def on_btn_menu(message: Message, state: FSMContext):
    if not await _table(state):
        await state.set_state(OrderFSM.table)
        await message.answer(ASK_TABLE)
        return
    await _show_menu(message, state)


@router.message(F.text == BTN_CART)
async def on_btn_cart(message: Message, state: FSMContext):
    cart = await _cart(state)
    table = await _table(state)
    if not cart:
        await message.answer(CART_EMPTY, reply_markup=kb_main(table))
        return
    await message.answer(cart_text(cart, table), reply_markup=kb_cart(cart))


@router.message(F.text.startswith("🪑 Стол:"))
async def on_btn_change_table(message: Message, state: FSMContext):
    await state.set_state(OrderFSM.table)
    await message.answer(ASK_TABLE)


# ------------------------------------------------- навигация по меню ---------
@router.callback_query(F.data == "m")
async def cb_menu(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    await _show_menu(cb.message, state, edit=True)


_photo_file_ids: dict[str, str] = {}  # кэш file_id: фото грузится в Telegram один раз


async def _send_info_photo(cb: CallbackQuery, which: str) -> None:
    """Фото «Выгода»/«Правила» из файла рядом с bot.py (после первой отправки —
    по file_id из кэша, без повторной загрузки)."""
    path = PROMO_PHOTO if which == "promo" else RULES_PHOTO
    kb = kb_info_photo(promo=(which == "promo"))
    fid = _photo_file_ids.get(which)
    if fid:
        try:
            await cb.message.answer_photo(fid, reply_markup=kb)
            return
        except TelegramBadRequest:
            _photo_file_ids.pop(which, None)
    if not path.is_file():
        logger.warning("Нет файла %s рядом с bot.py — картинку не показать", path.name)
        await cb.message.answer("🖼 Картинка скоро появится — загляните чуть позже.",
                                reply_markup=kb)
        return
    sent = await cb.message.answer_photo(FSInputFile(path), reply_markup=kb)
    try:
        _photo_file_ids[which] = sent.photo[-1].file_id
    except (AttributeError, IndexError, TypeError):
        pass


@router.callback_query(F.data == "info:benefit")
async def cb_info_benefit(cb: CallbackQuery):
    await cb.answer()
    try:
        await _send_info_photo(cb, "promo")
    except Exception:
        logger.exception("Не удалось отправить фото «Выгода»")


@router.callback_query(F.data == "info:rules")
async def cb_info_rules(cb: CallbackQuery):
    await cb.answer()
    try:
        await _send_info_photo(cb, "rules")
    except Exception:
        logger.exception("Не удалось отправить фото «Правила»")


@router.callback_query(F.data == "info:close")
async def cb_info_close(cb: CallbackQuery):
    await cb.answer("Меню выше 👇")
    try:
        await cb.message.delete()
    except Exception:
        logger.info("Не удалось удалить фото-сообщение (info:close)")


@router.callback_query(F.data.startswith("c:"))
async def cb_category(cb: CallbackQuery, state: FSMContext):
    cat = f_category(cb.data[2:])
    if not cat:
        await cb.answer(ERR_MENU, show_alert=True)
        return
    await cb.answer()
    cart = await _cart(state)
    await cb.message.edit_text(
        f"{cat['emoji']} <b>{esc(cat['name'])}</b> — нажмите, чтобы добавить:",
        reply_markup=kb_category(cat["id"], _cart_count(cart)),
    )


def _category_of_item(item_id: str) -> str | None:
    for c in MENU["categories"]:
        if any(it["id"] == item_id for it in c["items"]):
            return c["id"]
    return None


@router.callback_query(F.data.startswith("a:"))
async def cb_add(cb: CallbackQuery, state: FSMContext):
    item = f_find_item(cb.data[2:])
    if not item:
        await cb.answer(ERR_MENU, show_alert=True)
        return
    if not await _table(state):
        await cb.answer(ASK_TABLE_FIRST, show_alert=True)
        return
    cart = await _cart(state)
    if not _is_admin(cb.from_user.id) and _cart_count(cart) >= ORDER_MAX_ITEMS:
        await cb.answer(
            f"⚠️ В одном заказе не больше {ORDER_MAX_ITEMS} позиций. "
            "Отправьте этот заказ — и сделайте следующий.",
            show_alert=True,
        )
        return
    # лимит по остаткам: столько позиции можно набрать по рецепту
    limit = max_available(item["id"])
    in_cart = next((i["qty"] for i in cart if i["id"] == item["id"]), 0)
    if limit is not None and in_cart + 1 > limit:
        await cb.answer(
            f"⚠️ «{item['name']}»: по остаткам на складе можно не более {limit} шт. "
            "Ингредиентов на большее не хватает.",
            show_alert=True,
        )
        return
    found = next((i for i in cart if i["id"] == item["id"]), None)
    if found:
        found["qty"] += 1
    else:
        cart.append({"id": item["id"], "name": item["name"], "price": item["price"], "qty": 1})
    await state.update_data(cart=cart)
    qty = next(i["qty"] for i in cart if i["id"] == item["id"])
    await cb.answer(f"«{item['name']}» × {qty} — в корзине ({_cart_count(cart)} шт.)")
    # сразу обновляем счётчик на кнопке «🛒 Корзина (N)», не дожидаясь клика
    cat_id = _category_of_item(item["id"])
    if cat_id:
        try:
            await cb.message.edit_reply_markup(
                reply_markup=kb_category(cat_id, _cart_count(cart))
            )
        except TelegramBadRequest:
            pass


# ------------------------------------------------------- корзина --------------
async def _render_cart(cb: CallbackQuery, state: FSMContext) -> None:
    cart = await _cart(state)
    table = await _table(state)
    if not cart:
        await cb.message.edit_text(CART_EMPTY, reply_markup=kb_categories(0))
        return
    try:
        await cb.message.edit_text(cart_text(cart, table), reply_markup=kb_cart(cart))
    except TelegramBadRequest:
        pass  # сообщение не изменилось


def _cart_index(cart: list[dict], data: str) -> int | None:
    try:
        idx = int(data.split(":", 1)[1])
    except (ValueError, IndexError):
        return None
    return idx if 0 <= idx < len(cart) else None


@router.callback_query(F.data == "noop")
async def cb_noop(cb: CallbackQuery):
    await cb.answer()


@router.callback_query(F.data == "cv")
async def cb_cart(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    await _render_cart(cb, state)


@router.callback_query(F.data.startswith("ci:"))
async def cb_cart_inc(cb: CallbackQuery, state: FSMContext):
    cart = await _cart(state)
    idx = _cart_index(cart, cb.data)
    if idx is None:
        await cb.answer(ERR_MENU, show_alert=True)
        return
    if not _is_admin(cb.from_user.id) and _cart_count(cart) >= ORDER_MAX_ITEMS:
        await cb.answer(f"⚠️ В одном заказе не больше {ORDER_MAX_ITEMS} позиций", show_alert=True)
        return
    limit = max_available(cart[idx]["id"]) if cart[idx].get("id") else None
    if limit is not None and cart[idx]["qty"] + 1 > limit:
        await cb.answer(
            f"⚠️ «{cart[idx]['name']}»: по остаткам можно не более {limit} шт.",
            show_alert=True,
        )
        return
    cart[idx]["qty"] += 1
    await state.update_data(cart=cart)
    await cb.answer(f"{cart[idx]['name']}: {cart[idx]['qty']} шт.")
    await _render_cart(cb, state)


@router.callback_query(F.data.startswith("cd:"))
async def cb_cart_dec(cb: CallbackQuery, state: FSMContext):
    cart = await _cart(state)
    idx = _cart_index(cart, cb.data)
    if idx is None:
        await cb.answer(ERR_MENU, show_alert=True)
        return
    cart[idx]["qty"] -= 1
    if cart[idx]["qty"] <= 0:
        cart.pop(idx)
    await state.update_data(cart=cart)
    await cb.answer("Уменьшено")
    await _render_cart(cb, state)


@router.callback_query(F.data.startswith("cr:"))
async def cb_cart_remove(cb: CallbackQuery, state: FSMContext):
    cart = await _cart(state)
    idx = _cart_index(cart, cb.data)
    if idx is None:
        await cb.answer(ERR_MENU, show_alert=True)
        return
    name = cart[idx]["name"]
    cart.pop(idx)
    await state.update_data(cart=cart)
    await cb.answer(f"«{name}» убрано")
    await _render_cart(cb, state)


@router.callback_query(F.data == "cx")
async def cb_cart_clear(cb: CallbackQuery, state: FSMContext):
    await state.update_data(cart=[])
    await cb.answer("Корзина очищена")
    await _render_cart(cb, state)


# -------------------------------------------------- оформление заказа ---------
@router.callback_query(F.data == "co")
async def cb_checkout(cb: CallbackQuery, state: FSMContext):
    cart = await _cart(state)
    if not cart:
        await cb.answer(CART_EMPTY, show_alert=True)
        return
    table = await _table(state)
    if not table:
        await cb.answer()
        await state.set_state(OrderFSM.table)
        await cb.message.edit_text(ASK_TABLE)
        return
    await cb.answer()
    await cb.message.edit_text(confirm_text(cart, table), reply_markup=kb_confirm())


@router.callback_query(F.data == "cb")
async def cb_back_to_cart(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    await _render_cart(cb, state)


_user_locks: dict[int, asyncio.Lock] = {}


def _user_lock(uid: int) -> asyncio.Lock:
    """Замок на гостя: двойной клик «Отправить» не создаст два заказа/брони."""
    lk = _user_locks.get(uid)
    if lk is None:
        lk = _user_locks[uid] = asyncio.Lock()
    return lk


def _order_limits_error(uid: int, cart: list[dict]) -> str | None:
    """Анти-спам лимиты заказа (админы из ADMIN_IDS не ограничены)."""
    if _is_admin(uid):
        return None
    if _cart_count(cart) > ORDER_MAX_ITEMS:
        return f"⚠️ В одном заказе не больше {ORDER_MAX_ITEMS} позиций — уменьшите количество."
    age = db_last_order_age_sec(uid)
    if age is not None and 0 <= age < ORDER_COOLDOWN_SEC:
        return (
            f"⏳ Следующий заказ можно отправить через {int(ORDER_COOLDOWN_SEC - age) + 1} сек. "
            "Это защита от случайных повторов."
        )
    return None


async def _submit_order(cb: CallbackQuery, state: FSMContext, kind: str) -> None:
    """Отправка заказа в канал — обычного (корзина cart) или по акции (pcart)."""
    key = "pcart" if kind == "promo" else "cart"
    uid = cb.from_user.id
    lock = _user_lock(uid)
    if lock.locked():
        await cb.answer("⏳ Уже отправляем заказ…")
        return
    async with lock:
        data = await state.get_data()
        cart = list(data.get(key) or [])
        table = data.get("table")
        render = _render_pcart if kind == "promo" else _render_cart
        if not cart:
            await cb.answer(CART_EMPTY, show_alert=True)
            return
        if not table:
            await cb.answer()
            await state.set_state(OrderFSM.table)
            await cb.message.edit_text(ASK_TABLE)
            return
        if kind == "promo":
            cart, bad = _promo_revalidate(cart)
            if bad:
                await state.update_data(pcart=cart)
                await cb.answer("⚠️ Часть акций сейчас недоступна", show_alert=True)
                await cb.message.answer("⚠️ <b>Убрали из заказа:</b>\n" + "\n".join(bad))
                await render(cb, state)
                return
        err = _order_limits_error(uid, cart)
        if err:
            await cb.answer(err, show_alert=True)
            return
        # финальная проверка остатков по всей корзине — заказ сверх склада не уйдёт
        problems = cart_stock_problems(cart)
        if problems:
            await cb.answer("⚠️ Не хватает на складе", show_alert=True)
            await cb.message.answer(
                "⚠️ <b>Не хватает на складе:</b>\n"
                + "\n".join(problems)
                + "\n\nУменьшите количество в корзине или выберите другой товар."
            )
            await render(cb, state)
            return

        await cb.answer("Отправляем…")
        total = _cart_total(cart)
        user = cb.from_user
        card = guest_card(uid)
        order_id = db_add_order(
            table_no=table,
            total=total,
            user_id=uid,
            username=user.username if user else None,
            phone=card["phone"] if card else data.get("phone"),
            kind=kind,
        )
        db_add_items(order_id, cart)

        order = db_get_order(order_id)
        try:
            sent = await cb.bot.send_message(
                CHANNEL_ID,
                order_text(order),
                reply_markup=kb_status(order_id),
            )
        except Exception:
            logger.exception("Не удалось отправить заказ #%s в канал (CHANNEL_ID=%s)",
                             order_id, CHANNEL_ID)
            db_delete_order(order_id)
            await cb.message.edit_text(
                "⚠️ Не удалось отправить заказ в канал. Сообщите администратору.",
                reply_markup=kb_pconfirm() if kind == "promo" else kb_confirm(),
            )
            return
        db_set_message_id(order_id, sent.message_id)
        await state.update_data(**{key: []})
        db_guest_touch(user)
        text = sent_text(order_id, table, total)
        if kind == "promo":
            text = "🎁 Заказ по акции.\n" + text
        elif card and LOYALTY_PERCENT:
            text += f"\n\n💎 После оплаты начислим кешбэк {LOYALTY_PERCENT}% баллами."
        await cb.message.edit_text(text)


@router.callback_query(F.data == "so")
async def cb_send_order(cb: CallbackQuery, state: FSMContext):
    await _submit_order(cb, state, "regular")


# ============================== БРОНИРОВАНИЕ ================================
# Кнопка «📅 Бронирование» → тип → дата → время → гости → подтверждение →
# сообщение в канал со статусами (редактируется тем же сообщением).
# Вместимость, часы работы, длительности и цены аренды — в НАСТРОЙКИ → «Бронирование».


def _fmt_minutes(m: int) -> str:
    h, mi = divmod(int(m), 60)
    if h and mi:
        return f"{h} ч {mi} мин"
    if h:
        return f"{h} ч"
    return f"{mi} мин"


def _rsv_actual_min(r: dict) -> int:
    if r.get("actual_min"):
        return int(r["actual_min"])
    return RSV_BOOKED_MIN.get(r.get("duration") or "", 0)


def _rsv_price(r: dict) -> int | None:
    """Цена брони по факту времени. None — виды без почасовой оплаты (стол, бар)."""
    code = r.get("duration") or ""
    base = RSV_BASE_PRICE.get(code)
    if base is None or r.get("kind") not in RSV_ADJ_KINDS:
        return None
    steps = (_rsv_actual_min(r) - RSV_BOOKED_MIN[code]) // RSV_STEP_MIN
    return max(0, base + steps * RSV_STEP_PRICE[r["kind"]])


def _slots() -> list[str]:
    """Слоты времени: 15:00 → 02:30 с шагом 30 минут (работаем 15:00–03:00)."""
    out: list[str] = []
    for h in (15, 16, 17, 18, 19, 20, 21, 22, 23, 0, 1, 2):
        out.append(f"{h:02d}:00")
        out.append(f"{h:02d}:30")
    return out


def _arrival(date_s: str, time_s: str) -> datetime:
    """Фактическое время прихода: ночные слоты (00:00–03:00) — это утро
    СЛЕДУЮЩЕГО дня после выбранной даты (работаем до 03:00)."""
    d = datetime.strptime(date_s, "%Y-%m-%d").date()
    h, m = map(int, time_s.split(":"))
    if h < 15:
        d = d + timedelta(days=1)
    return datetime(d.year, d.month, d.day, h, m)


def _wd(d) -> str:
    return WDAYS[d.weekday()]


def _rsv_when_text(date_s: str, time_s: str) -> str:
    d0 = datetime.strptime(date_s, "%Y-%m-%d").date()
    d1 = _arrival(date_s, time_s).date()
    if d1 != d0:
        return f"{_wd(d0)} {d0:%d.%m} → {_wd(d1)} {d1:%d.%m} · {time_s}"
    return f"{_wd(d0)} {d0:%d.%m} · {time_s}"


def _guests_word(n: int) -> str:
    if n % 10 == 1 and n % 100 != 11:
        return "гость"
    if 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        return "гостя"
    return "гостей"


def _biz_today():
    """Текущий рабочий день: до 03:00 ещё идёт вчерашняя смена (15:00–03:00).
    В 01:00 вторника это понедельник — его ночные слоты ещё можно бронировать."""
    return (_now_dt() - timedelta(hours=RSV_CLOSE_H)).date()


def _date_label(d) -> str:
    today = _now_dt().date()
    if d == _biz_today() and d != today:
        return f"🌙 {d:%d.%m} (эта ночь)"
    if d == today:
        return f"📅 {d:%d.%m} (сегодня)"
    if d == today + timedelta(days=1):
        return f"📅 {d:%d.%m} (завтра)"
    return f"📅 {d:%d.%m} {_wd(d)}"


def rsv_text(r: dict) -> str:
    lines = [
        f"📅 <b>Бронь #{r['id']}</b>",
        RSV_KINDS.get(r["kind"], r["kind"]),
        f"📆 {_rsv_when_text(r['date'], r['time'])}",
        f"👥 {r['guests']} {_guests_word(r['guests'])}",
    ]
    if r.get("duration"):
        lines.append(f"⏱ {RSV_DURATIONS.get(r['duration'], r['duration'])}")
    try:
        t0, t1 = rsv_interval(r["kind"], r["date"], r["time"], r.get("duration") or "",
                              r.get("actual_min"))
        approx = " (расчётно)" if r["kind"] not in RSV_ADJ_KINDS else ""
        lines.append(f"🕒 {t0:%H:%M}–{t1:%H:%M}{approx}")
    except (ValueError, TypeError, KeyError):
        pass
    price = _rsv_price(r)
    if price is not None:
        if r["status"] == "paid":
            lines.append(
                f"💳 Оплачено: {_fmt_money(price)}"
                f" · {PAY_METHODS.get(r.get('pay_method') or '', '—')}"
            )
        else:
            fact = _fmt_minutes(_rsv_actual_min(r))
            booked = RSV_BOOKED_MIN.get(r.get("duration") or "")
            if booked and _rsv_actual_min(r) != booked:
                lines.append(f"🧾 К оплате: {_fmt_money(price)} (факт {fact})")
            else:
                lines.append(f"🧾 К оплате: {_fmt_money(price)}")
    name = r["guest_name"]
    if r.get("username") and "(@" not in str(name):
        name = f"{name} (@{r['username']})"
    lines += [f"👤 {esc(name)}"]
    if r.get("phone"):
        lines.append(f"📞 {r['phone']}")
    lines += [""]
    status = RSV_STATUSES.get(r["status"], r["status"])
    ts = (r.get("updated_at") or "")[11:16]
    lines.append(f"Статус: {status} · {ts}")
    return "\n".join(lines)


def _rsv_notify_text(r: dict) -> str:
    """Сообщение гостю о смене статуса его брони."""
    base = (
        f"📅 Бронь <b>#{r['id']}</b>\n"
        f"{RSV_KINDS.get(r['kind'], r['kind'])} · "
        f"{_rsv_when_text(r['date'], r['time'])}"
    )
    if r["status"] == "confirmed":
        return (
            f"✅ <b>Ваша бронь подтверждена!</b>\n\n{base}\n\n"
            "Ждём вас! Если планы изменятся — сообщите заранее."
        )
    if r["status"] == "paid":
        return (
            f"💳 <b>Бронь #{r['id']} оплачена</b>\n\n{base}\n\n"
            "Спасибо, что выбрали нас! Ждём снова ⚡"
        )
    if r["status"] == "cancelled":
        return (
            f"❌ <b>Ваша бронь отменена</b>\n\n{base}\n\n"
            "Если что-то изменилось — забронируйте снова через /start."
        )
    if r["status"] == "no_show":
        return (
            f"🚫 <b>Вы не пришли на бронь</b>\n\n{base}\n\n"
            "Если захотите снова — забронируйте через /start."
        )
    return (
        f"⏳ <b>Ваша бронь #{r['id']}</b> снова в статусе «Ожидает».\n"
        "Администратор вернётся к ней."
    )


# --- клавиатуры брони ---
def kb_rsv_kind() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [_btn(RSV_KINDS["table"], "rsv:kind:table")],
            [_btn(RSV_KINDS["bar"], "rsv:kind:bar")],
            [_btn(RSV_KINDS["vip"], "rsv:kind:vip")],
            [_btn(RSV_KINDS["ps"], "rsv:kind:ps")],
            [_btn("🏠 В начало", "go:home")],
        ]
    )


def kb_rsv_dates() -> InlineKeyboardMarkup:
    today = _biz_today()
    dates = [today + timedelta(days=i) for i in range(RSV_DAYS)]
    rows: list[list[InlineKeyboardButton]] = []
    for i in range(0, len(dates), 3):
        row = [_btn(_date_label(d), f"rsv:date:{d.isoformat()}") for d in dates[i : i + 3]]
        rows.append(row)
    rows.append([_btn("🔙 Назад", "rsv:back:kind")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_rsv_times(avail: list[str]) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    for i in range(0, len(avail), 4):
        rows.append([_btn(s, f"rsv:time:{s}") for s in avail[i : i + 4]])
    rows.append([_btn("🔙 Другая дата", "rsv:back:date")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_rsv_guests(kind: str | None = None) -> InlineKeyboardMarkup:
    if kind == "bar":
        return InlineKeyboardMarkup(
            inline_keyboard=[
                [_btn(str(n), f"rsv:g:{n}") for n in range(1, 6)],
                [_btn("🔙 Назад", "rsv:back:time")],
            ]
        )
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [_btn(str(n), f"rsv:g:{n}") for n in range(1, 6)],
            [_btn(str(n), f"rsv:g:{n}") for n in range(6, 11)],
            [_btn("🔙 Назад", "rsv:back:time")],
        ]
    )


def kb_rsv_duration(kind: str = "table", codes: list[str] | None = None) -> InlineKeyboardMarkup:
    """codes — доступные варианты (помещаются до закрытия и свободны)."""
    if codes is None:
        codes = RSV_DUR_BY_KIND.get(kind, RSV_DUR_BY_KIND["table"])
    btns = [_btn(RSV_DURATIONS[c], f"rsv:dur:{c}") for c in codes]
    rows = [btns[i : i + 2] for i in range(0, len(btns), 2)]
    rows.append([_btn("🔙 Назад", "rsv:back:name")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def _rsv_dur_screen(rsv: dict) -> tuple[str, InlineKeyboardMarkup]:
    """Экран выбора длительности: только варианты, которые помещаются до
    закрытия и свободны. Если всё заняли — предлагаем другое время."""
    kind = rsv.get("kind") or "table"
    codes = rsv_options(kind, rsv["date"], rsv["time"])
    if not codes:
        return (
            "😢 Пока вы заполняли форму, это время заняли.\nВыберите другое время 👇",
            InlineKeyboardMarkup(inline_keyboard=[[_btn("🔙 Другое время", "rsv:back:time")]]),
        )
    text = _rsv_dur_text(rsv)
    if len(codes) < len(RSV_DUR_BY_KIND.get(kind, [])):
        text += "\n\nℹ️ Показаны варианты, которые свободны и заканчиваются до закрытия (03:00)."
    return text, kb_rsv_duration(kind, codes)


def kb_rsv_name() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[[_btn("🔙 Назад", "rsv:back:guests")]]
    )


def kb_rsv_confirm(rsv: dict | None = None) -> InlineKeyboardMarkup:
    # у барной стойки часов нет — «Назад» возвращает к вводу имени
    back = "rsv:back:name" if (rsv or {}).get("kind") == "bar" else "rsv:back:dur"
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [_btn("✅ Отправить бронь", "rsv:ok")],
            [
                _btn("🔙 Назад", back),
                _btn("⏹ Отменить", "rsv:no"),
            ],
        ]
    )


def kb_rsv_status(rid: int, r: dict | None = None) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    if r and r.get("kind") in RSV_ADJ_KINDS and r.get("status") in ("new", "confirmed"):
        rows.append([
            _btn("➖ 15 мин", f"rsa:{rid}:-15"),
            _btn("➕ 15 мин", f"rsa:{rid}:+15"),
        ])
        rows.append([_btn("💳 Закрыть чек", f"rsp:{rid}")])
    if r and r.get("status") == "paid":
        rows.append([_btn("↩️ Оформить возврат", f"rs:{rid}:cancelled")])
    else:
        rows.append([
            _btn(RSV_STATUSES["confirmed"], f"rs:{rid}:confirmed"),
            _btn(RSV_STATUSES["cancelled"], f"rs:{rid}:cancelled"),
        ])
        rows.append([
            _btn(RSV_STATUSES["no_show"], f"rs:{rid}:no_show"),
            _btn(RSV_STATUSES["new"], f"rs:{rid}:new"),
        ])
    return InlineKeyboardMarkup(inline_keyboard=rows)


# --- экраны ---
async def _rsv_show_dates(target, edit: bool = False):
    text = "📅 <b>Бронирование</b>\n\nНа какую дату?"
    kb = kb_rsv_dates()
    if edit:
        try:
            await target.edit_text(text, reply_markup=kb)
        except TelegramBadRequest:
            pass
    else:
        await target.answer(text, reply_markup=kb)


_RSV_TIME_HINTS = {
    "vip": "Показано время, когда комната свободна и аренда заканчивается до 03:00 "
           "(между бронями — 30 мин на уборку).",
    "ps": "Показано время, когда есть свободная приставка и аренда заканчивается до 03:00.",
}


async def _rsv_show_times(target, date_s: str, kind: str | None = None, note: str = ""):
    if kind:
        avail = rsv_slots_available(kind, date_s)
    else:
        now = _now_dt() + timedelta(minutes=RSV_LEAD_MIN)
        avail = [s for s in _slots() if _arrival(date_s, s) > now]
    d_obj = datetime.strptime(date_s, "%Y-%m-%d").date()
    text = (
        f"{note}⏰ Выберите время — {_date_label(d_obj)}\n"
        "Кальянная работает <b>с 15:00 до 03:00</b>"
    )
    if kind in _RSV_TIME_HINTS and avail:
        text += f"\nℹ️ {_RSV_TIME_HINTS[kind]}"
    if not avail:
        text += "\n\n😢 На эту дату свободного времени уже нет — выберите другую."
        kb = InlineKeyboardMarkup(
            inline_keyboard=[[_btn("🔙 Другая дата", "rsv:back:date")]]
        )
    else:
        kb = kb_rsv_times(avail)
    try:
        await target.edit_text(text, reply_markup=kb)
    except TelegramBadRequest:
        pass


def _rsv_summary_lines(rsv: dict) -> list[str]:
    lines = []
    if rsv.get("kind"):
        lines.append(f"📌 {RSV_KINDS.get(rsv['kind'], rsv['kind'])}")
    if rsv.get("date") and rsv.get("time"):
        lines.append(f"📆 {_rsv_when_text(rsv['date'], rsv['time'])}")
    if rsv.get("guests"):
        lines.append(f"👥 {rsv['guests']} {_guests_word(rsv['guests'])}")
    return lines


def _rsv_guests_text(rsv: dict) -> str:
    return "\n".join(["👥 Сколько гостей?", ""] + _rsv_summary_lines(rsv))


def _rsv_dur_text(rsv: dict) -> str:
    lines = ["⏱ На сколько часов бронируем?", ""] + _rsv_summary_lines(rsv)
    cap = RSV_DUR_CAPTIONS.get(rsv.get("kind") or "")
    if cap:
        lines += ["", f"ℹ️ {cap}"]
    return "\n".join(lines)


def _rsv_name_text(rsv: dict) -> str:
    return "\n".join(
        ["✍️ Имя, на кого бронь?", ""]
        + _rsv_summary_lines(rsv)
        + ["", "Напишите имя обычным сообщением ↩️"]
    )


def _rsv_confirm_text(rsv: dict, name: str) -> str:
    lines = [
        "✅ <b>Подтверждение брони</b>",
        "",
        f"📌 {RSV_KINDS.get(rsv.get('kind'), '—')}",
        f"📆 {_rsv_when_text(rsv['date'], rsv['time'])}",
        f"👥 {rsv['guests']} {_guests_word(rsv['guests'])}",
    ]
    if rsv.get("duration"):
        lines.append(f"⏱ {RSV_DURATIONS.get(rsv['duration'], '—')}")
        if rsv.get("kind") in RSV_ADJ_KINDS:
            try:
                t0, t1 = rsv_interval(rsv["kind"], rsv["date"], rsv["time"], rsv["duration"])
                lines.append(f"🕒 {t0:%H:%M}–{t1:%H:%M}")
            except (ValueError, KeyError):
                pass
    lines += [
        f"👤 {esc(name)}",
        "",
        "Отправить бронь администратору?",
    ]
    return "\n".join(lines)


# --- хендлеры ---
@router.message(F.text == BTN_BOOK)
async def on_btn_book(message: Message, state: FSMContext):
    await _begin_flow("rsv", message, state)


@router.callback_query(F.data.startswith("rsv:kind:"))
async def rsv_kind(cb: CallbackQuery, state: FSMContext):
    kind = cb.data.split(":", 2)[2]
    if kind not in RSV_KINDS:
        await cb.answer(ERR_MENU, show_alert=True)
        return
    if not _is_admin(cb.from_user.id) and db_rsv_user_active(cb.from_user.id) >= RSV_MAX_ACTIVE_PER_USER:
        await cb.answer(_RSV_LIMIT_TEXT, show_alert=True)
        return
    await state.update_data(rsv={"kind": kind})
    await cb.answer()
    await _rsv_show_dates(cb.message, edit=True)


@router.callback_query(F.data.startswith("rsv:date:"))
async def rsv_date(cb: CallbackQuery, state: FSMContext):
    date_s = cb.data.split(":", 2)[2]
    try:
        d_obj = datetime.strptime(date_s, "%Y-%m-%d").date()
    except ValueError:
        await cb.answer("Некорректная дата", show_alert=True)
        return
    if not _biz_today() <= d_obj < _biz_today() + timedelta(days=RSV_DAYS):
        await cb.answer("Эта дата недоступна для брони", show_alert=True)
        return
    data = await state.get_data()
    rsv = dict(data.get("rsv") or {})
    if rsv.get("kind") is None:
        await cb.answer("Начните бронирование заново", show_alert=True)
        await cb.message.edit_text(
            "📅 <b>Бронирование</b>\n\nЧто будем бронировать?",
            reply_markup=kb_rsv_kind(),
        )
        return
    # свободное время на эту дату (интервалы, закрытие 03:00, уже прошедшее)
    if rsv_capacity_full(rsv["kind"], date_s):
        await cb.answer(
            f"На {_date_label(d_obj)[2:]} для «{RSV_KINDS[rsv['kind']]}» "
            "свободного времени нет. Выберите другую дату",
            show_alert=True,
        )
        return
    rsv["date"] = date_s
    await state.update_data(rsv=rsv)
    await cb.answer()
    await _rsv_show_times(cb.message, date_s, rsv["kind"])


@router.callback_query(F.data.startswith("rsv:time:"))
async def rsv_time(cb: CallbackQuery, state: FSMContext):
    time_s = cb.data.split(":", 2)[2]
    if time_s not in _slots():
        await cb.answer("Некорректное время", show_alert=True)
        return
    data = await state.get_data()
    rsv = dict(data.get("rsv") or {})
    if not rsv.get("kind") or not rsv.get("date"):
        await cb.answer("Начните бронирование заново", show_alert=True)
        await cb.message.edit_text(
            "📅 <b>Бронирование</b>\n\nЧто будем бронировать?",
            reply_markup=kb_rsv_kind(),
        )
        return
    if time_s not in rsv_slots_available(rsv["kind"], rsv["date"]):
        await cb.answer("Это время уже занято или недоступно — выберите другое", show_alert=True)
        await _rsv_show_times(cb.message, rsv["date"], rsv["kind"])
        return
    rsv["time"] = time_s
    rsv.pop("duration", None)
    await state.update_data(rsv=rsv)
    await cb.answer()
    await cb.message.edit_text(
        _rsv_guests_text(rsv),
        reply_markup=kb_rsv_guests(rsv.get("kind")),
    )


@router.callback_query(F.data.startswith("rsv:g:"))
async def rsv_guests(cb: CallbackQuery, state: FSMContext):
    try:
        n = int(cb.data.split(":", 2)[2])
    except ValueError:
        n = 0
    data = await state.get_data()
    rsv = dict(data.get("rsv") or {})
    gmax = 5 if rsv.get("kind") == "bar" else 10
    if not 1 <= n <= gmax:
        await cb.answer(
            "За барной стойкой — от 1 до 5 мест" if gmax == 5 else "От 1 до 10 гостей",
            show_alert=True,
        )
        return
    if not (rsv.get("kind") and rsv.get("date") and rsv.get("time")):
        await cb.answer("Начните бронирование заново", show_alert=True)
        await cb.message.edit_text(
            "📅 <b>Бронирование</b>\n\nЧто будем бронировать?",
            reply_markup=kb_rsv_kind(),
        )
        return
    rsv["guests"] = n
    await state.update_data(rsv=rsv)
    # следующий шаг — имя, на кого бронь
    await state.set_state(OrderFSM.rsv_name)
    await cb.answer()
    try:
        await cb.message.edit_text(_rsv_name_text(rsv), reply_markup=kb_rsv_name())
        await state.update_data(rsv_name_msg=cb.message.message_id)
    except TelegramBadRequest:
        pass


@router.callback_query(F.data.startswith("rsv:dur:"))
async def rsv_duration(cb: CallbackQuery, state: FSMContext):
    code = cb.data.split(":", 2)[2]
    if code not in RSV_DURATIONS:
        await cb.answer("Некорректная длительность", show_alert=True)
        return
    data = await state.get_data()
    rsv = dict(data.get("rsv") or {})
    if not (
        rsv.get("kind")
        and rsv.get("date")
        and rsv.get("time")
        and rsv.get("guests")
        and rsv.get("name")
    ):
        await cb.answer("Начните бронирование заново", show_alert=True)
        await cb.message.edit_text(
            "📅 <b>Бронирование</b>\n\nЧто будем бронировать?",
            reply_markup=kb_rsv_kind(),
        )
        return
    allowed = RSV_DUR_BY_KIND.get(rsv.get("kind") or "", [])
    if code not in allowed:
        await cb.answer("Выберите длительность из списка", show_alert=True)
        return
    if code not in rsv_options(rsv["kind"], rsv["date"], rsv["time"]):
        await cb.answer(
            "Этот вариант недоступен: не помещается до закрытия (03:00) или уже занят",
            show_alert=True,
        )
        text, kb = _rsv_dur_screen(rsv)
        try:
            await cb.message.edit_text(text, reply_markup=kb)
        except TelegramBadRequest:
            pass
        return
    rsv["duration"] = code
    await state.update_data(rsv=rsv)
    await state.set_state(None)
    await cb.answer()
    nm = rsv["name"]
    if cb.from_user and cb.from_user.username and "(@" not in nm:
        nm += f" (@{cb.from_user.username})"
    await cb.message.edit_text(
        _rsv_confirm_text(rsv, nm),
        reply_markup=kb_rsv_confirm(rsv),
    )


_RSV_LIMIT_TEXT = (
    f"⚠️ У вас уже {RSV_MAX_ACTIVE_PER_USER} активные брони — это максимум. "
    "Дождитесь визита или попросите администратора отменить лишнюю."
)


@router.callback_query(F.data == "rsv:ok")
async def rsv_send(cb: CallbackQuery, state: FSMContext):
    lock = _user_lock(cb.from_user.id)
    if lock.locked():
        await cb.answer("⏳ Уже отправляем бронь…")
        return
    async with lock:
        await _rsv_send_locked(cb, state)


async def _rsv_send_locked(cb: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    rsv = dict(data.get("rsv") or {})
    need_dur = rsv.get("kind") != "bar"  # барная стойка — место без часов
    if not (
        rsv.get("kind")
        and rsv.get("date")
        and rsv.get("time")
        and rsv.get("guests")
        and rsv.get("name")
        and (rsv.get("duration") or not need_dur)
    ):
        await cb.answer("Данные брони неполные — начните заново", show_alert=True)
        await cb.message.edit_text(
            "📅 <b>Бронирование</b>\n\nЧто будем бронировать?",
            reply_markup=kb_rsv_kind(),
        )
        return

    user = cb.from_user
    card = guest_card(user.id)
    # проверка лимитов и запись — одной операцией (без гонки за последнее место)
    rid, why = db_add_reservation_checked(
        None if _is_admin(user.id) else RSV_MAX_ACTIVE_PER_USER,
        kind=rsv["kind"],
        date_s=rsv["date"],
        time_s=rsv["time"],
        guests=rsv["guests"],
        duration=rsv.get("duration") or "",
        guest_name=rsv["name"],
        user_id=user.id,
        username=user.username,
        phone=card["phone"] if card else data.get("phone"),
    )
    if rid is None and why == "user_limit":
        await state.update_data(rsv={})
        await cb.answer(_RSV_LIMIT_TEXT, show_alert=True)
        try:
            await cb.message.edit_text("⚠️ <b>Бронь не отправлена</b>\n\n" + _RSV_LIMIT_TEXT)
        except TelegramBadRequest:
            pass
        return
    if rid is None:
        # пока заполняли форму, это время заняли (или оно уже прошло)
        rsv.pop("time", None)
        rsv.pop("duration", None)
        await state.update_data(rsv=rsv)
        await cb.answer()
        await _rsv_show_times(
            cb.message, rsv["date"], rsv["kind"],
            note="⚠️ <b>Бронь не отправлена</b> — пока заполняли форму, это время "
                 "заняли или оно уже прошло. Выберите другое 👇\n\n",
        )
        return

    await cb.answer("Отправляем…")
    r = db_get_reservation(rid)
    try:
        sent = await cb.bot.send_message(
            CHANNEL_ID,
            rsv_text(r),
            reply_markup=kb_rsv_status(rid, r),
        )
    except Exception:
        logger.exception(
            "Не удалось отправить бронь #%s в канал (CHANNEL_ID=%s)", rid, CHANNEL_ID
        )
        db_delete_reservation(rid)
        # cb.answer уже был («Отправляем…») — второй ответ Telegram отклонит,
        # поэтому сообщаем об ошибке в самом сообщении (кнопки остаются)
        try:
            await cb.message.edit_text(
                "⚠️ Не удалось отправить бронь администратору. "
                "Попробуйте ещё раз или сообщите персоналу.",
                reply_markup=kb_rsv_confirm(rsv),
            )
        except TelegramBadRequest:
            pass
        return
    db_set_rsv_message_id(rid, sent.message_id)
    await state.update_data(rsv={})
    await cb.message.edit_text(
        f"✅ Бронь <b>#{rid}</b> отправлена!\n\n"
        f"📌 {RSV_KINDS.get(rsv['kind'], '')}\n"
        f"📆 {_rsv_when_text(rsv['date'], rsv['time'])}\n"
        f"👥 {rsv['guests']} {_guests_word(rsv['guests'])}\n"
        + (f"⏱ {RSV_DURATIONS.get(rsv['duration'], '')}\n" if rsv.get("duration") else "")
        + "\n"
        "Администратор подтвердит бронь — статус сообщим здесь. Спасибо!"
    )


@router.callback_query(F.data.startswith("rsv:back:"))
async def rsv_back(cb: CallbackQuery, state: FSMContext):
    step = cb.data.split(":", 2)[2]
    data = await state.get_data()
    rsv = dict(data.get("rsv") or {})
    await cb.answer()
    await state.set_state(None)
    if step == "kind":
        await cb.message.edit_text(
            "📅 <b>Бронирование</b>\n\nЧто будем бронировать?",
            reply_markup=kb_rsv_kind(),
        )
    elif step == "date" and rsv.get("kind"):
        await _rsv_show_dates(cb.message, edit=True)
    elif step == "time" and rsv.get("date"):
        await _rsv_show_times(cb.message, rsv["date"], rsv.get("kind"))
    elif step == "guests" and rsv.get("time"):
        await cb.message.edit_text(
            _rsv_guests_text(rsv),
            reply_markup=kb_rsv_guests(rsv.get("kind")),
        )
    elif step == "name" and rsv.get("guests"):
        await state.set_state(OrderFSM.rsv_name)
        try:
            await cb.message.edit_text(_rsv_name_text(rsv), reply_markup=kb_rsv_name())
            await state.update_data(rsv_name_msg=cb.message.message_id)
        except TelegramBadRequest:
            pass
    elif step == "dur" and rsv.get("name") and rsv.get("time"):
        text, kb = _rsv_dur_screen(rsv)
        await cb.message.edit_text(text, reply_markup=kb)
    else:
        await cb.message.edit_text(
            "📅 <b>Бронирование</b>\n\nЧто будем бронировать?",
            reply_markup=kb_rsv_kind(),
        )


@router.callback_query(F.data == "rsv:no")
async def rsv_cancel(cb: CallbackQuery, state: FSMContext):
    await state.update_data(rsv={})
    await state.set_state(None)
    await cb.answer("Отменено")
    try:
        await cb.message.edit_text("Бронирование отменено ✅")
    except TelegramBadRequest:
        pass


@router.callback_query(F.data.startswith("rs:"))
async def cb_rsv_status(cb: CallbackQuery):
    try:
        _, rid_raw, code = cb.data.split(":")
        rid = _cb_id(rid_raw)
    except (ValueError, AttributeError):
        await cb.answer("Некорректный запрос", show_alert=True)
        return

    if code == "paid":
        await cb.answer("Оплату оформляйте кнопкой «💳 Закрыть чек»", show_alert=True)
        return
    if code not in RSV_STATUSES:
        await cb.answer("Неизвестный статус", show_alert=True)
        return

    if not await _staff_guard(cb, "⚠️ Статус меняют только администраторы канала"):
        return

    r = db_get_reservation(rid)
    if not r:
        await cb.answer("Бронь не найдена", show_alert=True)
        return

    label = RSV_STATUSES[code]
    if r["status"] == code:
        await cb.answer(f"Уже: {label}")
        return

    # Финальные статусы «заперты»: из «Отменена» / «Не пришёл»
    # активные кнопки не работают — вернуть бронь можно только через «Ожидает».
    if r["status"] in ("cancelled", "no_show") and code != "new":
        await cb.answer(
            f"🔒 Бронь: {RSV_STATUSES[r['status']]} — активные статусы не ставятся. "
            "Возврат — только «🟡 Ожидает»",
            show_alert=True,
        )
        return

    if r["status"] == "paid" and code != "cancelled":
        await cb.answer(
            "🔒 Бронь оплачена — активные статусы не ставятся. "
            "Возврат — только «❌ Отменена»",
            show_alert=True,
        )
        return

    # возврат отменённой брони в работу — предупредим, если её время уже заняли
    warn = ""
    if r["status"] in ("cancelled", "no_show") and code in ("new", "confirmed"):
        try:
            if not _rsv_free_in(db_rsv_active_rows(r["kind"], r["date"], exclude_id=rid),
                                r["kind"], r["date"], r["time"], r.get("duration") or "",
                                r.get("actual_min")):
                warn = "\n⚠️ На это время мест уже нет — проверьте, не будет ли накладки"
        except (ValueError, TypeError):
            pass

    db_set_rsv_status(rid, code)
    r = db_get_reservation(rid)
    try:
        await cb.message.edit_text(rsv_text(r), reply_markup=kb_rsv_status(rid, r))
    except Exception:
        logger.exception("Не удалось отредактировать бронь #%s", rid)

    # --- уведомление гостю о смене статуса ---
    if r.get("user_id"):
        try:
            await cb.bot.send_message(r["user_id"], _rsv_notify_text(r))
        except Exception:
            logger.info(
                "Не удалось уведомить гостя брони #%s (user_id=%s) — возможно, "
                "он не начинал диалог с ботом или заблокировал его",
                rid,
                r.get("user_id"),
            )
    await cb.answer(f"{label}{warn}", show_alert=bool(warn))


@router.callback_query(F.data.startswith("rsa:"))
async def cb_rsv_adjust(cb: CallbackQuery):
    """Правка фактического времени аренды VIP/PlayStation: ±15 минут."""
    try:
        _, rid_raw, sign = cb.data.split(":")
        rid = _cb_id(rid_raw)
    except (ValueError, AttributeError):
        await cb.answer("Некорректный запрос", show_alert=True)
        return
    if sign not in ("+15", "-15"):
        await cb.answer("Некорректная правка", show_alert=True)
        return
    if not await _staff_guard(cb, "⚠️ Правку выполняют только администраторы канала"):
        return
    r = db_get_reservation(rid)
    if not r:
        await cb.answer("Бронь не найдена", show_alert=True)
        return
    if r["kind"] not in RSV_ADJ_KINDS or r["status"] not in ("new", "confirmed"):
        await cb.answer(
            "🔒 Правка доступна только для активной аренды VIP/PlayStation",
            show_alert=True,
        )
        return
    if not RSV_BOOKED_MIN.get(r.get("duration") or ""):
        await cb.answer("У этой брони нет почасовой оплаты", show_alert=True)
        return
    actual = max(
        RSV_MIN_MIN,
        _rsv_actual_min(r) + (RSV_STEP_MIN if sign == "+15" else -RSV_STEP_MIN),
    )
    rr = dict(r)
    rr["actual_min"] = actual
    price = _rsv_price(rr) or 0
    # продление: предупредим о закрытии и о следующей брони (не блокируем)
    warn = ""
    if sign == "+15":
        try:
            _, t1 = rsv_interval(r["kind"], r["date"], r["time"], r.get("duration") or "", actual)
            if t1 > _rsv_close_dt(r["date"]):
                warn += "\n⚠️ Аренда заканчивается после закрытия (03:00)"
            if not _rsv_free_in(db_rsv_active_rows(r["kind"], r["date"], exclude_id=rid),
                                r["kind"], r["date"], r["time"], r.get("duration") or "", actual):
                warn += "\n⚠️ Пересекается со следующей бронью"
        except (ValueError, TypeError):
            pass
    db_set_rsv_actual(rid, actual, price)
    r = db_get_reservation(rid)
    try:
        await cb.message.edit_text(rsv_text(r), reply_markup=kb_rsv_status(rid, r))
    except Exception:
        logger.exception("Не удалось обновить бронь #%s после правки", rid)
    await cb.answer(f"Факт: {_fmt_minutes(actual)} · к оплате {_fmt_money(price)}{warn}",
                    show_alert=bool(warn))


@router.callback_query(F.data.startswith("rsp:"))
async def cb_rsv_pay(cb: CallbackQuery):
    """Закрытие чека аренды: способ оплаты → статус «Оплачен»."""
    parts = cb.data.split(":")
    try:
        rid = _cb_id(parts[1])
    except (ValueError, IndexError):
        await cb.answer("Некорректный запрос", show_alert=True)
        return
    if not await _staff_guard(cb, "⚠️ Оплату принимают только администраторы канала"):
        return
    r = db_get_reservation(rid)
    if not r:
        await cb.answer("Бронь не найдена", show_alert=True)
        return
    if r["kind"] not in RSV_ADJ_KINDS or r["status"] not in ("new", "confirmed"):
        await cb.answer(
            "🔒 Закрытие чека доступно только для активной аренды VIP/PlayStation",
            show_alert=True,
        )
        return
    if str(r["date"]) > _biz_today().isoformat():
        await cb.answer(
            f"⏳ Аренда ещё не наступила ({_ddmm(r['date'])} в {r['time']}). "
            "Чек закрывают в день аренды — по факту времени.",
            show_alert=True,
        )
        return
    if len(parts) == 2:
        kb = InlineKeyboardMarkup(
            inline_keyboard=[
                [
                    _btn(PAY_BUTTONS["cash"], f"rsp:{rid}:cash"),
                    _btn(PAY_BUTTONS["card"], f"rsp:{rid}:card"),
                ],
                [_btn(PAY_BUTTONS["transfer"], f"rsp:{rid}:transfer")],
                [_btn("🔙 Назад", f"rsp:{rid}:back")],
            ]
        )
        try:
            await cb.message.edit_reply_markup(reply_markup=kb)
        except TelegramBadRequest:
            pass
        await cb.answer(f"К оплате: {_fmt_money(_rsv_price(r) or 0)} — способ?")
        return
    action = parts[2]
    if action == "back":
        try:
            await cb.message.edit_reply_markup(reply_markup=kb_rsv_status(rid, r))
        except TelegramBadRequest:
            pass
        await cb.answer("Отменено")
        return
    if action not in PAY_METHODS:
        await cb.answer("Неизвестный способ", show_alert=True)
        return
    price = _rsv_price(r) or 0
    actual = _rsv_actual_min(r)
    db_set_rsv_paid(rid, action, actual, price)
    r = db_get_reservation(rid)
    try:
        await cb.message.edit_text(rsv_text(r), reply_markup=kb_rsv_status(rid, r))
    except Exception:
        logger.exception("Не удалось обновить бронь #%s после оплаты", rid)
    if r.get("user_id"):
        try:
            await cb.bot.send_message(r["user_id"], _rsv_notify_text(r))
        except Exception:
            logger.info("Не удалось уведомить гостя брони #%s об оплате", rid)
    await cb.answer(f"Оплачено: {_fmt_money(price)} ({PAY_METHODS[action]})")


# ============================ КАРТА ЛОЯЛЬНОСТИ ==============================
def kb_card(has_card: bool) -> InlineKeyboardMarkup | None:
    """Номер карты гость оставляет один раз; сменить его может только админ (/guest)."""
    if has_card:
        return None
    return InlineKeyboardMarkup(inline_keyboard=[[_btn("💳 Подключить карту", "lc:phone")]])


CARD_PHONE_LOCKED = "Номер карты уже привязан. Сменить его может администратор заведения."


async def _send_card(message: Message, state: FSMContext, uid: int) -> None:
    card = guest_card(uid)
    if card:
        await message.answer(card_text(uid), reply_markup=kb_card(True))
        return
    await message.answer(
        "💎 <b>Карта лояльности Zig Zag</b>\n\n"
        "У вас пока нет карты. Оставьте номер телефона — он станет номером карты.\n"
        f"🎁 Кешбэк <b>{LOYALTY_PERCENT}%</b> баллами с каждого оплаченного заказа "
        f"(1 балл = 1 {CURRENCY}), оплатить баллами можно до {BONUS_MAX_PAY_PCT}% чека.",
        reply_markup=kb_card(False),
    )


@router.message(F.text == BTN_CARD)
@router.message(Command("card"))
async def on_my_card(message: Message, state: FSMContext):
    if message.chat.type != "private" or not message.from_user:
        return
    db_guest_touch(message.from_user)
    await _send_card(message, state, message.from_user.id)


@router.callback_query(F.data == "lc:phone")
async def cb_card_phone(cb: CallbackQuery, state: FSMContext):
    if guest_card(cb.from_user.id):   # старая кнопка «Сменить номер» в прошлых сообщениях
        await cb.answer(f"💎 {CARD_PHONE_LOCKED}", show_alert=True)
        return
    await cb.answer()
    await state.update_data(next="card")
    await state.set_state(OrderFSM.phone)
    await cb.message.answer(
        "📱 Номер телефона станет номером карты. Надёжнее всего — кнопкой "
        "«📱 Отправить мой номер» (номер подтверждается Telegram).",
        reply_markup=kb_phone(),
    )


# --------------------------- /guest — карта гостя (админ) --------------------
class GuestFSM(StatesGroup):
    search = State()   # ввод телефона / @ника
    accrue = State()   # сумма чека вне бота → кешбэк
    spend = State()    # сколько баллов списать
    phone = State()    # новый номер карты


def kb_guest_admin(uid: int) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [_btn("📊 Аналитика гостя", f"ga:an:{uid}")],
        [_btn("➕ Начислить за чек", f"ga:acc:{uid}"), _btn("➖ Списать баллы", f"ga:sp:{uid}")],
        [_btn("📱 Сменить номер", f"ga:ph:{uid}")],
        [_btn("🔄 Обновить", f"ga:v:{uid}"), _btn("✖ Закрыть", "ga:close")],
    ])


async def _admin_guest_search(message: Message, query: str) -> None:
    found = db_guest_search(query)
    if not found:
        await message.answer(
            f"🔍 По запросу «{esc(query)}» карт не найдено.\n"
            "Ищите по телефону (можно последние цифры) или @нику: /guest +79001234567"
        )
        return
    if len(found) == 1:
        uid = int(found[0]["user_id"])
        await message.answer(card_text(uid, for_admin=True), reply_markup=kb_guest_admin(uid))
        return
    rows = [
        [_btn(f"{g.get('phone') or '—'} · {g.get('first_name') or ''}"
              + (f" @{g['username']}" if g.get("username") else ""), f"ga:v:{g['user_id']}")]
        for g in found[:10]
    ]
    await message.answer(
        f"🔍 Найдено карт: {len(found)}. Выберите:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=rows),
    )


@router.message(Command("guest"))
async def cmd_guest(message: Message, command: CommandObject, state: FSMContext):
    if not message.from_user or not _is_admin(message.from_user.id):
        return
    await state.clear()
    q = (command.args or "").strip()
    if q:
        await _admin_guest_search(message, q)
        return
    await state.set_state(GuestFSM.search)
    await message.answer(
        "🔍 Введите телефон гостя (можно последние 4+ цифры) или @ник:",
        reply_markup=_cancel_kb("ga:close"),
    )


@router.message(GuestFSM.search, F.text)
async def guest_search_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    await state.clear()
    await _admin_guest_search(message, (message.text or "").strip())


@router.callback_query(F.data.startswith("ga:"))
async def cb_guest_admin(cb: CallbackQuery, state: FSMContext):
    if not _is_admin(cb.from_user.id):
        await cb.answer("⛔ Только для администратора", show_alert=True)
        return
    parts = cb.data.split(":")
    try:
        if parts[1] == "close":
            await state.clear()
            await cb.answer("Закрыли")
            try:
                await cb.message.edit_reply_markup(reply_markup=None)
            except TelegramBadRequest:
                pass
            return
        uid = int(parts[2])
        if not guest_card(uid):
            await cb.answer("Карта не найдена", show_alert=True)
            return
        if parts[1] == "v":
            await state.clear()
            await cb.answer()
            try:
                await cb.message.edit_text(card_text(uid, for_admin=True),
                                           reply_markup=kb_guest_admin(uid))
            except TelegramBadRequest:
                pass
            return
        if parts[1] == "o":  # из /loyalty — карта отдельным сообщением
            await state.clear()
            await cb.answer()
            await cb.message.answer(card_text(uid, for_admin=True),
                                    reply_markup=kb_guest_admin(uid))
            return
        if parts[1] == "an":
            await state.clear()
            await cb.answer()
            kb = InlineKeyboardMarkup(inline_keyboard=[
                [_btn("🔙 К карте", f"ga:v:{uid}"), _btn("✖ Закрыть", "ga:close")],
            ])
            try:
                await cb.message.edit_text(guest_analytics_text(uid), reply_markup=kb)
            except TelegramBadRequest:
                pass
            return
        if parts[1] in ("acc", "sp"):
            await cb.answer()
            await state.set_state(GuestFSM.accrue if parts[1] == "acc" else GuestFSM.spend)
            await state.update_data(ga_uid=uid, ed_msg=cb.message.message_id)
            prompt = (
                f"➕ Сумма чека в {CURRENCY}, оплаченного мимо бота — начислим "
                f"{LOYALTY_PERCENT}% баллами:"
                if parts[1] == "acc"
                else f"➖ Сколько баллов списать? Баланс: {_fmt_pts(db_bonus_balance(uid))}"
            )
            await cb.message.edit_text(prompt, reply_markup=_cancel_kb(f"ga:v:{uid}"))
            return
        if parts[1] == "ph":
            g = guest_card(uid)
            await cb.answer()
            await state.set_state(GuestFSM.phone)
            await state.update_data(ga_uid=uid, ed_msg=cb.message.message_id)
            await cb.message.edit_text(
                f"📱 <b>Новый номер карты</b> для {esc(_guest_label(g, phone=False))}\n"
                f"Сейчас: <b>{esc(g['phone'])}</b>\n\n"
                "Отправьте новый номер (например +7 900 123-45-67). Баллы и история "
                "останутся на карте, гостю придёт уведомление.",
                reply_markup=_cancel_kb(f"ga:v:{uid}"),
            )
            return
        await cb.answer()
    except (ValueError, IndexError):
        await cb.answer("Некорректный запрос", show_alert=True)
    except Exception:
        logger.exception("Ошибка в карте гостя: %s", cb.data)
        try:
            await cb.answer("Ошибка", show_alert=True)
        except Exception:
            pass


@router.message(Command("loyalty"))
async def cmd_loyalty(message: Message, state: FSMContext):
    if not message.from_user or not _is_admin(message.from_user.id):
        return
    await state.clear()
    o = loyalty_overview(30)
    png = await _render_chart(loyalty_chart_png, o)
    await _answer_chart_and_text(message, png, "💎 Лояльность · 30 дн.",
                                 loyalty_text(o), kb_loyalty(o, "30"))


@router.callback_query(F.data.startswith("la:"))
async def cb_loyalty(cb: CallbackQuery):
    if not _is_admin(cb.from_user.id):
        await cb.answer("⛔ Только для администратора", show_alert=True)
        return
    key = cb.data.split(":", 1)[1]
    if key not in _LA_PERIODS:
        await cb.answer("Некорректный период", show_alert=True)
        return
    o = loyalty_overview(_LA_PERIODS[key])
    await cb.answer()
    try:
        await cb.message.edit_text(loyalty_text(o), reply_markup=kb_loyalty(o, key))
    except TelegramBadRequest:
        return  # тот же период — картинка уже актуальна
    png = await _render_chart(loyalty_chart_png, o)
    await _update_chart(cb.bot, cb.message.chat.id, cb.message.message_id, png,
                        "💎 Лояльность · " + ("всё время" if key == "all" else f"{key} дн."))


async def _notify_guest(bot, uid: int, text: str) -> None:
    try:
        await bot.send_message(uid, text)
    except Exception:
        logger.info("Не удалось уведомить гостя %s о баллах", uid)


@router.message(GuestFSM.accrue, F.text)
async def guest_accrue_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    data = await state.get_data()
    uid = data.get("ga_uid")
    v = _parse_num((message.text or "").strip())
    if not uid or not guest_card(uid):
        await state.clear()
        await message.answer("Карта не найдена — откройте /guest заново.")
        return
    if v is None or v < 1 or v > 1_000_000 or v != int(v):
        await message.answer("Введите сумму чека целым числом (например 3500), либо ✖ Отмена.")
        return
    pts = loyalty_offline_accrue(int(uid), int(v), message.from_user.id)
    await state.clear()
    await _edit_or_send(
        message, data,
        f"✅ Начислено +{_fmt_pts(pts)} баллов за чек {_fmt_money(int(v))}\n\n"
        + card_text(int(uid), for_admin=True),
        kb_guest_admin(int(uid)),
    )
    if pts:
        await _notify_guest(
            message.bot, int(uid),
            f"💎 Начислено +{_fmt_pts(pts)} баллов за чек {_fmt_money(int(v))}. "
            f"Баланс: {_fmt_pts(db_bonus_balance(int(uid)))}",
        )


@router.message(GuestFSM.phone, F.text)
async def guest_phone_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    data = await state.get_data()
    uid = data.get("ga_uid")
    g = guest_card(uid) if uid else None
    if not g:
        await state.clear()
        await message.answer("Карта не найдена — откройте /guest заново.")
        return
    phone = normalize_phone((message.text or "").strip())
    if not phone:
        await message.answer("Не похоже на номер. Пример: +7 900 123-45-67 — либо ✖ Отмена.")
        return
    if phone == g["phone"]:
        await message.answer("Это и есть текущий номер карты. Введите другой, либо ✖ Отмена.")
        return
    ok, owner = db_guest_admin_set_phone(int(uid), phone)
    if not ok:
        who = f": {esc(_guest_label(owner))}" if owner else ""
        await message.answer(f"⚠️ Номер {esc(phone)} уже привязан к другой карте{who}. "
                             "Введите другой, либо ✖ Отмена.")
        return
    await state.clear()
    logger.info("Админ %s сменил номер карты %s: %s → %s",
                message.from_user.id, uid, g["phone"], phone)
    await _edit_or_send(
        message, data,
        f"✅ Номер карты: {esc(g['phone'])} → <b>{esc(phone)}</b>\n\n"
        + card_text(int(uid), for_admin=True),
        kb_guest_admin(int(uid)),
    )
    await _notify_guest(
        message.bot, int(uid),
        f"📱 Номер вашей карты изменён администратором: <b>{esc(phone)}</b>. "
        "Баллы и история сохранены.",
    )


@router.message(GuestFSM.spend, F.text)
async def guest_spend_input(message: Message, state: FSMContext):
    if not await _admin_msg_guard(message, state):
        return
    data = await state.get_data()
    uid = data.get("ga_uid")
    v = _parse_num((message.text or "").strip())
    if not uid or not guest_card(uid):
        await state.clear()
        await message.answer("Карта не найдена — откройте /guest заново.")
        return
    if v is None or v < 1 or v != int(v):
        await message.answer("Введите количество баллов целым числом, либо ✖ Отмена.")
        return
    if not loyalty_manual_spend(int(uid), int(v), message.from_user.id):
        await message.answer(
            f"Недостаточно баллов: баланс {_fmt_pts(db_bonus_balance(int(uid)))}. "
            "Введите меньше, либо ✖ Отмена."
        )
        return
    await state.clear()
    await _edit_or_send(
        message, data,
        f"✅ Списано {_fmt_pts(int(v))} баллов\n\n" + card_text(int(uid), for_admin=True),
        kb_guest_admin(int(uid)),
    )
    await _notify_guest(
        message.bot, int(uid),
        f"💎 Списано {_fmt_pts(int(v))} баллов. Баланс: {_fmt_pts(db_bonus_balance(int(uid)))}",
    )


# ========================= ЗАКАЗ ПО АКЦИИ («Выгода») ========================
# Отдельная корзина pcart и отдельный заказ kind='promo'.
def _promo_text() -> str:
    hh = hh_active()
    lines = [
        "🎁 <b>Заказ по акции</b>",
        "Отдельный заказ: акции не суммируются, баллы не начисляются и не списываются.",
        "",
        "☀️ <b>Счастливые часы</b> — по будням 15:00–18:00: "
        + ("<b>сейчас действуют</b> ✅" if hh else "сейчас не действуют"),
        "🔥 <b>Комбо на электронной чаше ХУКА Про</b> — в любое время",
    ]
    return "\n".join(lines)


def kb_promo(pcart_count: int = 0) -> InlineKeyboardMarkup:
    reserved = reserved_snapshot()
    rows: list[list[InlineKeyboardButton]] = []
    hh = hh_active()

    def _row(pid: str):
        p = promo_item(pid)
        if not p or not item_in_stock(pid, reserved):
            return None
        old = f" (вместо {_fmt_money(p['old'])})" if p["old"] > p["price"] else ""
        return [_btn(f"＋ {p['name']} — {_fmt_money(p['price'])}{old}", f"p:a:{pid}")]

    for pid, p in PROMO_ITEMS.items():
        if p["hh"] and not hh:
            continue
        r = _row(pid)
        if r:
            rows.append(r)
    if (hh or PROMO_TEA_ALWAYS) and _items_in_cats(PROMO_TEA_HOOKAH_CATS) and _items_in_cats(PROMO_TEA_CATS):
        rows.append([_btn("🍵 Любой кальян по полной цене + чай в подарок", "p:tea")])
    rows.append([_btn("📅 Аренда VIP-комнаты / PlayStation", "p:rent")])
    label = "🛒 Корзина акций" + (f" ({pcart_count})" if pcart_count else "")
    rows.append([_btn(label, "p:cart")])
    rows.append([_btn("🔙 В меню", "m")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_pcart(cart: list[dict]) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    for idx in range(len(cart)):
        rows.append([
            _btn("➖", f"p:cd:{idx}"),
            _btn(f"{cart[idx]['qty']} шт.", "noop"),
            _btn("➕", f"p:ci:{idx}"),
            _btn("🗑", f"p:cr:{idx}"),
        ])
    rows.append([_btn("✅ Оформить заказ по акции", "p:co")])
    rows.append([_btn("🗑 Очистить", "p:cx")])
    rows.append([_btn("🔙 К акциям", "p:home")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_pconfirm() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [_btn("✅ Отправить заказ", "p:so")],
        [_btn("🔙 Назад к корзине", "p:cart")],
    ])


def _promo_revalidate(cart: list[dict]) -> tuple[list[dict], list[str]]:
    """Перед отправкой: акция существует, цена актуальна, счастливые часы
    ещё идут (с запасом HH_GRACE_MIN). Возвращает (корзина, что убрали)."""
    ok: list[dict] = []
    bad: list[str] = []
    hh = hh_active(HH_GRACE_MIN)
    for it in cart:
        p = promo_item(str(it.get("id") or ""))
        if not p:
            bad.append(f"• {esc(it.get('name', '?'))} — акция недоступна")
            continue
        if p["hh"] and not hh:
            bad.append(f"• {esc(p['name'])} — счастливые часы закончились")
            continue
        ok.append({"id": p["id"], "name": p["name"], "price": p["price"],
                   "qty": max(1, int(it.get("qty") or 1))})
    return ok, bad


async def _pcart(state: FSMContext) -> list[dict]:
    return list((await state.get_data()).get("pcart") or [])


async def _render_pcart(cb: CallbackQuery, state: FSMContext) -> None:
    cart = await _pcart(state)
    table = await _table(state)
    if not cart:
        try:
            await cb.message.edit_text(_promo_text() + "\n\n🛒 Корзина акций пуста.",
                                       reply_markup=kb_promo(0))
        except TelegramBadRequest:
            pass
        return
    try:
        await cb.message.edit_text(
            cart_text(cart, table, title="🎁 <b>Корзина акций</b>"), reply_markup=kb_pcart(cart)
        )
    except TelegramBadRequest:
        pass


async def _promo_add(cb: CallbackQuery, state: FSMContext, pid: str) -> None:
    p = promo_item(pid)
    if not p:
        await cb.answer(ERR_MENU, show_alert=True)
        return
    if p["hh"] and not hh_active():
        await cb.answer("☀️ Счастливые часы сейчас не действуют (будни 15:00–18:00)",
                        show_alert=True)
        return
    if not await _table(state):
        await cb.answer(ASK_TABLE_FIRST, show_alert=True)
        return
    cart = await _pcart(state)
    if not _is_admin(cb.from_user.id) and _cart_count(cart) >= ORDER_MAX_ITEMS:
        await cb.answer(f"⚠️ В одном заказе не больше {ORDER_MAX_ITEMS} позиций", show_alert=True)
        return
    limit = max_available(pid)
    in_cart = next((i["qty"] for i in cart if i["id"] == pid), 0)
    if limit is not None and in_cart + 1 > limit:
        await cb.answer(f"⚠️ По остаткам на складе можно не более {limit} шт.", show_alert=True)
        return
    found = next((i for i in cart if i["id"] == pid), None)
    if found:
        found["qty"] += 1
    else:
        cart.append({"id": pid, "name": p["name"], "price": p["price"], "qty": 1})
    await state.update_data(pcart=cart)
    await cb.answer(f"«{p['name']}» — в корзине акций ({_cart_count(cart)} шт.)")


@router.callback_query(F.data.startswith("p:"))
async def cb_promo(cb: CallbackQuery, state: FSMContext):
    parts = cb.data.split(":")
    act = parts[1] if len(parts) > 1 else ""
    try:
        if act in ("home", "new"):
            await cb.answer()
            cart = await _pcart(state)
            text, kb = _promo_text(), kb_promo(_cart_count(cart))
            if act == "new":  # из-под фото «Выгода» — новым сообщением
                await cb.message.answer(text, reply_markup=kb)
            else:
                try:
                    await cb.message.edit_text(text, reply_markup=kb)
                except TelegramBadRequest:
                    pass
            return
        if act == "rent":
            await _begin_flow("rsv", cb.message, state, cb=cb, strip=False)
            return
        if act == "a" and len(parts) == 3:
            await _promo_add(cb, state, parts[2])
            if _is_promo_id(parts[2]) and not parts[2].startswith("pr_t:"):
                try:
                    await cb.message.edit_reply_markup(
                        reply_markup=kb_promo(_cart_count(await _pcart(state)))
                    )
                except TelegramBadRequest:
                    pass
            return
        if act == "tea":
            if not (hh_active() or PROMO_TEA_ALWAYS):
                await cb.answer("☀️ Акция действует в счастливые часы (будни 15:00–18:00)",
                                show_alert=True)
                return
            await cb.answer()
            reserved = reserved_snapshot()
            rows = [
                [_btn(f"{h['name']} — {_fmt_money(h['price'])}", f"p:th:{h['id']}")]
                for h in _items_in_cats(PROMO_TEA_HOOKAH_CATS)
                if item_in_stock(h["id"], reserved)
            ]
            rows.append([_btn("🔙 К акциям", "p:home")])
            await cb.message.edit_text(
                "🍵 <b>Кальян по полной цене + чай в подарок</b>\n\nШаг 1 — выберите кальян:",
                reply_markup=InlineKeyboardMarkup(inline_keyboard=rows),
            )
            return
        if act == "th" and len(parts) == 3:
            hid = parts[2]
            hookah = next((i for i in _items_in_cats(PROMO_TEA_HOOKAH_CATS) if i["id"] == hid), None)
            if not hookah:
                await cb.answer(ERR_MENU, show_alert=True)
                return
            await cb.answer()
            reserved = reserved_snapshot()
            rows = [
                [_btn(f"🎁 {t['name']}", f"p:tt:{hid}:{t['id']}")]
                for t in _items_in_cats(PROMO_TEA_CATS)
                if item_in_stock(t["id"], reserved)
            ]
            rows.append([_btn("🔙 Другой кальян", "p:tea")])
            await cb.message.edit_text(
                f"🍵 {esc(hookah['name'])} — {_fmt_money(hookah['price'])}\n\n"
                "Шаг 2 — выберите чай в подарок:",
                reply_markup=InlineKeyboardMarkup(inline_keyboard=rows),
            )
            return
        if act == "tt" and len(parts) == 4:
            await _promo_add(cb, state, f"pr_t:{parts[2]}:{parts[3]}")
            try:
                await cb.message.edit_text(_promo_text(),
                                           reply_markup=kb_promo(_cart_count(await _pcart(state))))
            except TelegramBadRequest:
                pass
            return
        if act == "cart":
            await cb.answer()
            await _render_pcart(cb, state)
            return
        if act in ("ci", "cd", "cr") and len(parts) == 3:
            cart = await _pcart(state)
            try:
                idx = int(parts[2])
            except ValueError:
                idx = -1
            if not 0 <= idx < len(cart):
                await cb.answer(ERR_MENU, show_alert=True)
                return
            if act == "ci":
                p = promo_item(str(cart[idx].get("id")))
                if not p:
                    await cb.answer(ERR_MENU, show_alert=True)
                    return
                if p["hh"] and not hh_active():
                    await cb.answer("☀️ Счастливые часы закончились", show_alert=True)
                    return
                if not _is_admin(cb.from_user.id) and _cart_count(cart) >= ORDER_MAX_ITEMS:
                    await cb.answer(f"⚠️ В одном заказе не больше {ORDER_MAX_ITEMS} позиций",
                                    show_alert=True)
                    return
                limit = max_available(p["id"])
                if limit is not None and cart[idx]["qty"] + 1 > limit:
                    await cb.answer(f"⚠️ По остаткам можно не более {limit} шт.", show_alert=True)
                    return
                cart[idx]["qty"] += 1
            elif act == "cd":
                cart[idx]["qty"] -= 1
                if cart[idx]["qty"] <= 0:
                    cart.pop(idx)
            else:
                cart.pop(idx)
            await state.update_data(pcart=cart)
            await cb.answer()
            await _render_pcart(cb, state)
            return
        if act == "cx":
            await state.update_data(pcart=[])
            await cb.answer("Корзина акций очищена")
            await _render_pcart(cb, state)
            return
        if act == "co":
            cart = await _pcart(state)
            if not cart:
                await cb.answer(CART_EMPTY, show_alert=True)
                return
            table = await _table(state)
            if not table:
                await cb.answer()
                await state.set_state(OrderFSM.table)
                await cb.message.edit_text(ASK_TABLE)
                return
            await cb.answer()
            await cb.message.edit_text(confirm_text(cart, table, promo=True),
                                       reply_markup=kb_pconfirm())
            return
        if act == "so":
            await _submit_order(cb, state, "promo")
            return
        await cb.answer()
    except Exception:
        logger.exception("Ошибка в заказе по акции: %s", cb.data)
        try:
            await cb.answer("⚠️ Ошибка — попробуйте ещё раз", show_alert=True)
        except Exception:
            pass


# =============================== ПРЕДЗАКАЗ ==================================
# Гость с картой собирает заказ заранее — к времени прихода (сегодня/завтра)
# или к своей активной брони. Отдельная корзина precart, заказ kind='pre'
# со временем прихода arrive_at. Оплата — на месте, кешбэк как у обычного.
# Отменить сам гость может, пока заказ в статусе «Принят» (до «Готов»).
# Сроки и лимиты — в НАСТРОЙКИ → «Предзаказ».
PRE_TABLE = "предзаказ"    # orders.table_no для предзаказа


def _place_short(table_no: str) -> str:
    return PRE_TABLE if table_no == PRE_TABLE else f"стол {table_no}"


def _pre_when(o: dict) -> str:
    """«пн 12.10 · 20:00» по arrive_at заказа."""
    try:
        a = datetime.strptime(str(o.get("arrive_at"))[:16], "%Y-%m-%d %H:%M")
    except (TypeError, ValueError):
        return "—"
    return f"{_wd(a)} {a:%d.%m} · {a:%H:%M}"


def _pre_too_early(order: dict) -> bool:
    """Предзаказ на БУДУЩИЙ рабочий день (смена 15:00–03:00): «Готов», «Выдан»
    и «Оплачен» ставить ещё рано — гость не пришёл. Приход в 01:00 ночи —
    это ещё смена предыдущей даты."""
    if order.get("kind") != "pre" or not order.get("arrive_at"):
        return False
    try:
        a = datetime.strptime(str(order["arrive_at"])[:16], "%Y-%m-%d %H:%M")
    except ValueError:
        return False
    return (a - timedelta(hours=RSV_CLOSE_H)).date() > _biz_today()


def _pre_early_msg(order: dict) -> str:
    return (f"⏳ Предзаказ к {_pre_when(order)} — день прихода ещё не наступил. "
            "«Готов», «Выдан» и «Оплачен» — в день прихода. Отменить можно сейчас.")


def _pre_days() -> list:
    """Рабочие дни для предзаказа: текущий (включая «эту ночь») и следующий."""
    b = _biz_today()
    return [b, b + timedelta(days=1)]


def _pre_slots(d) -> list[str]:
    """Свободные для предзаказа слоты дня d: в часы работы, ≥ сейчас + PRE_LEAD_MIN."""
    edge = _now_dt() + timedelta(minutes=PRE_LEAD_MIN)
    close = datetime(d.year, d.month, d.day) + timedelta(days=1, hours=RSV_CLOSE_H)
    last = close - timedelta(minutes=PRE_LAST_BEFORE_CLOSE)
    out = []
    for t in _slots():
        a = _arrival(f"{d:%Y-%m-%d}", t)
        if edge <= a <= last:
            out.append(t)
    return out


def db_pre_active(uid: int) -> list[dict]:
    """Активные предзаказы гостя (не оплачены, не отменены, приход не в прошлом)."""
    since = (_now_dt() - timedelta(hours=6)).strftime("%Y-%m-%d %H:%M:%S")
    with _connect() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT id, total, status, arrive_at, rsv_id, message_id FROM orders"
            " WHERE user_id = ? AND kind = 'pre' AND status IN ('accepted', 'ready', 'issued')"
            " AND arrive_at >= ? ORDER BY arrive_at",
            (uid, since),
        )]


def db_pre_cancel(order_id: int, uid: int) -> bool:
    """Отмена предзаказа гостем — только свой и только пока «Принят». Атомарно:
    если персонал уже нажал «Готов», ничего не меняем."""
    with _connect() as conn:
        cur = conn.execute(
            "UPDATE orders SET status = 'cancelled', updated_at = ?, issued_at = NULL"
            " WHERE id = ? AND user_id = ? AND kind = 'pre' AND status = 'accepted'",
            (_now(), int(order_id), int(uid)),
        )
        return cur.rowcount == 1


def _pre_rsv_options(uid: int) -> list[dict]:
    """Активные брони гостя, к которым ещё можно сделать предзаказ."""
    days = {f"{d:%Y-%m-%d}" for d in _pre_days()}
    edge = _now_dt() + timedelta(minutes=PRE_LEAD_MIN)
    with _connect() as conn:
        rows = [dict(r) for r in conn.execute(
            "SELECT id, kind, date, time, status FROM reservations"
            " WHERE user_id = ? AND status IN ('new', 'confirmed') ORDER BY date, time",
            (uid,),
        )]
    out = []
    for r in rows:
        try:
            if r["date"] in days and _arrival(r["date"], r["time"]) >= edge:
                out.append(r)
        except (ValueError, TypeError):
            continue
    return out


def _pre_draft_when(pre: dict) -> str:
    return _rsv_when_text(pre["date"], pre["time"])


def pre_home_text(uid: int, data: dict) -> tuple[str, InlineKeyboardMarkup]:
    active = db_pre_active(uid)
    L = ["📝 <b>Предзаказ</b>", "",
         "Соберите заказ заранее — к вашему приходу всё будет готово. "
         "Оплата — на месте" + (f", кешбэк {LOYALTY_PERCENT}% как обычно." if LOYALTY_PERCENT
                                else ".")]
    rows: list[list[InlineKeyboardButton]] = []
    if active:
        L += ["", "<b>Ваши предзаказы:</b>"]
        for o in active:
            link = f" · к брони #{o['rsv_id']}" if o.get("rsv_id") else ""
            L.append(f"• #{o['id']} · {_pre_when(o)}{link} · {_fmt_money(o['total'])} · "
                     f"{STATUSES.get(o['status'], o['status'])}")
            if o["status"] == "accepted":
                rows.append([_btn(f"❌ Отменить #{o['id']}", f"po:x:{o['id']}")])
    pre = data.get("pre") or {}
    pcart = data.get("precart") or []
    if pre.get("date") and pcart:
        L += ["", f"🛒 Черновик к {_pre_draft_when(pre)}: {_cart_count(pcart)} шт."]
        rows.append([_btn("🛒 Продолжить черновик", "po:cart")])
    if len(active) >= PRE_MAX_ACTIVE and not _is_admin(uid):
        L += ["", f"ℹ️ Одновременно можно держать не больше {PRE_MAX_ACTIVE} предзаказов."]
    else:
        for r in _pre_rsv_options(uid):
            rows.append([_btn(
                f"📅 К брони #{r['id']} · {_rsv_when_text(r['date'], r['time'])}",
                f"po:r:{r['id']}")])
        if any(_pre_slots(d) for d in _pre_days()):
            rows.append([_btn("⏰ Выбрать время прихода", "po:time")])
        else:
            L += ["", "⏰ На сегодня время для предзаказа уже вышло — загляните завтра."]
    rows.append([_btn("✖ Закрыть", "po:close")])
    return "\n".join(L), InlineKeyboardMarkup(inline_keyboard=rows)


def kb_pre_days() -> InlineKeyboardMarkup:
    rows = [[_btn(_date_label(d), f"po:d:{d:%Y-%m-%d}")] for d in _pre_days() if _pre_slots(d)]
    rows.append([_btn("🔙 Назад", "po:home")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_pre_slots(d) -> InlineKeyboardMarkup:
    btns = [_btn(t, f"po:t:{d:%Y%m%d}{t.replace(':', '')}") for t in _pre_slots(d)]
    rows = [btns[i : i + 4] for i in range(0, len(btns), 4)]
    rows.append([_btn("🔙 Другой день", "po:time")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_pre_cats(n: int) -> InlineKeyboardMarkup:
    rows = [[_btn(f"{c['emoji']} {c['name']}", f"po:c:{c['id']}")] for c in MENU["categories"]]
    rows.append([_btn("🛒 Корзина предзаказа" + (f" ({n})" if n else ""), "po:cart")])
    rows.append([_btn("⏰ Изменить время", "po:time"), _btn("🔙 Назад", "po:home")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_pre_cat(cat_id: str, n: int) -> InlineKeyboardMarkup:
    cat = f_category(cat_id) or {"items": []}
    reserved = reserved_snapshot()
    rows = [[_btn(f"＋ {it['name']} — {_fmt_money(it['price'])}", f"po:a:{it['id']}")]
            for it in cat["items"] if item_in_stock(it["id"], reserved)]
    rows.append([_btn("🛒 Корзина предзаказа" + (f" ({n})" if n else ""), "po:cart")])
    rows.append([_btn("🔙 Разделы", "po:cats")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def kb_pre_cart(cart: list[dict]) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    for idx in range(len(cart)):
        rows.append([
            _btn("➖", f"po:cd:{idx}"),
            _btn(f"{cart[idx]['qty']} шт.", "noop"),
            _btn("➕", f"po:ci:{idx}"),
            _btn("🗑", f"po:cr:{idx}"),
        ])
    rows.append([_btn("✅ Оформить предзаказ", "po:co")])
    rows.append([_btn("➕ Добавить ещё", "po:cats"), _btn("🗑 Очистить", "po:cx")])
    rows.append([_btn("🔙 Назад", "po:home")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def _pre_head(pre: dict) -> str:
    s = f"📝 <b>Предзаказ</b> к <b>{_pre_draft_when(pre)}</b>"
    if pre.get("rsv_id"):
        s += f"\n📅 К брони #{pre['rsv_id']}"
    return s


def pre_confirm_text(cart: list[dict], pre: dict) -> str:
    L = ["✅ <b>Подтверждение предзаказа</b>", "",
         f"⏰ К приходу: <b>{_pre_draft_when(pre)}</b>"]
    if pre.get("rsv_id"):
        L.append(f"📅 К брони #{pre['rsv_id']}")
    L.append("")
    for i in cart:
        L.append(f"• {esc(i['name'])} × {i['qty']} — {_fmt_money(i['price'] * i['qty'])}")
    L += ["", f"💰 <b>К оплате на месте: {_fmt_money(_cart_total(cart))}</b>", "",
          "Отменить можно здесь же, пока персонал не отметил заказ «Готов».", "",
          "Отправить предзаказ?"]
    return "\n".join(L)


async def _pre_edit(cb: CallbackQuery, text: str, kb) -> None:
    try:
        await cb.message.edit_text(text, reply_markup=kb)
    except TelegramBadRequest as e:
        if "not modified" not in str(e).lower():
            await cb.message.answer(text, reply_markup=kb)


def _pre_valid(pre: dict | None) -> bool:
    """Черновик ещё актуален: день в горизонте, время в будущем."""
    if not pre or not pre.get("date") or not pre.get("time"):
        return False
    try:
        a = _arrival(pre["date"], pre["time"])
    except (ValueError, TypeError):
        return False
    return (pre["date"] in {f"{d:%Y-%m-%d}" for d in _pre_days()}
            and a >= _now_dt() + timedelta(minutes=PRE_SUBMIT_MIN))


def _pre_need_card_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[[_btn("📱 Подключить карту", "po:card")]])


PRE_NEED_CARD = ("📝 <b>Предзаказ</b> доступен гостям с картой лояльности — так мы "
                 "узнаем вас при приходе. Подключение — один шаг: номер телефона.")


async def _pre_home(message: Message, state: FSMContext, uid: int,
                    cb: CallbackQuery | None = None) -> None:
    if not guest_card(uid):
        if cb:
            await _pre_edit(cb, PRE_NEED_CARD, _pre_need_card_kb())
        else:
            await message.answer(PRE_NEED_CARD, reply_markup=_pre_need_card_kb())
        return
    text, kb = pre_home_text(uid, await state.get_data())
    if cb:
        await _pre_edit(cb, text, kb)
    else:
        await message.answer(text, reply_markup=kb)


async def _pre_render_cart(cb: CallbackQuery, state: FSMContext) -> None:
    data = await state.get_data()
    cart = list(data.get("precart") or [])
    pre = data.get("pre") or {}
    if not cart:
        await _pre_edit(cb, _pre_head(pre) + "\n\n🛒 Корзина предзаказа пуста. Выберите раздел:",
                        kb_pre_cats(0))
        return
    await _pre_edit(cb, cart_text(cart, None, title=_pre_head(pre)), kb_pre_cart(cart))


async def _submit_preorder(cb: CallbackQuery, state: FSMContext) -> None:
    uid = cb.from_user.id
    lock = _user_lock(uid)
    if lock.locked():
        await cb.answer("⏳ Уже отправляем…")
        return
    async with lock:
        card = guest_card(uid)
        if not card:
            await cb.answer("Нужна карта лояльности", show_alert=True)
            await _pre_edit(cb, PRE_NEED_CARD, _pre_need_card_kb())
            return
        data = await state.get_data()
        cart = list(data.get("precart") or [])
        pre = dict(data.get("pre") or {})
        if not cart:
            await cb.answer(CART_EMPTY, show_alert=True)
            return
        if not _pre_valid(pre):
            await cb.answer("⏰ Это время уже не подходит — выберите другое", show_alert=True)
            await state.update_data(pre={})
            await _pre_edit(cb, "⏰ <b>Когда придёте?</b>", kb_pre_days())
            return
        rsv_id = pre.get("rsv_id")
        if rsv_id and not any(r["id"] == rsv_id for r in _pre_rsv_options(uid)):
            await cb.answer("📅 Эта бронь уже неактивна — выберите время прихода",
                            show_alert=True)
            await state.update_data(pre={})
            await _pre_edit(cb, "⏰ <b>Когда придёте?</b>", kb_pre_days())
            return
        if not _is_admin(uid) and len(db_pre_active(uid)) >= PRE_MAX_ACTIVE:
            await cb.answer(f"Одновременно можно держать не больше {PRE_MAX_ACTIVE} "
                            "предзаказов", show_alert=True)
            return
        err = _order_limits_error(uid, cart)
        if err:
            await cb.answer(err, show_alert=True)
            return
        problems = cart_stock_problems(cart)
        if problems:
            await cb.answer("⚠️ Не хватает на складе", show_alert=True)
            await cb.message.answer("⚠️ <b>Не хватает на складе:</b>\n" + "\n".join(problems)
                                    + "\n\nУменьшите количество или выберите другое.")
            await _pre_render_cart(cb, state)
            return
        await cb.answer("Отправляем…")
        total = _cart_total(cart)
        arrive = _arrival(pre["date"], pre["time"]).strftime("%Y-%m-%d %H:%M:%S")
        user = cb.from_user
        order_id = db_add_order(table_no=PRE_TABLE, total=total, user_id=uid,
                                username=user.username if user else None,
                                phone=card["phone"], kind="pre",
                                arrive_at=arrive, rsv_id=rsv_id)
        db_add_items(order_id, cart)
        order = db_get_order(order_id)
        try:
            sent = await cb.bot.send_message(CHANNEL_ID, order_text(order),
                                             reply_markup=kb_status(order_id))
        except Exception:
            logger.exception("Не удалось отправить предзаказ #%s в канал", order_id)
            db_delete_order(order_id)
            await _pre_edit(cb, "⚠️ Не удалось отправить предзаказ. Попробуйте ещё раз "
                                "или сообщите администратору.",
                            InlineKeyboardMarkup(inline_keyboard=[
                                [_btn("✅ Отправить ещё раз", "po:so")],
                                [_btn("🔙 К корзине", "po:cart")]]))
            return
        db_set_message_id(order_id, sent.message_id)
        await state.update_data(precart=[], pre={})
        db_guest_touch(user)
        text = (f"✅ Предзаказ <b>#{order_id}</b> отправлен!\n"
                f"⏰ К приходу: <b>{_pre_when(order)}</b>\n"
                + (f"📅 К брони #{rsv_id}\n" if rsv_id else "")
                + f"💰 Сумма: <b>{_fmt_money(total)}</b> — оплата на месте\n\n"
                "Персонал подготовит всё к вашему приходу. Статус пришлём сюда.")
        if LOYALTY_PERCENT:
            text += f"\n💎 После оплаты начислим кешбэк {LOYALTY_PERCENT}% баллами."
        await _pre_edit(cb, text, InlineKeyboardMarkup(inline_keyboard=[
            [_btn("❌ Отменить предзаказ", f"po:x:{order_id}")],
            [_btn("📝 Мои предзаказы", "po:home")]]))


async def _pre_cancel(cb: CallbackQuery, order_id: int) -> None:
    uid = cb.from_user.id
    o = db_get_order(order_id)
    if not o or o.get("user_id") != uid or o.get("kind") != "pre":
        await cb.answer("Предзаказ не найден", show_alert=True)
        return
    if not db_pre_cancel(order_id, uid):
        st = STATUSES.get(o["status"], o["status"])
        await cb.answer(f"Отменить уже нельзя (статус: {st}). Если планы изменились — "
                        "напишите администратору.", show_alert=True)
        return
    # как отмена персоналом: склад (если списан) и баллы — ровно один раз
    if db_claim_stock_flag(order_id, 0):
        await _stock_alerts(cb.bot, stock_apply_order(o, +1))
    loyalty_on_refund(order_id)
    db_set_bonus_pending(order_id, 0)
    o = db_get_order(order_id)
    await cb.answer("Предзаказ отменён")
    if o.get("message_id"):
        try:
            await cb.bot.edit_message_text(order_text(o), chat_id=CHANNEL_ID,
                                           message_id=o["message_id"],
                                           reply_markup=kb_status(order_id))
        except Exception:
            logger.info("Не удалось обновить сообщение предзаказа #%s", order_id)
    try:
        await cb.bot.send_message(
            CHANNEL_ID, f"❌ Гость отменил предзаказ <b>#{order_id}</b> ({_pre_when(o)})",
            reply_to_message_id=o.get("message_id") or None)
    except Exception:
        logger.info("Не удалось уведомить канал об отмене предзаказа #%s", order_id)
    await _pre_edit(cb, f"❌ Предзаказ <b>#{order_id}</b> отменён.",
                    InlineKeyboardMarkup(inline_keyboard=[[_btn("📝 Предзаказ", "po:home")]]))


@router.message(F.text == BTN_PRE)
@router.message(Command("preorder"))
async def on_btn_pre(message: Message, state: FSMContext):
    if message.chat.type != "private" or not message.from_user:
        return
    await state.set_state(None)
    await _begin_flow("pre", message, state)


@router.callback_query(F.data.startswith("po:"))
async def cb_preorder(cb: CallbackQuery, state: FSMContext):
    parts = cb.data.split(":")
    act = parts[1] if len(parts) > 1 else ""
    uid = cb.from_user.id
    try:
        if act == "card":
            await cb.answer()
            await state.update_data(next="pre")
            await state.set_state(OrderFSM.phone)
            await cb.message.answer(
                "📱 Номер телефона станет номером карты. Надёжнее всего — кнопкой "
                "«📱 Отправить мой номер» (номер подтверждается Telegram).",
                reply_markup=kb_phone(),
            )
            return
        if act == "close":
            await cb.answer()
            try:
                await cb.message.delete()
            except TelegramBadRequest:
                try:
                    await cb.message.edit_reply_markup(reply_markup=None)
                except TelegramBadRequest:
                    pass
            return
        if act == "x" and len(parts) == 3:
            await _pre_cancel(cb, int(parts[2]))
            return
        # дальше — только с картой
        if not guest_card(uid):
            await cb.answer()
            await _pre_edit(cb, PRE_NEED_CARD, _pre_need_card_kb())
            return
        if act == "home":
            await cb.answer()
            await _pre_home(cb.message, state, uid, cb=cb)
            return
        if act == "time":
            await cb.answer()
            if not any(_pre_slots(d) for d in _pre_days()):
                await _pre_edit(cb, "⏰ Время для предзаказа на сегодня и завтра вышло.",
                                InlineKeyboardMarkup(inline_keyboard=[
                                    [_btn("🔙 Назад", "po:home")]]))
                return
            await _pre_edit(cb, "⏰ <b>Когда придёте?</b>\n\nВыберите день:", kb_pre_days())
            return
        if act == "d" and len(parts) == 3:
            try:
                d = datetime.strptime(parts[2], "%Y-%m-%d").date()
            except ValueError:
                await cb.answer("Некорректная дата", show_alert=True)
                return
            if d not in _pre_days() or not _pre_slots(d):
                await cb.answer("На этот день предзаказ уже недоступен", show_alert=True)
                return
            await cb.answer()
            await _pre_edit(cb, f"⏰ <b>{_date_label(d)}</b>\n\nВо сколько придёте?",
                            kb_pre_slots(d))
            return
        if act == "t" and len(parts) == 3 and re.fullmatch(r"\d{12}", parts[2]):
            raw = parts[2]
            date_s = f"{raw[:4]}-{raw[4:6]}-{raw[6:8]}"
            time_s = f"{raw[8:10]}:{raw[10:12]}"
            try:
                d = datetime.strptime(date_s, "%Y-%m-%d").date()
            except ValueError:
                await cb.answer("Некорректная дата", show_alert=True)
                return
            if d not in _pre_days() or time_s not in _pre_slots(d):
                await cb.answer("Это время уже недоступно — выберите другое", show_alert=True)
                return
            pre = {"date": date_s, "time": time_s}
            await state.update_data(pre=pre)
            await cb.answer(f"Ждём вас к {time_s}")
            n = _cart_count((await state.get_data()).get("precart") or [])
            await _pre_edit(cb, _pre_head(pre) + "\n\nВыберите раздел:", kb_pre_cats(n))
            return
        if act == "r" and len(parts) == 3 and parts[2].isdigit():
            r = next((x for x in _pre_rsv_options(uid) if x["id"] == int(parts[2])), None)
            if not r:
                await cb.answer("Эта бронь недоступна для предзаказа", show_alert=True)
                return
            pre = {"date": r["date"], "time": r["time"], "rsv_id": r["id"]}
            await state.update_data(pre=pre)
            await cb.answer()
            n = _cart_count((await state.get_data()).get("precart") or [])
            await _pre_edit(cb, _pre_head(pre) + "\n\nВыберите раздел:", kb_pre_cats(n))
            return
        # каталог/корзина — нужен выбранный момент прихода
        data = await state.get_data()
        pre = data.get("pre") or {}
        if not _pre_valid(pre):
            await cb.answer("Сначала выберите время прихода", show_alert=True)
            await state.update_data(pre={})
            await _pre_home(cb.message, state, uid, cb=cb)
            return
        cart = list(data.get("precart") or [])
        if act == "cats":
            await cb.answer()
            await _pre_edit(cb, _pre_head(pre) + "\n\nВыберите раздел:",
                            kb_pre_cats(_cart_count(cart)))
            return
        if act == "c" and len(parts) == 3:
            cat = f_category(parts[2])
            if not cat:
                await cb.answer(ERR_MENU, show_alert=True)
                return
            await cb.answer()
            await _pre_edit(cb, f"{_pre_head(pre)}\n\n{cat['emoji']} <b>{esc(cat['name'])}</b>",
                            kb_pre_cat(cat["id"], _cart_count(cart)))
            return
        if act == "a" and len(parts) == 3:
            item = f_find_item(parts[2])
            if not item:
                await cb.answer(ERR_MENU, show_alert=True)
                return
            if not _is_admin(uid) and _cart_count(cart) >= ORDER_MAX_ITEMS:
                await cb.answer(f"⚠️ В одном заказе не больше {ORDER_MAX_ITEMS} позиций",
                                show_alert=True)
                return
            limit = max_available(item["id"])
            in_cart = next((i["qty"] for i in cart if i["id"] == item["id"]), 0)
            if limit is not None and in_cart + 1 > limit:
                await cb.answer(f"⚠️ По остаткам на складе можно не более {limit} шт.",
                                show_alert=True)
                return
            found = next((i for i in cart if i["id"] == item["id"]), None)
            if found:
                found["qty"] += 1
            else:
                cart.append({"id": item["id"], "name": item["name"], "price": item["price"],
                             "qty": 1})
            await state.update_data(precart=cart)
            await cb.answer(f"«{item['name']}» — в предзаказе ({_cart_count(cart)} шт.)")
            cat_id = _category_of_item(item["id"])
            if cat_id:
                try:
                    await cb.message.edit_reply_markup(
                        reply_markup=kb_pre_cat(cat_id, _cart_count(cart)))
                except TelegramBadRequest:
                    pass
            return
        if act == "cart":
            await cb.answer()
            await _pre_render_cart(cb, state)
            return
        if act in ("ci", "cd", "cr") and len(parts) == 3 and parts[2].isdigit():
            idx = int(parts[2])
            if not 0 <= idx < len(cart):
                await cb.answer("Корзина изменилась", show_alert=True)
                await _pre_render_cart(cb, state)
                return
            it = cart[idx]
            if act == "ci":
                if not _is_admin(uid) and _cart_count(cart) >= ORDER_MAX_ITEMS:
                    await cb.answer(f"⚠️ Не больше {ORDER_MAX_ITEMS} позиций", show_alert=True)
                    return
                limit = max_available(it["id"])
                if limit is not None and it["qty"] + 1 > limit:
                    await cb.answer(f"⚠️ По остаткам можно не более {limit} шт.",
                                    show_alert=True)
                    return
                it["qty"] += 1
            elif act == "cd":
                it["qty"] -= 1
                if it["qty"] <= 0:
                    cart.pop(idx)
            else:
                cart.pop(idx)
            await state.update_data(precart=cart)
            await cb.answer()
            await _pre_render_cart(cb, state)
            return
        if act == "cx":
            await state.update_data(precart=[])
            await cb.answer("Корзина предзаказа очищена")
            await _pre_render_cart(cb, state)
            return
        if act == "co":
            if not cart:
                await cb.answer(CART_EMPTY, show_alert=True)
                return
            await cb.answer()
            await _pre_edit(cb, pre_confirm_text(cart, pre), InlineKeyboardMarkup(
                inline_keyboard=[[_btn("✅ Отправить предзаказ", "po:so")],
                                 [_btn("🔙 К корзине", "po:cart")]]))
            return
        if act == "so":
            await _submit_preorder(cb, state)
            return
        await cb.answer()
    except Exception:
        logger.exception("Ошибка в предзаказе: %s", cb.data)
        try:
            await cb.answer("⚠️ Ошибка — попробуйте ещё раз", show_alert=True)
        except Exception:
            pass


# ------------------------------------- ID канала и служебные -----------------
@router.message(F.forward_origin)
async def on_forward(message: Message):
    origin = message.forward_origin
    chat = getattr(origin, "chat", None)
    if chat is not None:
        await message.answer(
            f"🪪 ID чата: <code>{chat.id}</code>\n\n"
            "Укажите его в Bothost как CHANNEL_ID и перезапустите бота."
        )
    else:
        await message.answer("Это не сообщение из канала. Перешлите именно пост из канала.")


async def _register_commands(bot: Bot) -> None:
    """Чистое меню «/»: гостям 5 команд, админам — 14. Вызывается при старте
    и при /help админа — список само-восстанавливается, если клиент Telegram
    показывает устаревший вариант или прошлая версия бота поставила мусор."""
    guest_cmds = [
        BotCommand(command="start", description="Меню / в начало"),
        BotCommand(command="help", description="Как пользоваться"),
        BotCommand(command="table", description="Сменить номер стола"),
        BotCommand(command="card", description="Моя карта лояльности"),
        BotCommand(command="preorder", description="Предзаказ к приходу"),
    ]
    admin_cmds = guest_cmds + [
        BotCommand(command="stats", description="Итоги дня"),
        BotCommand(command="orders", description="Последние заказы"),
        BotCommand(command="bookings", description="Ближайшие брони"),
        BotCommand(command="timing", description="Сроки выдачи заказов"),
        BotCommand(command="stock", description="Склад: остатки и себестоимость"),
        BotCommand(command="recipe", description="Рецепты позиций меню"),
        BotCommand(command="guest", description="Карта гостя: баллы по телефону/@нику"),
        BotCommand(command="loyalty", description="Аналитика программы лояльности"),
        BotCommand(command="backup", description="Резервная копия базы (файл)"),
    ]
    # Полная зачистка старых списков (от прежних версий, в любой области)
    for _scope in (
        BotCommandScopeDefault(),
        BotCommandScopeAllPrivateChats(),
        BotCommandScopeAllGroupChats(),
    ):
        try:
            await bot.delete_my_commands(scope=_scope)
        except Exception:
            pass
    for _aid in ADMIN_IDS:
        try:
            await bot.delete_my_commands(scope=BotCommandScopeChat(chat_id=_aid))
        except Exception:
            pass
    await bot.set_my_commands(guest_cmds)
    for _aid in ADMIN_IDS:
        try:
            await bot.set_my_commands(
                admin_cmds, scope=BotCommandScopeChat(chat_id=_aid)
            )
        except Exception:
            logger.exception("Не удалось поставить админ-команды для %s", _aid)
    logger.info(
        "Меню команд: гостям %d, админам %d", len(guest_cmds), len(admin_cmds)
    )


@router.message(Command("help"))
async def cmd_help(message: Message, state: FSMContext):
    table = await _table(state)
    await message.answer(
        "Как это работает:\n"
        "1. /start → выберите «Забронировать» или «Сделать заказ»\n"
        "2. Бот спросит про карту лояльности — оставьте телефон\n"
        "3. Заказ: укажите стол → выбирайте позиции → корзина → «Оформить»\n"
        "   Бронь: дата → время → гости → имя → длительность → подтверждение\n"
        "4. Заказ или бронь уйдут администратору, Вы получите смс с ответным статусом\n"
        f"5. 💎 Карта лояльности: кешбэк {LOYALTY_PERCENT}% баллами с оплаченных заказов, "
        f"оплатить баллами можно до {BONUS_MAX_PAY_PCT}% чека\n"
        "6. 🎁 Заказ по акции — в меню (отдельный заказ, баллы на него не действуют)\n"
        "7. 📝 Предзаказ (с картой) — соберите заказ заранее к времени прихода или к брони\n\n"
        "Команды: /start — в начало, /table — сменить стол, /card — моя карта, "
        "/preorder — предзаказ, "
        "📅 Бронирование — кнопка внизу",
        reply_markup=kb_main(table),
    )
    # Админ-команды показываем только админу
    if message.from_user.id in ADMIN_IDS:
        # Заодно заново ставим меню «/» — чтобы список само-восстановился,
        # если Telegram показывает устаревший вариант
        reg_ok = True
        try:
            await _register_commands(message.bot)
        except Exception:
            reg_ok = False
            logger.exception("Обновление меню команд (из /help) не удалось")
        reg_note = (
            "\n\n✅ Меню команд обновлено — наберите «/» в этом чате, "
            "список выпадет заново"
            if reg_ok
            else "\n\n⚠️ Не удалось обновить меню команд — проверьте логи Bothost"
        )
        await message.answer(
            "👤 <b>Для админа:</b>\n"
            "/stats — итоги дня\n"
            "/orders — последние заказы\n"
            "/bookings — ближайшие брони\n"
            "/timing — время выдачи заказов\n"
            "/stock — остатки, себестоимость, 📜 история прихода и списаний\n"
            "/recipe — состав позиций для списания\n"
            "/guest — карта гостя по телефону/@нику: аналитика, баллы, начислить, списать, сменить номер\n"
            "/loyalty — аналитика лояльности: топ гостей, «спящие», баллы\n"
            "/backup — резервная копия базы файлом (заказы, брони, карты, склад)\n\n"
            f"🤖 Версия бота: <b>{BOT_VERSION}</b>" + reg_note
        )


@router.message(F.text, ~F.text.startswith("/"))
async def fallback(message: Message, state: FSMContext):
    data = await state.get_data()
    table = data.get("table")
    uid = message.from_user.id if message.from_user else None
    if table and (data.get("loyal_done") or guest_card(uid)):
        await message.answer(
            "Выберите раздел каталога 👇 или /help",
            reply_markup=kb_main(table),
        )
    else:
        await message.answer(WELCOME, reply_markup=kb_welcome())


# ================================= ЗАПУСК ===================================
async def _startup_notice(bot: Bot) -> None:
    """Проблемы с файлами рядом с bot.py — сразу админам в личку."""
    problems = []
    if MENU_ERROR:
        problems.append(f"❌ Меню не загружено: {esc(MENU_ERROR)}. Каталог пуст.")
    if MENU_WARNINGS:
        problems.append("⚠️ В menu.json пропущены позиции:\n"
                        + "\n".join(f"• {esc(w)}" for w in MENU_WARNINGS[:15]))
    if STOCK_WARNINGS:
        problems.append("⚠️ " + "\n".join(esc(w) for w in STOCK_WARNINGS[:10]))
    for key, price in RSV_BASE_PRICE.items():
        label = RSV_DURATIONS.get(key, "")
        digits = re.sub(r"\D", "", label.split("-", 1)[-1])
        if digits != str(price):
            problems.append(f"💰 Цена аренды «{esc(label)}» не совпадает с RSV_BASE_PRICE "
                            f"({price} {CURRENCY}) — поправьте в НАСТРОЙКАХ.")
    for p in (PROMO_PHOTO, RULES_PHOTO):
        if not p.is_file():
            problems.append(f"🖼 Нет файла {p.name} рядом с bot.py — кнопка покажет заглушку.")
    if not problems:
        return
    for line in problems:
        logger.warning(re.sub(r"<[^>]+>", "", line))
    for admin_id in ADMIN_IDS:
        try:
            await bot.send_message(admin_id, "🤖 <b>Бот перезапущен, но есть проблемы:</b>\n\n"
                                   + "\n\n".join(problems))
        except Exception:
            logger.info("Не удалось написать админу %s о проблемах старта", admin_id)


async def main() -> None:
    # --- Предполётная проверка настроек -----------------------------------
    if not BOT_TOKEN:
        raise SystemExit(
            "❌ Не задан BOT_TOKEN — добавьте переменную окружения BOT_TOKEN "
            "(в Bothost: раздел «Переменные») и перезапустите бота"
        )
    if not re.match(r"^\d{5,}:[A-Za-z0-9_-]{20,}$", BOT_TOKEN):
        raise SystemExit(
            "❌ BOT_TOKEN имеет неверный формат (ожидается вида 123456789:AA...). "
            "Проверьте, что скопирован целиком, без пробелов и кавычек"
        )

    logger.info(
        "Настройки: канал=%s, админов=%s, валюта=%s, db=%s",
        CHANNEL_ID or "❌ НЕ ЗАДАН",
        len(ADMIN_IDS),
        CURRENCY,
        DB_PATH,
    )
    logger.info("✅ Версия бота: %s", BOT_VERSION)
    logger.info(
        "Лояльность: кешбэк %d%%, баллами до %d%% чека; лимиты: %d поз./заказ, "
        "1 заказ в %d сек, броней на гостя %d",
        LOYALTY_PERCENT, BONUS_MAX_PAY_PCT, ORDER_MAX_ITEMS, ORDER_COOLDOWN_SEC,
        RSV_MAX_ACTIVE_PER_USER,
    )
    if not ADMIN_IDS:
        logger.warning("ADMIN_IDS не задан — админ-команды (/stats, /stock, /guest…) недоступны")
    if not CHANNEL_ID:
        logger.warning("CHANNEL_ID не задан — заказы в канал отправляться не будут!")

    db_init()
    seed_stock_from_menu()
    bot = Bot(
        BOT_TOKEN,
        default=DefaultBotProperties(parse_mode=ParseMode.HTML),
    )
    storage = SQLiteStorage()
    try:
        purged = storage.purge_expired()
        if purged:
            logger.info("FSM: удалено устаревших диалогов: %d", purged)
    except Exception:
        logger.exception("Не удалось почистить устаревшие диалоги")
    dp = Dispatcher(storage=storage)
    dp.include_router(router)

    # Проверяем токен реально, до старта поллинга
    me = await bot.get_me()
    logger.info("Бот запущен: @%s (id=%s)", me.username, me.id)
    await _startup_notice(bot)

    # --- Меню команд Telegram (по ним кликают в поле ввода) ---
    try:
        await _register_commands(bot)
    except Exception:
        logger.exception("Не удалось зарегистрировать меню команд")

    await bot.delete_webhook(drop_pending_updates=True)
    logger.info("Polling запущен — бот должен отвечать на /start")
    # Страховка: исключение в хендлере не должно убивать бота насовсем
    while True:
        try:
            await dp.start_polling(bot)
            break
        except (KeyboardInterrupt, SystemExit):
            raise
        except Exception:
            logger.exception("Поллинг упал с ошибкой — перезапуск через 3 сек…")
            await asyncio.sleep(3)


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        logger.info("Остановлено")
    except SystemExit as e:
        if isinstance(e.code, str):      # ошибка настройки (нет BOT_TOKEN и т.п.) — показать причину
            logger.error(e.code)
            raise SystemExit(1)
        logger.info("Остановлено")
